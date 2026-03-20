import base64
from collections import deque
from datetime import datetime, timedelta
import hashlib
import hmac
import io
import json
import logging
from pathlib import Path
import re
import secrets
import string
import time
import traceback
import uuid
from urllib.parse import quote, urlparse

from django.conf import settings
from django.contrib.auth import authenticate
from django.contrib.auth.models import User
from django.core.cache import cache
from django.contrib.auth.password_validation import validate_password
from django.core import signing
from django.core.exceptions import ValidationError
from django.db import DataError, IntegrityError, OperationalError, connection, transaction
from django.http import HttpResponse
from drf_yasg import openapi
from drf_yasg.utils import swagger_auto_schema
from minio import Minio
from minio.error import S3Error
from rest_framework import status
from rest_framework.authtoken.models import Token
from rest_framework.decorators import api_view, authentication_classes, parser_classes, permission_classes
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import AllowAny, IsAuthenticated
from rest_framework.response import Response

from .models import WebsiteSetting

SUCCESS_CODE = '0000'
ERROR_CODE_INVALID_PARAMS = '1001'
ERROR_CODE_AUTH_FAILED = '1002'
ERROR_CODE_FORBIDDEN = '1003'
ERROR_CODE_NOT_FOUND = '1004'


def _as_bool(value, default=False):
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return value != 0
    normalized = str(value).strip().lower()
    if normalized in {'1', 'true', 'yes', 'y', 'on'}:
        return True
    if normalized in {'0', 'false', 'no', 'n', 'off'}:
        return False
    return default

REFRESH_TOKEN_SALT = 'api.refresh.token'
REFRESH_TOKEN_MAX_AGE = getattr(settings, 'REFRESH_TOKEN_MAX_AGE', 7 * 24 * 60 * 60)
CAPTCHA_TOKEN_SALT = 'api.captcha.token'
CAPTCHA_MAX_AGE = getattr(settings, 'CAPTCHA_MAX_AGE', 5 * 60)
CAPTCHA_SINGLE_USE = getattr(settings, 'CAPTCHA_SINGLE_USE', True)
LOGIN_MAX_FAIL_COUNT = getattr(settings, 'LOGIN_MAX_FAIL_COUNT', 5)
LOGIN_FAIL_WINDOW_SECONDS = getattr(settings, 'LOGIN_FAIL_WINDOW_SECONDS', 10 * 60)
LOGIN_LOCK_SECONDS = getattr(settings, 'LOGIN_LOCK_SECONDS', 15 * 60)
LOGIN_MAX_USERNAME_LENGTH = 150
LOGIN_MAX_PASSWORD_LENGTH = 128
LOGIN_MAX_CAPTCHA_ID_LENGTH = 1024
WEBSITE_NAME_MAX_LENGTH = 100
WATERMARK_TEXT_MAX_LENGTH = 100
WATERMARK_FONT_SIZE_MIN = 10
WATERMARK_FONT_SIZE_MAX = 36
WATERMARK_OPACITY_MIN = 0.05
WATERMARK_OPACITY_MAX = 0.5
WATERMARK_ROTATE_MIN = -90
WATERMARK_ROTATE_MAX = 90
WEBSITE_IMAGE_MAX_BYTES = getattr(settings, 'WEBSITE_IMAGE_MAX_BYTES', 2 * 1024 * 1024)
WEBSITE_IMAGE_MAX_LENGTH = WEBSITE_IMAGE_MAX_BYTES * 2
WEBSITE_IMAGE_ALLOWED_TYPES = {
    'image/avif',
    'image/bmp',
    'image/gif',
    'image/heic',
    'image/heif',
    'image/jpeg',
    'image/png',
    'image/svg+xml',
    'image/tiff',
    'image/vnd.microsoft.icon',
    'image/webp',
    'image/x-icon',
}
WEBSITE_IMAGE_EXTENSION_MAP = {
    'image/avif': 'avif',
    'image/bmp': 'bmp',
    'image/gif': 'gif',
    'image/heic': 'heic',
    'image/heif': 'heif',
    'image/jpeg': 'jpg',
    'image/png': 'png',
    'image/svg+xml': 'svg',
    'image/tiff': 'tiff',
    'image/vnd.microsoft.icon': 'ico',
    'image/webp': 'webp',
    'image/x-icon': 'ico',
}
WEBSITE_IMAGE_MIME_MAP = {
    'avif': 'image/avif',
    'bmp': 'image/bmp',
    'gif': 'image/gif',
    'heic': 'image/heic',
    'heif': 'image/heif',
    'ico': 'image/x-icon',
    'jfif': 'image/jpeg',
    'jpg': 'image/jpeg',
    'png': 'image/png',
    'svg': 'image/svg+xml',
    'tif': 'image/tiff',
    'tiff': 'image/tiff',
    'webp': 'image/webp',
}
WEBSITE_IMAGE_KIND_SET = {'favicon', 'logo'}
WEBSITE_ASSET_OBJECT_PREFIX = 'website-assets'
_MINIO_BUCKET_READY_CACHE_KEY = 'api:minio:bucket:ready'
WEBSITE_LOGO_AUTO_TRANSPARENT = _as_bool(getattr(settings, 'WEBSITE_LOGO_AUTO_TRANSPARENT', True), True)
WEBSITE_LOGO_TRANSPARENT_TOLERANCE = int(getattr(settings, 'WEBSITE_LOGO_TRANSPARENT_TOLERANCE', 28))
WEBSITE_LOGO_TRANSPARENT_MIN_RATIO = float(getattr(settings, 'WEBSITE_LOGO_TRANSPARENT_MIN_RATIO', 0.01))

logger = logging.getLogger(__name__)

ROLE_BUTTONS_MAP = {
    'R_SUPER': ['B_CODE1', 'B_CODE2', 'B_CODE3'],
    'R_ADMIN': ['B_CODE2', 'B_CODE3'],
    'R_USER': ['B_CODE3']
}

DATA_SCHEDULE_SCOPE_ALL = 'all'
DATA_SCHEDULE_SCOPE_COMPLETE = 'complete'
DATA_SCHEDULE_SCOPE_MISSING = 'missing'
DATA_SCHEDULE_SCOPE_SET = {
    DATA_SCHEDULE_SCOPE_ALL,
    DATA_SCHEDULE_SCOPE_COMPLETE,
    DATA_SCHEDULE_SCOPE_MISSING,
}
DATA_SCHEDULE_LOG_LEVEL_INFO = 'info'
DATA_SCHEDULE_LOG_LEVEL_WARN = 'warn'
DATA_SCHEDULE_LOG_LEVEL_ERROR = 'error'
DATA_SCHEDULE_LOG_LEVEL_SET = {
    DATA_SCHEDULE_LOG_LEVEL_INFO,
    DATA_SCHEDULE_LOG_LEVEL_WARN,
    DATA_SCHEDULE_LOG_LEVEL_ERROR,
}
DATA_SCHEDULE_LOG_ORDER_BY_TIME = 'time'
DATA_SCHEDULE_LOG_ORDER_DIRECTION_ASC = 'asc'
DATA_SCHEDULE_LOG_ORDER_DIRECTION_DESC = 'desc'
DATA_SCHEDULE_LOG_ORDER_DIRECTION_SET = {
    DATA_SCHEDULE_LOG_ORDER_DIRECTION_ASC,
    DATA_SCHEDULE_LOG_ORDER_DIRECTION_DESC,
}
DATA_SCHEDULE_LOG_DEFAULT_COUNT = 100
DATA_SCHEDULE_LOG_MAX_COUNT = 1000
DATA_SCHEDULE_TASK_DEFAULT_SIZE = 10
DATA_SCHEDULE_TASK_MAX_SIZE = 100
DATA_SCHEDULE_TASK_ACTION_MAX_SIZE = 200
DATA_SCHEDULE_TASK_STATUS_PENDING = 'pending'
DATA_SCHEDULE_TASK_STATUS_RUNNING = 'running'
DATA_SCHEDULE_TASK_STATUS_SUCCESS = 'success'
DATA_SCHEDULE_TASK_STATUS_FAILED = 'failed'
DATA_SCHEDULE_TASK_STATUS_PAUSED = 'paused'
DATA_SCHEDULE_TASK_STATUS_STOPPED = 'stopped'
DATA_SCHEDULE_TASK_STATUS_SET = {
    DATA_SCHEDULE_TASK_STATUS_PENDING,
    DATA_SCHEDULE_TASK_STATUS_RUNNING,
    DATA_SCHEDULE_TASK_STATUS_SUCCESS,
    DATA_SCHEDULE_TASK_STATUS_FAILED,
    DATA_SCHEDULE_TASK_STATUS_PAUSED,
    DATA_SCHEDULE_TASK_STATUS_STOPPED,
}
DATA_SCHEDULE_TASK_ACTION_RE_EXECUTE = 'reExecute'
DATA_SCHEDULE_TASK_ACTION_CONTINUE = 'continue'
DATA_SCHEDULE_TASK_ACTION_PAUSE = 'pause'
DATA_SCHEDULE_TASK_ACTION_STOP = 'stop'
DATA_SCHEDULE_TASK_ACTION_DELETE = 'delete'
DATA_SCHEDULE_TASK_ACTION_REFRESH = 'refresh'
DATA_SCHEDULE_TASK_ACTION_SET = {
    DATA_SCHEDULE_TASK_ACTION_RE_EXECUTE,
    DATA_SCHEDULE_TASK_ACTION_CONTINUE,
    DATA_SCHEDULE_TASK_ACTION_PAUSE,
    DATA_SCHEDULE_TASK_ACTION_STOP,
    DATA_SCHEDULE_TASK_ACTION_DELETE,
    DATA_SCHEDULE_TASK_ACTION_REFRESH,
}
DATA_SCHEDULE_LOG_TIME_FORMAT = '%Y-%m-%d %H:%M:%S'
DATA_SCHEDULE_FIELD_VALUE_MAX_LENGTH = 5000
DATA_SCHEDULE_DRILLDOWN_ROW_VALUE_MAX_LENGTH = 2000
DATA_SCHEDULE_EXPORT_LOG_DIR = Path(__file__).resolve().parents[1] / 'logs'
DATA_SCHEDULE_EXPORT_LOG_FILE = DATA_SCHEDULE_EXPORT_LOG_DIR / 'data_schedule_export.log'
API_FAILURE_LOG_FILE = DATA_SCHEDULE_EXPORT_LOG_DIR / 'api_failure.log'
DATA_SCHEDULE_REAL_SOURCE_ERROR = '数据调度真实数据源未接入，mock 兜底数据已移除'
_TABLE_COLUMN_CACHE = {}
_TABLE_COLUMN_META_CACHE = {}
_TABLE_REQUIRED_INSERT_COLUMNS_CACHE = {}
DATA_SCHEDULE_TASK_CREATOR_COLUMN_CANDIDATES = (
    'creator',
    'created_by',
    'create_by',
    'created_user',
    'create_user',
)
DATA_SCHEDULE_TASK_CREATE_TIME_COLUMN_CANDIDATES = (
    'created_at',
    'create_time',
    'created_time',
    'create_at',
)
DATA_SCHEDULE_PROJECT_CREATOR_COLUMN_CANDIDATES = (
    'creator',
    'created_by',
    'create_by',
    'owner',
)
DATA_SCHEDULE_PROJECT_CREATE_TIME_COLUMN_CANDIDATES = (
    'created_at',
    'create_time',
    'created_time',
    'create_at',
)
PROJECT_LOGIC_DELETE_COLUMN_CANDIDATES = (
    'is_deleted',
    'logic_deleted',
    'deleted',
)
DATA_SCHEDULE_DETAIL_SW = 'SW'
DATA_SCHEDULE_DETAIL_JS = 'JS'
DATA_SCHEDULE_DETAIL_INDUSTRY_SET = {
    DATA_SCHEDULE_DETAIL_SW,
    DATA_SCHEDULE_DETAIL_JS,
}
DATA_SCHEDULE_SW_PROJECT_INFO_TABLE = 'c_r_cm_engineering_sw_water_affairs_project_info'
DATA_SCHEDULE_SW_SECTION_TABLE = 'c_r_cm_engineering_sw_water_affairs_prj_section'
DATA_SCHEDULE_SW_TENDER_AGENT_TABLE = 'c_r_cm_engineering_sw_water_affairs_tender_agent'
DATA_SCHEDULE_SW_EXPERT_TABLE = 'c_r_cm_engineering_sw_water_affairs_expert_info'
DATA_SCHEDULE_SW_BID_SUBMISSION_TABLE = 'c_r_cm_engineering_sw_water_affairs_bid_submission'
DATA_SCHEDULE_JS_PROJECT_INFO_TABLE = 'c_r_cm_engineering_js_general_project_info'
DATA_SCHEDULE_JS_SECTION_TABLE = 'c_r_cm_engineering_js_general_prj_section'
DATA_SCHEDULE_JS_BID_SUBMISSION_TABLE = 'c_r_cm_engineering_js_general_bid_submission'
DATA_SCHEDULE_JS_EXPERT_TABLE = 'c_r_cm_engineering_js_general_expert_info'
DATA_SCHEDULE_JS_OPENING_ATTENDEE_TABLE = 'c_r_cm_engineering_js_general_opening_attendee'
DATA_SCHEDULE_DETAIL_EXCLUDED_COLUMN_SET = {
    'create_at',
    'create_time',
    'created_at',
    'created_time',
    'gmt_create',
    'gmt_modified',
    'update_at',
    'update_time',
    'updated_at',
    'updated_time',
}
DATA_SCHEDULE_DETAIL_NON_EDITABLE_COLUMN_SET = {
    'agent_id',
    'id',
    'pk',
    'project_id',
    'prj_section_id',
}
DATA_SCHEDULE_DRILLDOWN_NON_EDITABLE_COLUMN_SET = {
    'agent_id',
    'id',
    'pk',
    'project_id',
    'prj_section_id',
    'row_id',
    'record_id',
    'section_id',
}
DATA_SCHEDULE_DETAIL_ROW_LOCATOR_COLUMN_PRIORITY = (
    'id',
    'pk',
    'row_id',
    'record_id',
    'project_id',
    'prj_section_id',
    'section_id',
    'agent_id',
)
DATA_SCHEDULE_SELECT_DATA_ROOT = '/home/tj'
DATA_SCHEDULE_SELECT_DATA_CATALOG_ID = 'home_tj'
DATA_SCHEDULE_PROJECT_ROOT_PATH_COLUMN_CANDIDATES = (
    'root_path',
    'rootPath',
)
DATA_QUERY_DEFAULT_SIZE = 10
DATA_QUERY_MAX_SIZE = 50
DATA_QUERY_DRILLDOWN_DEFAULT_SIZE = 10
DATA_QUERY_DRILLDOWN_MAX_SIZE = 200
DATA_QUERY_EXPORT_DRILLDOWN_FIELD_SET = {
    'expert_info',
    'opening_attendee_info',
    'tender_agent_info',
}

PROJECT_STATUS_RUNNING_TASK_SET = {'RUNNING', 'CLAIMED', 'LOCKED'}
PROJECT_STATUS_FAILED_TASK_SET = {'FAILED', 'CANCELLED', 'CANCEL'}
PROJECT_STATUS_CATEGORY_COMPLETED = 'completed'
PROJECT_STATUS_CATEGORY_RUNNING = 'running'
PROJECT_STATUS_CATEGORY_REMAINING = 'remaining'
PROJECT_STATUS_CATEGORY_ABNORMAL = 'abnormal'
PROJECT_STATUS_CATEGORY_SET = {
    PROJECT_STATUS_CATEGORY_COMPLETED,
    PROJECT_STATUS_CATEGORY_RUNNING,
    PROJECT_STATUS_CATEGORY_REMAINING,
    PROJECT_STATUS_CATEGORY_ABNORMAL,
}
PROJECT_STATUS_RANKING_DEFAULT_LIMIT = 20
PROJECT_STATUS_RANKING_MAX_LIMIT = 100

SYSTEM_MANAGE_USER_DEFAULT_SIZE = 10
SYSTEM_MANAGE_USER_MAX_SIZE = 100
SYSTEM_MANAGE_USER_STATUS_ENABLED = '1'
SYSTEM_MANAGE_USER_STATUS_DISABLED = '2'


def _parse_positive_int(value, default=1, max_value=200):
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        return default

    if parsed <= 0:
        return default
    if parsed > max_value:
        return max_value

    return parsed


def _dictfetchall(cursor):
    columns = [col[0] for col in cursor.description or []]
    return [dict(zip(columns, row)) for row in cursor.fetchall()]


def _get_table_columns(table_name):
    cached = _TABLE_COLUMN_CACHE.get(table_name)
    if cached is not None:
        return cached

    try:
        with connection.cursor() as cursor:
            description = connection.introspection.get_table_description(cursor, table_name)
        columns = {
            _to_str(getattr(column, 'name', column[0])).lower()
            for column in (description or [])
            if _to_str(getattr(column, 'name', column[0]))
        }
    except Exception as exc:
        logger.warning('detect table columns failed, table=%s, error=%s', table_name, exc)
        columns = set()

    _TABLE_COLUMN_CACHE[table_name] = columns
    return columns


def _pick_existing_column(table_name, candidates):
    table_columns = _get_table_columns(table_name)
    if not table_columns:
        return None

    for column in candidates:
        normalized = _to_str(column).lower()
        if normalized and normalized in table_columns:
            return column

    return None


def _table_has_column(table_name, column_name):
    normalized_column = _to_str(column_name).lower()
    if not normalized_column:
        return False
    return normalized_column in _get_table_columns(table_name)


def _build_valid_project_where_clause(table_alias=''):
    alias_prefix = f'{table_alias}.' if table_alias else ''
    conditions = [
        f"{alias_prefix}project_id IS NOT NULL",
        f"{alias_prefix}project_id != ''",
    ]

    if _table_has_column('c_r_cm_project', 'industry'):
        conditions.append(f"{alias_prefix}industry IS NOT NULL")
        conditions.append(f"TRIM({alias_prefix}industry) != ''")

    logic_delete_column = _pick_existing_column('c_r_cm_project', PROJECT_LOGIC_DELETE_COLUMN_CANDIDATES)
    if logic_delete_column:
        conditions.append(f"COALESCE({alias_prefix}{logic_delete_column}, 0) = 0")

    return ' AND '.join(conditions)


def _normalize_data_schedule_detail_industry(raw_value):
    normalized = _to_str(raw_value).upper()
    if normalized in DATA_SCHEDULE_DETAIL_INDUSTRY_SET:
        return normalized
    return ''


def _normalize_data_schedule_project_root_path(raw_path):
    normalized = _to_str(raw_path).replace('\\', '/').strip()
    if not normalized:
        return ''

    while '//' in normalized:
        normalized = normalized.replace('//', '/')

    root_prefix = DATA_SCHEDULE_SELECT_DATA_ROOT.rstrip('/')
    normalized_no_leading = normalized.lstrip('/')
    root_no_leading = root_prefix.lstrip('/')

    if normalized == root_prefix or normalized_no_leading == root_no_leading:
        return ''

    if root_prefix and normalized.startswith(f'{root_prefix}/'):
        return normalized[len(root_prefix) + 1:].strip('/')

    if root_no_leading and normalized_no_leading.startswith(f'{root_no_leading}/'):
        return normalized_no_leading[len(root_no_leading) + 1:].strip('/')

    return normalized_no_leading.strip('/')


def _parse_data_schedule_project_root_segments(raw_path):
    relative_path = _normalize_data_schedule_project_root_path(raw_path)
    if not relative_path:
        return []
    return [item.strip() for item in relative_path.split('/') if item and item.strip()]


def _data_schedule_tree_segment_sort_key(value):
    normalized = _to_str(value)
    if normalized.isdigit():
        return 0, int(normalized)
    return 1, normalized


def _sort_data_schedule_tree_nodes(nodes):
    nodes.sort(key=lambda item: _data_schedule_tree_segment_sort_key(item.get('name')))
    for node in nodes:
        children = node.get('children') or []
        _sort_data_schedule_tree_nodes(children)


def _fill_data_schedule_tree_item_count(nodes):
    for node in nodes:
        children = node.get('children') or []
        _fill_data_schedule_tree_item_count(children)
        node['itemCount'] = len(children)


def _build_data_schedule_tree_nodes(raw_paths):
    roots = []
    node_index = {}

    for raw_path in raw_paths:
        segments = _parse_data_schedule_project_root_segments(raw_path)
        if not segments:
            continue

        parent_key = None
        siblings = roots
        for index, segment in enumerate(segments):
            node_key = '/'.join(segments[:index + 1])
            node = node_index.get(node_key)
            if node is None:
                node = {
                    'catalogId': DATA_SCHEDULE_SELECT_DATA_CATALOG_ID,
                    'children': [],
                    'id': node_key,
                    'itemCount': 0,
                    'name': segment,
                    'parentId': parent_key,
                    'type': 'folder',
                }
                node_index[node_key] = node
                siblings.append(node)

            parent_key = node_key
            children = node.get('children')
            if not isinstance(children, list):
                children = []
                node['children'] = children
            siblings = children

    _sort_data_schedule_tree_nodes(roots)
    _fill_data_schedule_tree_item_count(roots)
    return roots


def _fetch_data_schedule_project_root_paths():
    root_path_column = _pick_existing_column('c_r_cm_project', DATA_SCHEDULE_PROJECT_ROOT_PATH_COLUMN_CANDIDATES)
    if not root_path_column:
        return []

    project_scope_where_clause = _build_data_schedule_select_data_project_where_clause('p')
    sql = (
        f"SELECT DISTINCT p.{root_path_column} AS rootPath "
        "FROM c_r_cm_project p "
        f"WHERE {project_scope_where_clause} "
        f"AND p.{root_path_column} IS NOT NULL "
        f"AND TRIM(p.{root_path_column}) != ''"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return _dictfetchall(cursor)


def _build_data_schedule_select_data_project_where_clause(table_alias='p'):
    valid_project_where_clause = _build_valid_project_where_clause(table_alias)
    alias_prefix = f'{table_alias}.' if table_alias else ''
    has_file_prepare_table = _table_exists('c_r_cm_file_prepare')
    has_task_queue_table = _table_exists('c_r_cm_task_item_queue')
    if not has_file_prepare_table or not has_task_queue_table:
        return valid_project_where_clause

    reverse_conditions = [
        "("
        "NOT EXISTS ("
        "SELECT 1 FROM c_r_cm_file_prepare f "
        f"WHERE f.project_id = {alias_prefix}project_id"
        ")"
        ")",
        "("
        "EXISTS ("
        "SELECT 1 FROM c_r_cm_file_prepare f "
        f"WHERE f.project_id = {alias_prefix}project_id"
        ") "
        "AND NOT EXISTS ("
        "SELECT 1 FROM c_r_cm_task_item_queue q "
        f"WHERE q.project_id = {alias_prefix}project_id "
        "AND q.file_id LIKE 'PROJECT:%'"
        ")"
        ")",
    ]
    reverse_where_clause = '(' + ' OR '.join(reverse_conditions) + ')'
    if valid_project_where_clause:
        return f'{valid_project_where_clause} AND {reverse_where_clause}'
    return reverse_where_clause


def _fetch_data_schedule_select_data_projects():
    if not _table_exists('c_r_cm_project'):
        return []
    if not _table_has_column('c_r_cm_project', 'industry'):
        return []

    where_clause = _build_data_schedule_select_data_project_where_clause('p')
    sql = (
        "SELECT DISTINCT "
        "p.project_id AS projectId, "
        "p.project_name AS projectName, "
        "TRIM(p.industry) AS industry "
        "FROM c_r_cm_project p "
        f"WHERE {where_clause} "
        "ORDER BY UPPER(TRIM(COALESCE(p.project_name, p.project_id))) ASC, p.project_id ASC"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return _dictfetchall(cursor)


def _build_data_schedule_select_data_payload():
    rows = _fetch_data_schedule_select_data_projects()
    projects = []
    seen_project_ids = set()
    for row in rows:
        project_id = _to_str(row.get('projectId'))
        if not project_id or project_id in seen_project_ids:
            continue

        industry = _to_str(row.get('industry'))
        if not industry:
            continue

        seen_project_ids.add(project_id)
        projects.append({
            'industry': industry,
            'projectId': project_id,
            'projectName': _to_str(row.get('projectName')) or project_id,
        })

    return {'projects': projects}


def _fetch_data_query_industry_rows():
    if not _table_exists('c_r_cm_project'):
        return []

    valid_project_where_clause = _build_valid_project_where_clause('p')
    sql = (
        "SELECT DISTINCT p.industry AS industry "
        "FROM c_r_cm_project p "
        f"WHERE {valid_project_where_clause} "
        "ORDER BY UPPER(TRIM(p.industry)) ASC"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return _dictfetchall(cursor)


def _build_data_query_industry_options():
    rows = _fetch_data_query_industry_rows()
    industry_set = set()
    for row in rows:
        industry = _normalize_data_schedule_detail_industry(row.get('industry'))
        if not industry:
            continue
        industry_set.add(industry)

    options = []
    for industry in sorted(industry_set):
        options.append({'label': industry, 'value': industry})
    return options


def _parse_data_query_industry_pivot_filters(request):
    query_params = getattr(request, 'query_params', None)
    if query_params is None:
        query_params = getattr(request, 'GET', {})

    raw_industry = query_params.get('industry')
    if not _to_str(raw_industry):
        return None, 'industry 参数不能为空'

    industry = _normalize_data_schedule_detail_industry(raw_industry)
    if not industry:
        return None, 'industry 参数无效，仅支持 SW/JS'

    current = _parse_positive_int(query_params.get('current'), default=1, max_value=100000)
    size = _parse_positive_int(
        query_params.get('size'),
        default=DATA_QUERY_DEFAULT_SIZE,
        max_value=DATA_QUERY_MAX_SIZE
    )
    project_name = _to_str(query_params.get('projectName'))

    return {
        'current': current,
        'industry': industry,
        'projectName': project_name,
        'size': size,
    }, ''


def _parse_data_query_drilldown_filters(request):
    query_params = request.query_params or {}
    raw_project_id = _to_str(query_params.get('projectId'))
    if not raw_project_id:
        return None, 'projectId 参数不能为空'

    raw_industry = query_params.get('industry')
    if not _to_str(raw_industry):
        return None, 'industry 参数不能为空'
    industry = _normalize_data_schedule_detail_industry(raw_industry)
    if not industry:
        return None, 'industry 参数无效，仅支持 SW/JS'

    raw_field_key = _to_str(query_params.get('fieldKey'))
    if not raw_field_key:
        return None, 'fieldKey 参数不能为空'

    current = _parse_positive_int(query_params.get('current'), default=1, max_value=100000)
    size = _parse_positive_int(
        query_params.get('size'),
        default=DATA_QUERY_DRILLDOWN_DEFAULT_SIZE,
        max_value=DATA_QUERY_DRILLDOWN_MAX_SIZE
    )

    return {
        'current': current,
        'fieldKey': raw_field_key,
        'industry': industry,
        'projectId': raw_project_id,
        'size': size,
    }, ''


def _fetch_data_query_project_relation_map(file_prepare_project_ids):
    normalized_project_ids = [_to_str(item) for item in (file_prepare_project_ids or []) if _to_str(item)]
    if not normalized_project_ids:
        return {}
    if not _table_exists('c_r_cm_kb_project_relation'):
        return {}

    where_clause, params = _build_data_schedule_where_with_values('file_prepare_project_id', normalized_project_ids)
    if where_clause is None:
        return {}

    order_parts = []
    for column_name in ('update_time', 'create_time', 'id'):
        if _table_has_column('c_r_cm_kb_project_relation', column_name):
            order_parts.append(f'{column_name} DESC')
    order_sql = f" ORDER BY {', '.join(order_parts)}" if order_parts else ''

    sql = (
        "SELECT file_prepare_project_id AS filePrepareProjectId, "
        "project_id AS projectId "
        "FROM c_r_cm_kb_project_relation "
        f"WHERE {where_clause} "
        f"{order_sql}"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        rows = _dictfetchall(cursor)

    relation_map = {}
    for row in rows:
        file_prepare_project_id = _to_str(row.get('filePrepareProjectId'))
        kb_project_id = _to_str(row.get('projectId'))
        if not file_prepare_project_id or not kb_project_id:
            continue
        if file_prepare_project_id in relation_map:
            continue
        relation_map[file_prepare_project_id] = kb_project_id
    return relation_map


def _fetch_data_query_projects_by_industry(industry, project_name=''):
    normalized_industry = _normalize_data_schedule_detail_industry(industry)
    if not normalized_industry:
        return []

    valid_project_where_clause = _build_valid_project_where_clause('p')
    where_parts = [
        valid_project_where_clause,
        "UPPER(TRIM(p.industry)) = %s",
    ]
    params = [normalized_industry]

    normalized_project_name = _to_str(project_name)
    if normalized_project_name:
        where_parts.append("p.project_name LIKE %s")
        params.append(f'%{normalized_project_name}%')

    sql = (
        "SELECT p.project_id AS projectId, "
        "p.project_name AS projectName "
        "FROM c_r_cm_project p "
        f"WHERE {' AND '.join(where_parts)} "
        "ORDER BY p.project_name ASC, p.project_id ASC"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        rows = _dictfetchall(cursor)

    file_prepare_project_ids = [_to_str(item.get('projectId')) for item in rows if _to_str(item.get('projectId'))]
    relation_map = _fetch_data_query_project_relation_map(file_prepare_project_ids)

    projects = []
    for row in rows:
        file_prepare_project_id = _to_str(row.get('projectId'))
        if not file_prepare_project_id:
            continue

        projects.append({
            'kbProjectId': relation_map.get(file_prepare_project_id, ''),
            'projectId': file_prepare_project_id,
            'projectName': _to_str(row.get('projectName')) or file_prepare_project_id,
        })
    return projects


def _build_data_query_project_field_value_map(project_row, industry):
    kb_project_id = _to_str((project_row or {}).get('kbProjectId'))
    if not kb_project_id:
        return {}

    config = _get_data_schedule_industry_table_config(industry)
    if not config:
        return {}

    field_rows = _build_data_schedule_field_rows({
        'config': config,
        'projectId': kb_project_id,
    })

    field_name_order = []
    field_values_map = {}
    for item in field_rows:
        field_name = _to_str(item.get('fieldName'))
        if not field_name:
            continue
        if field_name not in field_values_map:
            field_name_order.append(field_name)
            field_values_map[field_name] = []

        field_value = _to_str(item.get('fieldValueDisplay')) or '--'
        if _is_blank_text_value(field_value):
            continue
        if field_value in field_values_map[field_name]:
            continue
        field_values_map[field_name].append(field_value)

    flattened = {}
    for field_name in field_name_order:
        values = field_values_map.get(field_name) or []
        flattened[field_name] = '；'.join(values) if values else '--'
    return flattened


def _normalize_data_query_drilldown_field_key(raw_field_key, config):
    field_key = _to_str(raw_field_key)
    if not field_key:
        return ''

    expert_field_key = _to_str((config or {}).get('expertFieldKey'))
    opening_attendee_field_key = _to_str((config or {}).get('openingAttendeeFieldKey'))
    bid_submission_field_key = _to_str((config or {}).get('bidSubmissionFieldKey'))
    alias_map = {
        '招标代理机构': 'tender_agent_info',
        '评标专家信息': expert_field_key,
        '开标人员信息': opening_attendee_field_key,
        '中标候选人公示': bid_submission_field_key,
    }
    mapped_field_key = _to_str(alias_map.get(field_key))
    return mapped_field_key or field_key


def _resolve_data_query_drilldown_table(field_key, config):
    normalized_field_key = _to_str(field_key)
    if not normalized_field_key:
        return '', False

    if normalized_field_key == 'tender_agent_info':
        return _to_str((config or {}).get('tenderAgentTable')), True
    if normalized_field_key == _to_str((config or {}).get('expertFieldKey')):
        return _to_str((config or {}).get('expertTable')), False
    if normalized_field_key == _to_str((config or {}).get('bidSubmissionFieldKey')):
        return _to_str((config or {}).get('bidSubmissionTable')), False
    if normalized_field_key == _to_str((config or {}).get('openingAttendeeFieldKey')):
        return _to_str((config or {}).get('openingAttendeeTable')), False
    return '', False


def _build_data_query_drilldown_title(field_key, config):
    normalized_field_key = _to_str(field_key)
    if normalized_field_key == 'tender_agent_info':
        return '字段详情 - 招标代理机构'
    if normalized_field_key == _to_str((config or {}).get('expertFieldKey')):
        return f"字段详情 - {_to_str((config or {}).get('expertFieldName')) or '评标专家信息'}"
    if normalized_field_key == _to_str((config or {}).get('bidSubmissionFieldKey')):
        return f"字段详情 - {_to_str((config or {}).get('bidSubmissionFieldName')) or '中标候选人公示'}"
    if normalized_field_key == _to_str((config or {}).get('openingAttendeeFieldKey')):
        return f"字段详情 - {_to_str((config or {}).get('openingAttendeeFieldName')) or '开标人员信息'}"
    return f'字段详情 - {normalized_field_key}'


def _to_readonly_data_query_drilldown_columns(columns):
    readonly_columns = []
    for column in (columns or []):
        if not isinstance(column, dict):
            continue
        readonly_columns.append({
            'editable': False,
            'key': _to_str(column.get('key')),
            'title': _to_str(column.get('title')),
            'width': column.get('width'),
        })
    return readonly_columns


def _build_data_query_drilldown_payload(filters):
    project_id = _to_str((filters or {}).get('projectId'))
    industry = _to_str((filters or {}).get('industry'))
    raw_field_key = _to_str((filters or {}).get('fieldKey'))
    current = _to_int((filters or {}).get('current'), default=1)
    size = _to_int((filters or {}).get('size'), default=DATA_QUERY_DRILLDOWN_DEFAULT_SIZE)

    config = _get_data_schedule_industry_table_config(industry)
    if not config:
        raise RuntimeError(f'暂不支持行业 {industry} 的详情查询')

    relation_map = _fetch_data_query_project_relation_map([project_id])
    kb_project_id = _to_str(relation_map.get(project_id))
    if not kb_project_id:
        raise ValueError('项目映射不存在')

    context = {
        'agentIds': set(),
        'config': config,
        'projectId': kb_project_id,
        'sectionIds': set(),
    }
    # Initialize section/agent scopes used by tender-agent drilldown filtering.
    _build_data_schedule_field_rows(context)

    field_key = _normalize_data_query_drilldown_field_key(raw_field_key, config)
    table_name, prefer_agent = _resolve_data_query_drilldown_table(field_key, config)
    if not table_name:
        raise ValueError('下钻字段不存在')

    title = _build_data_query_drilldown_title(field_key, config)
    payload = _build_data_schedule_drilldown_from_table(table_name, title, context, prefer_agent=prefer_agent)
    all_records = payload.get('records') or []
    paginated_data = _paginate_data_schedule_records(all_records, current, size)

    return {
        'actions': {
            'canCreate': False,
            'canDelete': False,
            'canEdit': False,
        },
        'columns': _to_readonly_data_query_drilldown_columns(payload.get('columns') or []),
        'current': paginated_data.get('current') or current,
        'fieldKey': field_key,
        'records': paginated_data.get('records') or [],
        'size': paginated_data.get('size') or size,
        'title': payload.get('title') or title,
        'total': paginated_data.get('total') or 0,
    }


def _build_data_query_pivot_columns_and_records(project_rows, industry):
    dynamic_field_keys = []
    dynamic_field_set = set()
    records = []

    for project_row in (project_rows or []):
        project_field_value_map = _build_data_query_project_field_value_map(project_row, industry)
        record = {
            'projectId': _to_str(project_row.get('projectId')),
            'projectName': _to_str(project_row.get('projectName')) or _to_str(project_row.get('projectId')),
        }

        for field_name, field_value in project_field_value_map.items():
            if field_name not in dynamic_field_set:
                dynamic_field_set.add(field_name)
                dynamic_field_keys.append(field_name)
            record[field_name] = field_value if not _is_blank_text_value(field_value) else '--'

        records.append(record)

    for record in records:
        for field_name in dynamic_field_keys:
            if field_name not in record:
                record[field_name] = '--'

    columns = [{'fixed': 'left', 'key': 'projectName', 'title': '项目名称'}]
    columns.extend([{'key': field_name, 'title': field_name} for field_name in dynamic_field_keys])

    return columns, records


def _build_data_query_industry_pivot_payload(filters):
    current = _to_int((filters or {}).get('current'), default=1)
    size = _to_int((filters or {}).get('size'), default=DATA_QUERY_DEFAULT_SIZE)
    industry = _to_str((filters or {}).get('industry'))
    project_name = _to_str((filters or {}).get('projectName'))

    all_projects = _fetch_data_query_projects_by_industry(industry, project_name)
    total = len(all_projects)
    start_idx = max(0, (current - 1) * size)
    end_idx = start_idx + size
    paged_projects = all_projects[start_idx:end_idx]

    columns, records = _build_data_query_pivot_columns_and_records(paged_projects, industry)

    return {
        'columns': columns,
        'current': current,
        'records': records,
        'size': size,
        'total': total,
    }


def _build_data_query_export_binary(filters):
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.hyperlink import Hyperlink
    except Exception as exc:
        logger.exception('openpyxl import failed when export data query pivot: %s', exc)
        raise

    industry = _to_str((filters or {}).get('industry'))
    project_name = _to_str((filters or {}).get('projectName'))
    config = _get_data_schedule_industry_table_config(industry) or {}
    all_projects = _fetch_data_query_projects_by_industry(industry, project_name)
    columns, records = _build_data_query_pivot_columns_and_records(all_projects, industry)

    workbook = Workbook()
    used_sheet_names = set()
    sheet = workbook.active
    sheet.title = _safe_excel_sheet_name('数据查询', used_sheet_names)
    sheet.freeze_panes = 'A2'

    sheet.append([_to_str(column.get('title')) or _to_str(column.get('key')) for column in columns])

    drilldown_sheet_map = {}
    for record in records:
        project_id = _to_str(record.get('projectId'))
        project_name_text = _to_str(record.get('projectName')) or project_id
        if not project_id:
            continue

        for column in columns:
            column_key = _to_str(column.get('key'))
            field_key = _normalize_data_query_drilldown_field_key(column_key, config)
            if field_key not in DATA_QUERY_EXPORT_DRILLDOWN_FIELD_SET:
                continue

            map_key = (project_id, column_key)
            if map_key in drilldown_sheet_map:
                continue

            try:
                drilldown_payload = _build_data_query_drilldown_payload(
                    {
                        'current': 1,
                        'fieldKey': field_key,
                        'industry': industry,
                        'projectId': project_id,
                        'size': DATA_QUERY_DRILLDOWN_MAX_SIZE,
                    }
                )
            except Exception:
                continue

            drilldown_columns = drilldown_payload.get('columns') or []
            drilldown_records = drilldown_payload.get('records') or []
            drilldown_title = _to_str(drilldown_payload.get('title')) or _to_str(column.get('title')) or column_key
            sheet_name = _safe_excel_sheet_name(f'{project_name_text}_{drilldown_title}', used_sheet_names)
            drilldown_sheet_map[map_key] = sheet_name

            drilldown_sheet = workbook.create_sheet(sheet_name)
            drilldown_sheet.append([_to_str(item.get('title')) or _to_str(item.get('key')) for item in drilldown_columns])
            for drilldown_row in drilldown_records:
                row_values = []
                for drilldown_column in drilldown_columns:
                    drilldown_column_key = _to_str(drilldown_column.get('key'))
                    row_values.append(_to_excel_cell_value(drilldown_row.get(drilldown_column_key)))
                drilldown_sheet.append(row_values)
            drilldown_sheet.freeze_panes = 'A2'
            _style_data_schedule_drilldown_sheet(drilldown_sheet)

    for row_index, record in enumerate(records, start=2):
        row_values = []
        project_id = _to_str(record.get('projectId'))
        for column in columns:
            column_key = _to_str(column.get('key'))
            map_key = (project_id, column_key)
            if map_key in drilldown_sheet_map:
                row_values.append('查看')
            else:
                row_values.append(_to_excel_cell_value(record.get(column_key)))
        sheet.append(row_values)

        for col_index, column in enumerate(columns, start=1):
            column_key = _to_str(column.get('key'))
            map_key = (project_id, column_key)
            if map_key not in drilldown_sheet_map:
                continue
            link_cell = sheet.cell(row=row_index, column=col_index)
            link_cell.value = '查看'
            link_cell.hyperlink = Hyperlink(
                ref=link_cell.coordinate,
                location=_build_excel_internal_link(drilldown_sheet_map[map_key], 'A1').lstrip('#'),
                display='查看'
            )

    _style_data_query_pivot_export_sheet(sheet)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output.getvalue()


def _table_exists(table_name):
    return bool(_get_table_columns(table_name))


def _get_table_column_meta(table_name):
    cached = _TABLE_COLUMN_META_CACHE.get(table_name)
    if cached is not None:
        return cached

    if not _table_exists(table_name):
        _TABLE_COLUMN_META_CACHE[table_name] = []
        return []

    rows = []
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COLUMN_NAME AS columnName, "
                "COALESCE(COLUMN_COMMENT, '') AS columnComment, "
                "ORDINAL_POSITION AS ordinalPosition "
                "FROM INFORMATION_SCHEMA.COLUMNS "
                "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s "
                "ORDER BY ORDINAL_POSITION ASC",
                [table_name]
            )
            rows = _dictfetchall(cursor)
    except Exception as exc:
        logger.warning('query table column meta failed, table=%s, error=%s', table_name, exc)

    if rows:
        normalized = []
        for row in rows:
            column_name = _to_str(row.get('columnName'))
            if not column_name:
                continue
            normalized.append({
                'comment': _to_str(row.get('columnComment')),
                'name': column_name,
                'normalizedName': column_name.lower(),
                'ordinalPosition': _to_int(row.get('ordinalPosition'), default=0),
            })
        _TABLE_COLUMN_META_CACHE[table_name] = normalized
        return normalized

    fallback_meta = []
    for index, column_name in enumerate(sorted(_get_table_columns(table_name)), start=1):
        fallback_meta.append({
            'comment': '',
            'name': column_name,
            'normalizedName': column_name.lower(),
            'ordinalPosition': index,
        })
    _TABLE_COLUMN_META_CACHE[table_name] = fallback_meta
    return fallback_meta


def _get_table_required_insert_columns(table_name):
    cached = _TABLE_REQUIRED_INSERT_COLUMNS_CACHE.get(table_name)
    if cached is not None:
        return cached

    if not _table_exists(table_name):
        _TABLE_REQUIRED_INSERT_COLUMNS_CACHE[table_name] = []
        return []

    rows = []
    try:
        with connection.cursor() as cursor:
            cursor.execute(
                "SELECT COLUMN_NAME AS columnName, "
                "COALESCE(IS_NULLABLE, 'YES') AS isNullable, "
                "COLUMN_DEFAULT AS columnDefault, "
                "COALESCE(EXTRA, '') AS extraInfo "
                "FROM INFORMATION_SCHEMA.COLUMNS "
                "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s "
                "ORDER BY ORDINAL_POSITION ASC",
                [table_name]
            )
            rows = _dictfetchall(cursor)
    except Exception as exc:
        logger.warning('query table required insert columns failed, table=%s, error=%s', table_name, exc)

    required_columns = []
    for row in rows:
        column_name = _to_str(row.get('columnName'))
        if not column_name:
            continue

        is_nullable = _to_str(row.get('isNullable')).upper()
        has_default = row.get('columnDefault') is not None
        extra_info = _to_str(row.get('extraInfo')).lower()

        if is_nullable != 'NO':
            continue
        if has_default:
            continue
        if 'auto_increment' in extra_info:
            continue
        if 'generated' in extra_info:
            continue
        required_columns.append(column_name)

    _TABLE_REQUIRED_INSERT_COLUMNS_CACHE[table_name] = required_columns
    return required_columns


def _is_blank_text_value(value):
    if value is None:
        return True
    text = str(value).strip()
    return text in {'', '--'}


def _format_data_schedule_field_value(value):
    if value is None:
        return '--'
    if isinstance(value, bool):
        return '是' if value else '否'
    if isinstance(value, datetime):
        return value.strftime(DATA_SCHEDULE_LOG_TIME_FORMAT)
    text = str(value).strip()
    return text if text else '--'


def _build_data_schedule_field_name(column_meta):
    comment = _to_str((column_meta or {}).get('comment'))
    if comment:
        return comment
    return _to_str((column_meta or {}).get('name'))


def _resolve_data_schedule_project_id_from_task_id(task_id):
    task_key = _to_str(task_id)
    if not task_key:
        return ''

    normalized = task_key.upper()
    if normalized.startswith('PROJECT:'):
        return task_key.split(':', 1)[1].strip()

    sql = (
        "SELECT project_id AS projectId "
        "FROM c_r_cm_task_item_queue "
        "WHERE resource_id = %s "
        "ORDER BY COALESCE(locked_at, updated_at) DESC, updated_at DESC "
        "LIMIT 1"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, [task_key])
        rows = _dictfetchall(cursor)

    if rows:
        project_id = _to_str(rows[0].get('projectId'))
        if project_id:
            return project_id

    return task_key


def _fetch_data_schedule_project_base_row(project_id):
    normalized_project_id = _to_str(project_id)
    if not normalized_project_id:
        return None

    valid_project_where_clause = _build_valid_project_where_clause('p')
    sql = (
        "SELECT p.project_id AS projectId, "
        "p.project_name AS projectName, "
        "p.industry AS industry "
        "FROM c_r_cm_project p "
        f"WHERE {valid_project_where_clause} AND p.project_id = %s "
        "LIMIT 1"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, [normalized_project_id])
        rows = _dictfetchall(cursor)

    if not rows:
        return None
    return rows[0]


def _resolve_data_schedule_kb_project_id_from_file_prepare_project_id(file_prepare_project_id):
    normalized_project_id = _to_str(file_prepare_project_id)
    if not normalized_project_id:
        return ''
    if not _table_exists('c_r_cm_kb_project_relation'):
        return ''

    sql = (
        "SELECT project_id AS projectId "
        "FROM c_r_cm_kb_project_relation "
        "WHERE file_prepare_project_id = %s "
        "ORDER BY update_time DESC, create_time DESC, id DESC "
        "LIMIT 1"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, [normalized_project_id])
        rows = _dictfetchall(cursor)
    if not rows:
        return ''
    return _to_str(rows[0].get('projectId'))


def _find_data_schedule_task_row(task_id):
    normalized_task_id = _to_str(task_id)
    if not normalized_task_id:
        return None

    rows = _build_data_schedule_task_rows()
    for row in rows:
        if _to_str(row.get('taskId')) == normalized_task_id:
            return row

    resolved_project_id = _resolve_data_schedule_project_id_from_task_id(normalized_task_id)
    for row in rows:
        project_id = _to_str(row.get('projectId'))
        if project_id and project_id == resolved_project_id:
            return row
        if project_id and normalized_task_id.upper() == f'PROJECT:{project_id}'.upper():
            return row
    return None


def _fetch_data_schedule_table_records(table_name, where_parts=None, params=None):
    if not _table_exists(table_name):
        return []

    column_meta = _get_table_column_meta(table_name)
    column_names = [item.get('name') for item in column_meta if _to_str(item.get('name'))]
    if not column_names:
        return []

    where_sql = ''
    if where_parts:
        where_sql = ' WHERE ' + ' AND '.join(where_parts)

    select_sql = ', '.join([f"`{name}`" for name in column_names])
    sql = f"SELECT {select_sql} FROM `{table_name}`{where_sql}"
    with connection.cursor() as cursor:
        cursor.execute(sql, params or [])
        rows = _dictfetchall(cursor)
    return rows


def _normalize_data_schedule_detail_rows(table_name, rows):
    column_meta = _get_table_column_meta(table_name)
    available_columns = {
        _to_str(item.get('normalizedName')).lower()
        for item in column_meta
        if _to_str(item.get('normalizedName'))
    }
    normalized_to_column_name = {}
    for item in column_meta:
        normalized_name = _to_str(item.get('normalizedName')).lower()
        column_name = _to_str(item.get('name'))
        if normalized_name and column_name and normalized_name not in normalized_to_column_name:
            normalized_to_column_name[normalized_name] = column_name

    included_meta = [
        item
        for item in column_meta
        if item.get('normalizedName') not in DATA_SCHEDULE_DETAIL_EXCLUDED_COLUMN_SET
    ]

    field_rows = []
    agent_ids = set()
    section_ids = set()
    for row_index, row in enumerate(rows):
        row_data = row or {}
        agent_id_value = row_data.get('agent_id') if 'agent_id' in available_columns else None
        if not _is_blank_text_value(agent_id_value):
            agent_ids.add(_to_str(agent_id_value))

        section_id_value = row_data.get('prj_section_id') if 'prj_section_id' in available_columns else None
        if not _is_blank_text_value(section_id_value):
            section_ids.add(_to_str(section_id_value))

        source_locator = _build_data_schedule_row_locator(normalized_to_column_name, row_data)
        for meta in included_meta:
            column_name = _to_str(meta.get('name'))
            normalized_name = _to_str(meta.get('normalizedName')).lower()
            raw_value = row_data.get(column_name)
            field_name = _build_data_schedule_field_name(meta)

            if normalized_name == 'agent_id' or field_name in {'招标代理机构ID', '招标代理机构Id', '招标代理机构id'}:
                continue

            value_display = _format_data_schedule_field_value(raw_value)
            status_value = DATA_SCHEDULE_SCOPE_MISSING if _is_blank_text_value(value_display) else DATA_SCHEDULE_SCOPE_COMPLETE
            can_edit = (
                bool(source_locator)
                and _is_data_schedule_detail_column_editable(normalized_name)
            )
            field_rows.append({
                'canDrilldown': False,
                'canEdit': can_edit,
                'fieldKey': _build_data_schedule_field_key(table_name, column_name, row_index, source_locator),
                'fieldName': field_name or column_name,
                'fieldValueDisplay': value_display,
                'sourceColumn': column_name,
                'sourceLocator': source_locator,
                'sourceTable': table_name,
                'status': status_value,
            })

    return field_rows, agent_ids, section_ids


def _build_data_schedule_row_locator(normalized_to_column_name, row_data):
    locator = []
    used_columns = set()
    has_strong_unique_key = False

    for normalized_name in DATA_SCHEDULE_DETAIL_ROW_LOCATOR_COLUMN_PRIORITY:
        column_name = _to_str((normalized_to_column_name or {}).get(normalized_name))
        if not column_name or column_name in used_columns:
            continue

        raw_value = (row_data or {}).get(column_name)
        if _is_blank_text_value(raw_value):
            continue

        locator.append({'column': column_name, 'value': raw_value})
        used_columns.add(column_name)
        if normalized_name in {'id', 'pk', 'row_id', 'record_id'}:
            has_strong_unique_key = True

    if has_strong_unique_key:
        return locator

    for normalized_name in sorted((normalized_to_column_name or {}).keys()):
        if not normalized_name.endswith('_id'):
            continue
        column_name = _to_str((normalized_to_column_name or {}).get(normalized_name))
        if not column_name or column_name in used_columns:
            continue

        raw_value = (row_data or {}).get(column_name)
        if _is_blank_text_value(raw_value):
            continue

        locator.append({'column': column_name, 'value': raw_value})
        used_columns.add(column_name)

    for normalized_name in sorted((normalized_to_column_name or {}).keys()):
        column_name = _to_str((normalized_to_column_name or {}).get(normalized_name))
        if not column_name or column_name in used_columns:
            continue

        raw_value = (row_data or {}).get(column_name)
        if _is_blank_text_value(raw_value):
            continue

        locator.append({'column': column_name, 'value': raw_value})
        used_columns.add(column_name)

    return locator


def _build_data_schedule_field_key(table_name, column_name, row_index, source_locator):
    normalized_table = _to_str(table_name)
    normalized_column = _to_str(column_name)
    if not normalized_table or not normalized_column:
        return ''

    normalized_locator = []
    for locator_item in source_locator or []:
        if not isinstance(locator_item, dict):
            continue
        locator_column = _to_str(locator_item.get('column'))
        locator_value = locator_item.get('value')
        if not locator_column or _is_blank_text_value(locator_value):
            continue
        normalized_locator.append({
            'column': locator_column,
            'value': str(locator_value),
        })
    if normalized_locator:
        locator_text = json.dumps(normalized_locator, ensure_ascii=False, separators=(',', ':'))
        digest = hashlib.md5(locator_text.encode('utf-8')).hexdigest()[:12]
        return f'{normalized_table}.{normalized_column}.{digest}'

    return f'{normalized_table}.{normalized_column}.{row_index}'


def _is_data_schedule_detail_column_editable(normalized_column_name):
    normalized = _to_str(normalized_column_name).lower()
    if not normalized:
        return False
    if normalized in DATA_SCHEDULE_DETAIL_EXCLUDED_COLUMN_SET:
        return False
    if normalized in DATA_SCHEDULE_DETAIL_NON_EDITABLE_COLUMN_SET:
        return False
    if normalized.endswith('_id'):
        return False
    return True


def _is_data_schedule_drilldown_column_editable(normalized_column_name):
    normalized = _to_str(normalized_column_name).lower()
    if not normalized:
        return False
    if normalized in DATA_SCHEDULE_DETAIL_EXCLUDED_COLUMN_SET:
        return False
    if normalized in DATA_SCHEDULE_DRILLDOWN_NON_EDITABLE_COLUMN_SET:
        return False
    return True


def _is_safe_sql_identifier(name):
    normalized = _to_str(name)
    if not normalized:
        return False
    if not (normalized[0].isalpha() or normalized[0] == '_'):
        return False
    for char in normalized:
        if not (char.isalnum() or char == '_'):
            return False
    return True


def _to_int(value, default=0):
    if value is None:
        return default
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        return int(value)
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


def _to_str(value, default=''):
    if value is None:
        return default
    text = str(value).strip()
    return text if text else default


def _format_datetime_text(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.strftime(DATA_SCHEDULE_LOG_TIME_FORMAT)
    return str(value)


def _parse_datetime_value(value):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value

    text = str(value).strip()
    if not text:
        return None

    for fmt in (DATA_SCHEDULE_LOG_TIME_FORMAT, '%Y-%m-%d'):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            pass

    try:
        parsed = datetime.fromisoformat(text.replace('Z', '+00:00'))
        if parsed.tzinfo is not None:
            parsed = parsed.astimezone().replace(tzinfo=None)
        return parsed
    except ValueError:
        return None


def _fetch_project_name_rows():
    valid_project_where_clause = _build_valid_project_where_clause()
    sql = (
        "SELECT project_id AS projectId, project_name AS projectName "
        "FROM c_r_cm_project "
        f"WHERE {valid_project_where_clause}"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return _dictfetchall(cursor)


def _fetch_project_file_snapshot_rows():
    valid_project_where_clause = _build_valid_project_where_clause('p')
    sql = (
        "SELECT f.project_id AS projectId, "
        "COUNT(*) AS totalFiles, "
        "SUM(CASE WHEN f.status <> 'DRAFT' THEN 1 ELSE 0 END) AS nonDraftFiles, "
        "SUM(CASE WHEN f.status = 'DRAFT' THEN 1 ELSE 0 END) AS draftFiles, "
        "SUM(CASE WHEN f.status = 'PENDING' THEN 1 ELSE 0 END) AS pendingFiles, "
        "SUM(CASE WHEN f.status = 'RUNNING' THEN 1 ELSE 0 END) AS runningFiles, "
        "SUM(CASE WHEN f.status = 'SUCCESS' THEN 1 ELSE 0 END) AS successFiles, "
        "SUM(CASE WHEN f.status = 'FAILED' THEN 1 ELSE 0 END) AS failedFiles, "
        "SUM(CASE WHEN f.status = 'PARTIAL_FAILED' THEN 1 ELSE 0 END) AS partialFailedFiles, "
        "SUM(CASE WHEN f.status = 'SUCCESS' AND f.current_step = 3 THEN 1 ELSE 0 END) AS step3SuccessFiles, "
        "MAX(COALESCE(f.current_step, 0)) AS maxCurrentStep, "
        "MAX(CASE WHEN f.status = 'RUNNING' THEN COALESCE(f.current_step, 0) ELSE 0 END) AS runningStepNo, "
        "SUBSTRING_INDEX( "
        "  GROUP_CONCAT( "
        "    CASE WHEN f.last_error IS NOT NULL AND f.last_error != '' THEN f.last_error ELSE NULL END "
        "    ORDER BY f.updated_at DESC SEPARATOR '\n' "
        "  ), "
        "  '\n', "
        "  1 "
        ") AS latestFileError "
        "FROM c_r_cm_file_prepare f "
        "JOIN c_r_cm_project p ON p.project_id = f.project_id "
        f"WHERE {valid_project_where_clause} "
        "GROUP BY f.project_id"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return _dictfetchall(cursor)


def _fetch_latest_project_task_rows():
    creator_column = _pick_existing_column('c_r_cm_task_item_queue', DATA_SCHEDULE_TASK_CREATOR_COLUMN_CANDIDATES)
    create_time_column = _pick_existing_column('c_r_cm_task_item_queue', DATA_SCHEDULE_TASK_CREATE_TIME_COLUMN_CANDIDATES)

    select_parts = [
        "t.project_id AS projectId",
        "t.resource_id AS taskId",
        "t.step_no AS stepNo",
        "t.status AS status",
        "t.last_error AS lastError",
        "t.updated_at AS updatedAt",
        f"t.{creator_column} AS creator" if creator_column else "NULL AS creator",
        f"t.{create_time_column} AS createdAt" if create_time_column else "NULL AS createdAt",
    ]
    valid_project_where_clause = _build_valid_project_where_clause('p')
    sql = (
        "SELECT "
        + ", ".join(select_parts)
        + " FROM c_r_cm_task_item_queue t "
        "JOIN c_r_cm_project p ON p.project_id = t.project_id "
        "JOIN ( "
        "    SELECT project_id, "
        "           SUBSTRING_INDEX( "
        "               GROUP_CONCAT( "
        "                   resource_id "
        "                   ORDER BY COALESCE(locked_at, updated_at) DESC, updated_at DESC, resource_id DESC "
        "                   SEPARATOR ',' "
        "               ), "
        "               ',', "
        "               1 "
        "           ) AS latest_resource_id "
        "    FROM c_r_cm_task_item_queue "
        "    WHERE file_id LIKE 'PROJECT:%' "
        "    GROUP BY project_id "
        ") latest ON latest.latest_resource_id = t.resource_id "
        f"WHERE t.file_id LIKE 'PROJECT:%' AND {valid_project_where_clause}"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return _dictfetchall(cursor)


def _empty_file_snapshot(project_id):
    return {
        'projectId': project_id,
        'totalFiles': 0,
        'nonDraftFiles': 0,
        'draftFiles': 0,
        'pendingFiles': 0,
        'runningFiles': 0,
        'successFiles': 0,
        'failedFiles': 0,
        'partialFailedFiles': 0,
        'step3SuccessFiles': 0,
        'maxCurrentStep': 0,
        'runningStepNo': 0,
        'latestFileError': '',
    }


def _normalize_file_snapshot(row):
    project_id = _to_str(row.get('projectId'))
    return {
        'projectId': project_id,
        'totalFiles': _to_int(row.get('totalFiles')),
        'nonDraftFiles': _to_int(row.get('nonDraftFiles')),
        'draftFiles': _to_int(row.get('draftFiles')),
        'pendingFiles': _to_int(row.get('pendingFiles')),
        'runningFiles': _to_int(row.get('runningFiles')),
        'successFiles': _to_int(row.get('successFiles')),
        'failedFiles': _to_int(row.get('failedFiles')),
        'partialFailedFiles': _to_int(row.get('partialFailedFiles')),
        'step3SuccessFiles': _to_int(row.get('step3SuccessFiles')),
        'maxCurrentStep': _to_int(row.get('maxCurrentStep')),
        'runningStepNo': _to_int(row.get('runningStepNo')),
        'latestFileError': _to_str(row.get('latestFileError')),
    }


def _normalize_task_snapshot(row):
    return {
        'projectId': _to_str(row.get('projectId')),
        'taskId': _to_str(row.get('taskId')),
        'stepNo': _to_int(row.get('stepNo')),
        'status': _to_str(row.get('status')).upper(),
        'lastError': _to_str(row.get('lastError')),
        'creator': _to_str(row.get('creator')),
        'createdAt': _format_datetime_text(row.get('createdAt')),
        'updatedAt': _format_datetime_text(row.get('updatedAt')),
    }


def _resolve_project_current_step(file_snapshot, task_snapshot):
    running_step_no = _to_int(file_snapshot.get('runningStepNo'))
    if running_step_no > 0:
        return running_step_no

    if task_snapshot:
        task_step_no = _to_int(task_snapshot.get('stepNo'))
        if task_step_no > 0:
            return task_step_no

    max_current_step = _to_int(file_snapshot.get('maxCurrentStep'))
    if max_current_step > 0:
        return max_current_step

    return 0


def _resolve_project_status(file_snapshot, task_snapshot):
    running_files = _to_int(file_snapshot.get('runningFiles'))
    non_draft_files = _to_int(file_snapshot.get('nonDraftFiles'))
    step3_success_files = _to_int(file_snapshot.get('step3SuccessFiles'))
    failed_files = _to_int(file_snapshot.get('failedFiles'))
    partial_failed_files = _to_int(file_snapshot.get('partialFailedFiles'))
    failed_like_files = failed_files + partial_failed_files

    task_status = ''
    if task_snapshot:
        task_status = _to_str(task_snapshot.get('status')).upper()

    if running_files > 0 or task_status in PROJECT_STATUS_RUNNING_TASK_SET:
        return 'RUNNING'

    if non_draft_files > 0 and step3_success_files >= non_draft_files:
        return 'SUCCESS'

    if task_status == 'PARTIAL_FAILED' or (failed_like_files > 0 and failed_like_files < non_draft_files):
        return 'PARTIAL_FAILED'

    if task_status in PROJECT_STATUS_FAILED_TASK_SET or failed_like_files > 0:
        return 'FAILED'

    return 'PENDING'


def _resolve_project_last_error(file_snapshot, task_snapshot):
    if task_snapshot:
        task_error = _to_str(task_snapshot.get('lastError'))
        if task_error:
            return task_error

    file_error = _to_str(file_snapshot.get('latestFileError'))
    return file_error or None


def _project_status_label(status_value):
    if status_value == 'RUNNING':
        return '运行中'
    if status_value == 'SUCCESS':
        return '解析成功'
    if status_value == 'FAILED':
        return '解析失败'
    if status_value == 'PARTIAL_FAILED':
        return '部分失败'
    return '待处理'


def _project_step_label(step_no):
    if step_no == 1:
        return '步骤1'
    if step_no == 2:
        return '步骤2'
    if step_no == 3:
        return '步骤3'
    return '待处理'


def _build_project_status_rows():
    project_rows = _fetch_project_name_rows()
    file_rows = _fetch_project_file_snapshot_rows()
    task_rows = _fetch_latest_project_task_rows()

    project_name_map = {}
    for row in project_rows:
        project_id = _to_str(row.get('projectId'))
        if not project_id:
            continue
        project_name_map[project_id] = _to_str(row.get('projectName'))

    file_snapshot_map = {}
    for row in file_rows:
        normalized = _normalize_file_snapshot(row)
        project_id = normalized.get('projectId')
        if not project_id:
            continue
        file_snapshot_map[project_id] = normalized

    task_snapshot_map = {}
    for row in task_rows:
        normalized = _normalize_task_snapshot(row)
        project_id = normalized.get('projectId')
        if not project_id:
            continue
        task_snapshot_map[project_id] = normalized

    project_ids = set(project_name_map.keys()) | set(file_snapshot_map.keys()) | set(task_snapshot_map.keys())
    rows = []
    for project_id in sorted(project_ids):
        file_snapshot = file_snapshot_map.get(project_id, _empty_file_snapshot(project_id))
        task_snapshot = task_snapshot_map.get(project_id)
        status_value = _resolve_project_status(file_snapshot, task_snapshot)
        current_step_no = _resolve_project_current_step(file_snapshot, task_snapshot)
        latest_task_status = _to_str(task_snapshot.get('status')) if task_snapshot else None
        latest_task_step_no = _to_int(task_snapshot.get('stepNo')) if task_snapshot else None
        latest_task_id = _to_str(task_snapshot.get('taskId')) if task_snapshot else None
        latest_task_creator = _to_str(task_snapshot.get('creator')) if task_snapshot else None
        latest_task_created_at = task_snapshot.get('createdAt') if task_snapshot else None
        latest_task_updated_at = task_snapshot.get('updatedAt') if task_snapshot else None

        rows.append({
            'projectId': project_id,
            'projectName': project_name_map.get(project_id, ''),
            'status': status_value,
            'statusLabel': _project_status_label(status_value),
            'currentStepNo': current_step_no,
            'currentStepName': _project_step_label(current_step_no),
            'latestTaskId': latest_task_id,
            'latestTaskCreator': latest_task_creator,
            'latestTaskCreatedAt': latest_task_created_at,
            'latestTaskStatus': latest_task_status,
            'latestTaskStepNo': latest_task_step_no,
            'latestTaskUpdatedAt': latest_task_updated_at,
            'lastError': _resolve_project_last_error(file_snapshot, task_snapshot),
            'fileStats': {
                'totalFiles': _to_int(file_snapshot.get('totalFiles')),
                'nonDraftFiles': _to_int(file_snapshot.get('nonDraftFiles')),
                'draftFiles': _to_int(file_snapshot.get('draftFiles')),
                'pendingFiles': _to_int(file_snapshot.get('pendingFiles')),
                'runningFiles': _to_int(file_snapshot.get('runningFiles')),
                'successFiles': _to_int(file_snapshot.get('successFiles')),
                'failedFiles': _to_int(file_snapshot.get('failedFiles')),
                'partialFailedFiles': _to_int(file_snapshot.get('partialFailedFiles')),
                'step3SuccessFiles': _to_int(file_snapshot.get('step3SuccessFiles')),
            }
        })

    return rows


def _build_project_summary_payload(project_status_rows):
    payload = {
        'totalProjects': len(project_status_rows),
        'runningProjects': 0,
        'successProjects': 0,
        'failedProjects': 0,
        'pendingProjects': 0,
    }

    for row in project_status_rows:
        status_value = row.get('status')
        if status_value == 'RUNNING':
            payload['runningProjects'] += 1
        elif status_value == 'SUCCESS':
            payload['successProjects'] += 1
        elif status_value in {'FAILED', 'PARTIAL_FAILED'}:
            payload['failedProjects'] += 1
        else:
            payload['pendingProjects'] += 1

    return payload


def _normalize_ranking_category(raw_value):
    normalized = _to_str(raw_value).lower()
    if not normalized:
        return PROJECT_STATUS_CATEGORY_COMPLETED
    if normalized in {'success', PROJECT_STATUS_CATEGORY_COMPLETED}:
        return PROJECT_STATUS_CATEGORY_COMPLETED
    if normalized in {'pending', PROJECT_STATUS_CATEGORY_REMAINING}:
        return PROJECT_STATUS_CATEGORY_REMAINING
    if normalized in {'failed', 'partial_failed', PROJECT_STATUS_CATEGORY_ABNORMAL}:
        return PROJECT_STATUS_CATEGORY_ABNORMAL
    if normalized in PROJECT_STATUS_CATEGORY_SET:
        return normalized
    return None


def _status_match_ranking_category(category, status_value):
    if category == PROJECT_STATUS_CATEGORY_COMPLETED:
        return status_value == 'SUCCESS'
    if category == PROJECT_STATUS_CATEGORY_RUNNING:
        return status_value == 'RUNNING'
    if category == PROJECT_STATUS_CATEGORY_REMAINING:
        return status_value == 'PENDING'
    if category == PROJECT_STATUS_CATEGORY_ABNORMAL:
        return status_value in {'FAILED', 'PARTIAL_FAILED'}
    return False


def _filter_project_status_rows_by_time(rows, start_time, end_time):
    if start_time is None and end_time is None:
        return rows

    out = []
    for row in rows:
        updated_at = _parse_datetime_value(row.get('latestTaskUpdatedAt'))
        if updated_at is None:
            continue
        if start_time and updated_at < start_time:
            continue
        if end_time and updated_at > end_time:
            continue
        out.append(row)
    return out


def _split_multi_query_items(raw_items):
    values = []
    for raw_item in raw_items:
        text = _to_str(raw_item)
        if not text:
            continue
        parts = [segment.strip() for segment in text.replace('，', ',').split(',')]
        values.extend([segment for segment in parts if segment])

    deduplicated = []
    seen = set()
    for item in values:
        normalized = item.lower()
        if normalized in seen:
            continue
        seen.add(normalized)
        deduplicated.append(item)
    return deduplicated


def _get_multi_query_values(request, key):
    raw_values = []
    raw_values.extend(request.query_params.getlist(key))
    raw_values.extend(request.query_params.getlist(f'{key}[]'))
    indexed_prefix = f'{key}['
    for param_key in request.query_params.keys():
        if param_key.startswith(indexed_prefix):
            raw_values.extend(request.query_params.getlist(param_key))

    if not raw_values:
        single_value = request.query_params.get(key)
        if single_value is not None:
            raw_values.append(single_value)

    return _split_multi_query_items(raw_values)


def _parse_optional_percent(raw_value, field_name):
    text = _to_str(raw_value)
    if not text:
        return None, ''

    try:
        value = int(text)
    except (TypeError, ValueError):
        return None, f'{field_name} 参数无效，需为 0-100 的整数'

    if value < 0 or value > 100:
        return None, f'{field_name} 参数无效，需为 0-100 的整数'

    return value, ''


def _normalize_data_schedule_task_status(raw_status):
    status_value = _to_str(raw_status).lower()
    if not status_value:
        return None
    if status_value in DATA_SCHEDULE_TASK_STATUS_SET:
        return status_value
    if status_value in {'complete', 'completed'}:
        return DATA_SCHEDULE_TASK_STATUS_SUCCESS
    if status_value in {'remain', 'remaining'}:
        return DATA_SCHEDULE_TASK_STATUS_PENDING
    if status_value in {'abnormal', 'partial_failed'}:
        return DATA_SCHEDULE_TASK_STATUS_FAILED
    return None


def _normalize_data_schedule_task_action(raw_action):
    raw = _to_str(raw_action)
    if not raw:
        return ''

    normalized = raw.replace('-', '').replace('_', '').strip().lower()
    action_map = {
        'batchdelete': DATA_SCHEDULE_TASK_ACTION_DELETE,
        'continue': DATA_SCHEDULE_TASK_ACTION_CONTINUE,
        'delete': DATA_SCHEDULE_TASK_ACTION_DELETE,
        'pause': DATA_SCHEDULE_TASK_ACTION_PAUSE,
        'refresh': DATA_SCHEDULE_TASK_ACTION_REFRESH,
        'reexecute': DATA_SCHEDULE_TASK_ACTION_RE_EXECUTE,
        'retry': DATA_SCHEDULE_TASK_ACTION_RE_EXECUTE,
        'resume': DATA_SCHEDULE_TASK_ACTION_CONTINUE,
        'stop': DATA_SCHEDULE_TASK_ACTION_STOP,
    }
    return action_map.get(normalized, '')


def _normalize_data_schedule_task_ids(raw_task_ids):
    if raw_task_ids is None:
        return []

    if isinstance(raw_task_ids, str):
        items = [item for item in raw_task_ids.split(',')]
    elif isinstance(raw_task_ids, (list, tuple, set)):
        items = list(raw_task_ids)
    else:
        items = [raw_task_ids]

    normalized_ids = []
    seen = set()
    for item in items:
        if isinstance(item, (list, tuple, set)):
            nested = item
        else:
            nested = [item]

        for value in nested:
            task_id = _to_str(value)
            if not task_id or task_id in seen:
                continue
            seen.add(task_id)
            normalized_ids.append(task_id)
    return normalized_ids


def _fetch_data_schedule_action_task_rows(task_ids):
    normalized_task_ids = _normalize_data_schedule_task_ids(task_ids)
    if not normalized_task_ids:
        return []

    where_clause, params = _build_data_schedule_where_with_values('resource_id', normalized_task_ids)
    if where_clause is None:
        return []

    sql = (
        "SELECT resource_id AS taskId, "
        "project_id AS projectId, "
        "file_id AS fileId, "
        "status AS taskQueueStatus "
        "FROM c_r_cm_task_item_queue "
        f"WHERE {where_clause} "
        "AND file_id LIKE 'PROJECT:%%'"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        return _dictfetchall(cursor)


def _reset_data_schedule_file_prepare_failed_to_pending(project_ids):
    normalized_project_ids = [_to_str(item) for item in (project_ids or []) if _to_str(item)]
    if not normalized_project_ids:
        return 0

    where_clause, params = _build_data_schedule_where_with_values('project_id', normalized_project_ids)
    if where_clause is None:
        return 0

    sql = (
        "UPDATE c_r_cm_file_prepare "
        "SET status='PENDING', "
        "last_error=NULL, "
        "updated_at=CURRENT_TIMESTAMP "
        f"WHERE {where_clause} "
        "AND status IN ('FAILED','PARTIAL_FAILED')"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        return cursor.rowcount or 0


def _update_data_schedule_queue_items(task_ids, next_status, allow_statuses=None, reset_error=False):
    normalized_task_ids = [_to_str(item) for item in (task_ids or []) if _to_str(item)]
    if not normalized_task_ids:
        return 0

    where_clause, params = _build_data_schedule_where_with_values('resource_id', normalized_task_ids)
    if where_clause is None:
        return 0

    set_parts = [
        f"status='{next_status}'",
        "locked_by=NULL",
        "locked_at=NULL",
        "updated_at=CURRENT_TIMESTAMP",
    ]
    if _table_has_column('c_r_cm_task_item_queue', 'attempt_count') and next_status == 'PENDING':
        set_parts.append("attempt_count=0")
    if _table_has_column('c_r_cm_task_item_queue', 'available_at') and next_status == 'PENDING':
        set_parts.append("available_at=CURRENT_TIMESTAMP")
    if _table_has_column('c_r_cm_task_item_queue', 'last_error') and reset_error:
        set_parts.append("last_error=NULL")

    sql = (
        "UPDATE c_r_cm_task_item_queue "
        f"SET {', '.join(set_parts)} "
        f"WHERE {where_clause} "
        "AND file_id LIKE 'PROJECT:%%'"
    )
    if allow_statuses:
        allow_clause, allow_params = _build_data_schedule_where_with_values('status', allow_statuses)
        if allow_clause:
            sql += f" AND {allow_clause}"
            params.extend(allow_params)

    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        return cursor.rowcount or 0


def _delete_data_schedule_queue_items(task_ids):
    normalized_task_ids = [_to_str(item) for item in (task_ids or []) if _to_str(item)]
    if not normalized_task_ids:
        return 0

    where_clause, params = _build_data_schedule_where_with_values('resource_id', normalized_task_ids)
    if where_clause is None:
        return 0

    sql = (
        "DELETE FROM c_r_cm_task_item_queue "
        f"WHERE {where_clause} "
        "AND file_id LIKE 'PROJECT:%%'"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        return cursor.rowcount or 0


def _delete_data_schedule_queue_items_by_project_ids(project_ids):
    normalized_project_ids = [_to_str(item) for item in (project_ids or []) if _to_str(item)]
    if not normalized_project_ids:
        return 0

    where_clause, params = _build_data_schedule_where_with_values('project_id', normalized_project_ids)
    if where_clause is None:
        return 0

    sql = (
        "DELETE FROM c_r_cm_task_item_queue "
        f"WHERE {where_clause} "
        "AND file_id LIKE 'PROJECT:%%'"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, params)
        return cursor.rowcount or 0


def _build_data_schedule_task_refresh_payload():
    rows = _build_data_schedule_task_rows()
    stats = {
        DATA_SCHEDULE_TASK_STATUS_FAILED: 0,
        DATA_SCHEDULE_TASK_STATUS_PAUSED: 0,
        DATA_SCHEDULE_TASK_STATUS_PENDING: 0,
        DATA_SCHEDULE_TASK_STATUS_RUNNING: 0,
        DATA_SCHEDULE_TASK_STATUS_STOPPED: 0,
        DATA_SCHEDULE_TASK_STATUS_SUCCESS: 0,
    }
    for row in rows:
        status_value = _to_str(row.get('taskStatus')).lower()
        if status_value in stats:
            stats[status_value] += 1

    return {
        'refreshedAt': datetime.now().strftime(DATA_SCHEDULE_LOG_TIME_FORMAT),
        'statusStats': stats,
        'total': len(rows),
    }


def _operate_data_schedule_tasks(action, task_rows):
    action_value = _normalize_data_schedule_task_action(action)
    if action_value not in DATA_SCHEDULE_TASK_ACTION_SET:
        raise ValueError('action 参数无效')

    task_ids = [_to_str(item.get('taskId')) for item in (task_rows or []) if _to_str(item.get('taskId'))]
    project_ids = [_to_str(item.get('projectId')) for item in (task_rows or []) if _to_str(item.get('projectId'))]
    deduped_project_ids = []
    seen_projects = set()
    for project_id in project_ids:
        if project_id in seen_projects:
            continue
        seen_projects.add(project_id)
        deduped_project_ids.append(project_id)

    if action_value == DATA_SCHEDULE_TASK_ACTION_REFRESH:
        return {
            'deleted': 0,
            'fileReset': 0,
            'updated': 0,
        }

    if not task_ids:
        return {
            'deleted': 0,
            'fileReset': 0,
            'updated': 0,
        }

    with transaction.atomic():
        if action_value == DATA_SCHEDULE_TASK_ACTION_RE_EXECUTE:
            updated = _update_data_schedule_queue_items(task_ids, 'PENDING', allow_statuses=None, reset_error=True)
            file_reset = _reset_data_schedule_file_prepare_failed_to_pending(deduped_project_ids)
            return {
                'deleted': 0,
                'fileReset': file_reset,
                'updated': updated,
            }

        if action_value == DATA_SCHEDULE_TASK_ACTION_CONTINUE:
            updated = _update_data_schedule_queue_items(
                task_ids,
                'PENDING',
                allow_statuses=['PAUSED', 'STOPPED', 'CANCELLED', 'CANCEL', 'FAILED', 'PARTIAL_FAILED'],
                reset_error=True
            )
            file_reset = _reset_data_schedule_file_prepare_failed_to_pending(deduped_project_ids)
            return {
                'deleted': 0,
                'fileReset': file_reset,
                'updated': updated,
            }

        if action_value == DATA_SCHEDULE_TASK_ACTION_PAUSE:
            updated = _update_data_schedule_queue_items(
                task_ids,
                'PAUSED',
                allow_statuses=['PENDING', 'CLAIMED', 'RUNNING'],
                reset_error=False
            )
            return {
                'deleted': 0,
                'fileReset': 0,
                'updated': updated,
            }

        if action_value == DATA_SCHEDULE_TASK_ACTION_STOP:
            updated = _update_data_schedule_queue_items(
                task_ids,
                'CANCEL',
                allow_statuses=['PENDING', 'CLAIMED', 'RUNNING', 'PAUSED'],
                reset_error=False
            )
            return {
                'deleted': 0,
                'fileReset': 0,
                'updated': updated,
            }

        if action_value == DATA_SCHEDULE_TASK_ACTION_DELETE:
            deleted = 0
            if deduped_project_ids:
                deleted = _delete_data_schedule_queue_items_by_project_ids(deduped_project_ids)

            # Fallback for edge cases where queue rows are missing project_id values.
            if deleted <= 0:
                deleted = _delete_data_schedule_queue_items(task_ids)
            return {
                'deleted': deleted,
                'fileReset': 0,
                'updated': 0,
            }

    return {
        'deleted': 0,
        'fileReset': 0,
        'updated': 0,
    }


def _parse_data_schedule_task_filters(request):
    current = _parse_positive_int(request.query_params.get('current'), default=1, max_value=100000)
    size = _parse_positive_int(
        request.query_params.get('size'),
        default=DATA_SCHEDULE_TASK_DEFAULT_SIZE,
        max_value=DATA_SCHEDULE_TASK_MAX_SIZE
    )
    project_name = _to_str(request.query_params.get('projectName'))

    task_status_values = []
    raw_task_status_values = _get_multi_query_values(request, 'taskStatus')
    for raw_status in raw_task_status_values:
        normalized_status = _normalize_data_schedule_task_status(raw_status)
        if normalized_status is None:
            return None, 'taskStatus 参数无效，仅支持 pending/running/success/failed/paused/stopped'
        task_status_values.append(normalized_status)

    creator_values = [item for item in _get_multi_query_values(request, 'creator') if item]

    match_range_min, min_err_msg = _parse_optional_percent(request.query_params.get('matchRangeMin'), 'matchRangeMin')
    if min_err_msg:
        return None, min_err_msg
    match_range_max, max_err_msg = _parse_optional_percent(request.query_params.get('matchRangeMax'), 'matchRangeMax')
    if max_err_msg:
        return None, max_err_msg
    if match_range_min is not None and match_range_max is not None and match_range_min > match_range_max:
        return None, 'matchRangeMin 不能大于 matchRangeMax'

    raw_create_start_time = request.query_params.get('createStartTime') or request.query_params.get('startTime')
    raw_create_end_time = request.query_params.get('createEndTime') or request.query_params.get('endTime')
    create_start_time = _parse_datetime_value(raw_create_start_time)
    create_end_time = _parse_datetime_value(raw_create_end_time)

    if raw_create_start_time and create_start_time is None:
        return None, 'createStartTime 格式无效'
    if raw_create_end_time and create_end_time is None:
        return None, 'createEndTime 格式无效'
    if raw_create_end_time and create_end_time and len(_to_str(raw_create_end_time)) == 10:
        create_end_time = create_end_time + timedelta(days=1) - timedelta(seconds=1)
    if create_start_time and create_end_time and create_start_time > create_end_time:
        return None, 'createStartTime 不能晚于 createEndTime'

    return {
        'creatorValues': creator_values,
        'createEndTime': create_end_time,
        'createStartTime': create_start_time,
        'current': current,
        'matchRangeMax': match_range_max,
        'matchRangeMin': match_range_min,
        'projectName': project_name,
        'size': size,
        'taskStatusValues': task_status_values,
    }, ''


def _fetch_project_schedule_meta_rows():
    valid_project_where_clause = _build_valid_project_where_clause()
    creator_column = _pick_existing_column('c_r_cm_project', DATA_SCHEDULE_PROJECT_CREATOR_COLUMN_CANDIDATES)
    create_time_column = _pick_existing_column('c_r_cm_project', DATA_SCHEDULE_PROJECT_CREATE_TIME_COLUMN_CANDIDATES)
    select_parts = [
        "project_id AS projectId",
        f"{creator_column} AS creator" if creator_column else "NULL AS creator",
        f"{create_time_column} AS createTime" if create_time_column else "NULL AS createTime",
    ]
    sql = (
        "SELECT "
        + ", ".join(select_parts)
        + " FROM c_r_cm_project "
        f"WHERE {valid_project_where_clause}"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql)
        return _dictfetchall(cursor)


def _map_project_row_to_schedule_task_status(project_row):
    latest_task_status = _to_str(project_row.get('latestTaskStatus')).upper()
    if latest_task_status == 'PAUSED':
        return DATA_SCHEDULE_TASK_STATUS_PAUSED
    if latest_task_status in {'CANCELLED', 'CANCEL', 'STOPPED'}:
        return DATA_SCHEDULE_TASK_STATUS_STOPPED

    normalized_status = _to_str(project_row.get('status')).upper()
    if normalized_status == 'RUNNING':
        return DATA_SCHEDULE_TASK_STATUS_RUNNING
    if normalized_status == 'SUCCESS':
        return DATA_SCHEDULE_TASK_STATUS_SUCCESS
    if normalized_status in {'FAILED', 'PARTIAL_FAILED'}:
        return DATA_SCHEDULE_TASK_STATUS_FAILED
    return DATA_SCHEDULE_TASK_STATUS_PENDING


def _build_data_schedule_progress(file_stats):
    non_draft_files = _to_int((file_stats or {}).get('nonDraftFiles'))
    if non_draft_files <= 0:
        return 0
    step3_success_files = _to_int((file_stats or {}).get('step3SuccessFiles'))
    progress = round(step3_success_files * 100 / non_draft_files)
    return max(0, min(100, int(progress)))


def _build_data_schedule_task_rows():
    project_rows = _build_project_status_rows()
    project_meta_rows = _fetch_project_schedule_meta_rows()
    project_meta_map = {}
    for row in project_meta_rows:
        project_id = _to_str(row.get('projectId'))
        if not project_id:
            continue
        project_meta_map[project_id] = {
            'creator': _to_str(row.get('creator')),
            'createTime': _format_datetime_text(row.get('createTime'))
        }

    records = []
    for project_row in project_rows:
        file_stats = project_row.get('fileStats') or {}
        if _to_int(file_stats.get('totalFiles')) <= 0:
            continue
        if not _to_str(project_row.get('latestTaskId')):
            continue

        project_id = _to_str(project_row.get('projectId'))
        project_name = _to_str(project_row.get('projectName'))
        task_id = _to_str(project_row.get('latestTaskId')) or f'PROJECT:{project_id}'
        task_status = _map_project_row_to_schedule_task_status(project_row)
        progress = _build_data_schedule_progress(file_stats)

        project_meta = project_meta_map.get(project_id, {})
        creator = (
            _to_str(project_row.get('latestTaskCreator'))
            or _to_str(project_meta.get('creator'))
            or 'system'
        )
        create_time = (
            _to_str(project_row.get('latestTaskCreatedAt'))
            or _to_str(project_meta.get('createTime'))
            or _to_str(project_row.get('latestTaskUpdatedAt'))
        )
        complete_time = ''
        if task_status in {
            DATA_SCHEDULE_TASK_STATUS_SUCCESS,
            DATA_SCHEDULE_TASK_STATUS_FAILED,
            DATA_SCHEDULE_TASK_STATUS_PAUSED,
            DATA_SCHEDULE_TASK_STATUS_STOPPED,
        }:
            complete_time = _to_str(project_row.get('latestTaskUpdatedAt'))

        update_time = _to_str(project_row.get('latestTaskUpdatedAt')) or create_time

        records.append({
            'completeTime': complete_time,
            'createBy': creator,
            'createTime': create_time,
            'creator': creator,
            'projectId': project_id,
            'projectName': project_name or project_id,
            'progress': progress,
            'status': None,
            'taskId': task_id,
            'taskStatus': task_status,
            'updateBy': creator,
            'updateTime': update_time,
        })

    records.sort(
        key=lambda item: (
            _parse_datetime_value(item.get('createTime')) or datetime.min,
            _to_str(item.get('taskId')),
            _to_str(item.get('projectId'))
        ),
        reverse=True
    )

    for index, item in enumerate(records, start=1):
        item['id'] = index

    return records


def _filter_data_schedule_task_rows(rows, filters):
    project_name_keyword = _to_str((filters or {}).get('projectName')).lower()
    task_status_values = set((filters or {}).get('taskStatusValues') or [])
    creator_values = [item.lower() for item in ((filters or {}).get('creatorValues') or []) if item]
    creator_set = set(creator_values)
    match_range_min = (filters or {}).get('matchRangeMin')
    match_range_max = (filters or {}).get('matchRangeMax')
    create_start_time = (filters or {}).get('createStartTime')
    create_end_time = (filters or {}).get('createEndTime')

    filtered = []
    for row in rows:
        project_name = _to_str(row.get('projectName'))
        if project_name_keyword and project_name_keyword not in project_name.lower():
            continue

        task_status = _to_str(row.get('taskStatus')).lower()
        if task_status_values and task_status not in task_status_values:
            continue

        creator = _to_str(row.get('creator'))
        if creator_set and creator.lower() not in creator_set:
            continue

        progress = _to_int(row.get('progress'))
        if match_range_min is not None and progress < match_range_min:
            continue
        if match_range_max is not None and progress > match_range_max:
            continue

        create_time = _parse_datetime_value(row.get('createTime'))
        if create_start_time and (create_time is None or create_time < create_start_time):
            continue
        if create_end_time and (create_time is None or create_time > create_end_time):
            continue

        filtered.append(row)

    return filtered


def _build_data_schedule_creator_options(keyword=''):
    query_set = User.objects.filter(is_active=True)
    normalized_keyword = _to_str(keyword)
    if normalized_keyword:
        query_set = query_set.filter(username__icontains=normalized_keyword)

    users = query_set.order_by('username').values('username')
    options = []
    for user in users:
        username = _to_str(user.get('username'))
        if not username:
            continue
        options.append({'label': username, 'value': username})

    return options


def _to_excel_cell_value(value):
    if value is None:
        return ''
    if isinstance(value, (str, int, float, bool, datetime)):
        return value
    if isinstance(value, (list, tuple, set, dict)):
        try:
            return json.dumps(value, ensure_ascii=False)
        except Exception:
            return str(value)
    return str(value)


def _get_request_query_snapshot(request):
    return {key: request.query_params.get(key) for key in request.query_params.keys()}


def _append_data_schedule_export_log(level, event, *, task_id='', scope='', request_params=None, extra=None, error=None):
    record = {
        'event': event,
        'level': level,
        'scope': scope,
        'taskId': task_id,
        'time': datetime.now().strftime(DATA_SCHEDULE_LOG_TIME_FORMAT)
    }
    if request_params:
        record['requestParams'] = request_params
    if extra is not None:
        record['extra'] = extra
    if error is not None:
        record['error'] = f'{type(error).__name__}: {error}'
        record['traceback'] = traceback.format_exc()

    try:
        DATA_SCHEDULE_EXPORT_LOG_DIR.mkdir(parents=True, exist_ok=True)
        with DATA_SCHEDULE_EXPORT_LOG_FILE.open('a', encoding='utf-8') as log_file:
            log_file.write(f'{json.dumps(record, ensure_ascii=False)}\n')
    except Exception:
        logger.exception('append data schedule export log failed')


def _read_data_schedule_export_log_records(limit, keyword=''):
    if not DATA_SCHEDULE_EXPORT_LOG_FILE.exists():
        return []

    normalized_keyword = str(keyword or '').strip().lower()
    records = []
    with DATA_SCHEDULE_EXPORT_LOG_FILE.open('r', encoding='utf-8') as log_file:
        for line in log_file:
            raw_line = line.strip()
            if not raw_line:
                continue
            if normalized_keyword and normalized_keyword not in raw_line.lower():
                continue
            try:
                records.append(json.loads(raw_line))
            except json.JSONDecodeError:
                records.append({'raw': raw_line})

    if limit <= 0:
        return records
    return records[-limit:]


def _read_api_failure_log_records(limit, keyword=''):
    if not API_FAILURE_LOG_FILE.exists():
        return []

    normalized_keyword = str(keyword or '').strip().lower()
    records = []
    with API_FAILURE_LOG_FILE.open('r', encoding='utf-8') as log_file:
        for line in log_file:
            raw_line = line.strip()
            if not raw_line:
                continue
            if normalized_keyword and normalized_keyword not in raw_line.lower():
                continue
            try:
                records.append(json.loads(raw_line))
            except json.JSONDecodeError:
                records.append({'raw': raw_line})

    if limit <= 0:
        return records
    return records[-limit:]


def _raise_data_schedule_real_source_error():
    raise RuntimeError(DATA_SCHEDULE_REAL_SOURCE_ERROR)


def _build_data_schedule_mock_fields():
    fields = [
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'bid_no',
            'fieldName': '招标编号',
            'fieldValueDisplay': 'CG1201002025044003(1)',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'project_year',
            'fieldName': '项目年度',
            'fieldValueDisplay': '2025',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'bidder_name',
            'fieldName': '招标人',
            'fieldValueDisplay': '天津市城市道路桥梁管理事务中心',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'legal_person',
            'fieldName': '法定代表人',
            'fieldValueDisplay': '李民',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'project_name',
            'fieldName': '项目名称',
            'fieldValueDisplay': '道桥中心2025年度道路挖掘损害修复项目',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'approver_org',
            'fieldName': '项目批准机关',
            'fieldValueDisplay': '天津市城市管理委员会',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'approval_doc_no',
            'fieldName': '批准文号',
            'fieldValueDisplay': '津城管桥批〔2025〕33号',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'pricing_ceiling',
            'fieldName': '招标代理机构最高投标限价',
            'fieldValueDisplay': '39400000',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'authorized_letter',
            'fieldName': '授权委托书',
            'fieldValueDisplay': '--',
            'status': DATA_SCHEDULE_SCOPE_MISSING
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'insurance_proof',
            'fieldName': '参保缴费证明',
            'fieldValueDisplay': '--',
            'status': DATA_SCHEDULE_SCOPE_MISSING
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'project_bill_of_quantities',
            'fieldName': '工程量清单',
            'fieldValueDisplay': '--',
            'status': DATA_SCHEDULE_SCOPE_MISSING
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'bidder_representative',
            'fieldName': '唱标人代表',
            'fieldValueDisplay': '--',
            'status': DATA_SCHEDULE_SCOPE_MISSING
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'bid_opening_host',
            'fieldName': '唱标人',
            'fieldValueDisplay': '王文哲',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'recorder',
            'fieldName': '记录人',
            'fieldValueDisplay': '洪聪',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': True,
            'canEdit': True,
            'fieldKey': 'supervisor_info',
            'fieldName': '监督人信息',
            'fieldValueDisplay': '查看',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': True,
            'canEdit': True,
            'fieldKey': 'expert_info',
            'fieldName': '评标专家信息',
            'fieldValueDisplay': '查看',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'committee_member_list',
            'fieldName': '评标委员会组成人员名单',
            'fieldValueDisplay': '--',
            'status': DATA_SCHEDULE_SCOPE_MISSING
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'party_a_representative_name',
            'fieldName': '甲方评委代表人姓名',
            'fieldValueDisplay': '徐鹏',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'party_a_representative_id_no',
            'fieldName': '甲方评委代表人身份证号',
            'fieldValueDisplay': '120101197607132554',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'party_a_representative_mobile',
            'fieldName': '甲方评委代表人手机号',
            'fieldValueDisplay': '18622884518',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'construction_unit_name',
            'fieldName': '建设单位法人名称',
            'fieldValueDisplay': '天津市城市道路桥梁管理事务中心',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'construction_unit_legal_person',
            'fieldName': '建设单位法人代表姓名',
            'fieldValueDisplay': '李民',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'construction_unit_cert_no',
            'fieldName': '建设单位证件号',
            'fieldValueDisplay': '120108197710274037',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'qualification_review_form',
            'fieldName': '资格评审情况表',
            'fieldValueDisplay': '--',
            'status': DATA_SCHEDULE_SCOPE_MISSING
        },
        {
            'canDrilldown': True,
            'canEdit': True,
            'fieldKey': 'bid_ranking',
            'fieldName': '投标报价排名',
            'fieldValueDisplay': '查看',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': True,
            'canEdit': True,
            'fieldKey': 'winning_candidate_notice',
            'fieldName': '中标候选人公示',
            'fieldValueDisplay': '查看',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'winning_company',
            'fieldName': '中标人',
            'fieldValueDisplay': '天津路桥建设工程有限公司',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
        {
            'canDrilldown': False,
            'canEdit': True,
            'fieldKey': 'winning_amount',
            'fieldName': '中标金额',
            'fieldValueDisplay': '38296827',
            'status': DATA_SCHEDULE_SCOPE_COMPLETE
        },
    ]

    # Fillers to keep a data-dense detail table like the design reference.
    for index in range(1, 17):
        status_value = DATA_SCHEDULE_SCOPE_COMPLETE if index % 3 else DATA_SCHEDULE_SCOPE_MISSING
        fields.append(
            {
                'canDrilldown': False,
                'canEdit': True,
                'fieldKey': f'extended_field_{index}',
                'fieldName': f'扩展字段{index}',
                'fieldValueDisplay': f'扩展值{index}' if status_value == DATA_SCHEDULE_SCOPE_COMPLETE else '--',
                'status': status_value,
            }
        )

    return fields


DATA_SCHEDULE_DRILLDOWN_MOCK = {
    'expert_info': {
        'columns': [
            {'editable': True, 'key': 'expertName', 'title': '专家姓名'},
            {'editable': True, 'key': 'idNo', 'title': '证件号'},
            {'editable': True, 'key': 'major', 'title': '评标专业'},
            {'editable': True, 'key': 'expertType', 'title': '专家类型'},
            {'editable': True, 'key': 'workUnit', 'title': '工作单位'},
            {'editable': True, 'key': 'contact', 'title': '联系方式'},
        ],
        'records': [
            {'contact': '--', 'expertName': '周盈盈', 'expertType': '社会评委（正常）', 'id': 1, 'idNo': '--', 'major': '--', 'workUnit': '天津海泰市政绿化有限公司'},
            {'contact': '--', 'expertName': '莫鸿雁', 'expertType': '社会评委（正常）', 'id': 2, 'idNo': '--', 'major': '--', 'workUnit': '天津市蓟州区水务管理服务中心'},
            {'contact': '--', 'expertName': '高国妍', 'expertType': '社会评委（正常）', 'id': 3, 'idNo': '--', 'major': '--', 'workUnit': '天津市西青区公路建设养护中心'},
            {'contact': '--', 'expertName': '李晓明', 'expertType': '社会评委（正常）', 'id': 4, 'idNo': '--', 'major': '--', 'workUnit': '天津安纳赛能源科技有限公司'},
            {'contact': '--', 'expertName': '郭听燃', 'expertType': '招标人代表', 'id': 5, 'idNo': '120101198607303012', 'major': '--', 'workUnit': '天津市城市道路桥梁管理事务中心'},
            {'contact': '--', 'expertName': '徐鹏', 'expertType': '招标人代表', 'id': 6, 'idNo': '120101197607132554', 'major': '--', 'workUnit': '天津市城市道路桥梁管理事务中心'},
            {'contact': '--', 'expertName': '刘斌', 'expertType': '招标人代表', 'id': 7, 'idNo': '--', 'major': '--', 'workUnit': '天津市城市道路桥梁管理事务中心'},
            {'contact': '--', 'expertName': '李志壮', 'expertType': '社会评委（正常）', 'id': 8, 'idNo': '--', 'major': '--', 'workUnit': '天津市建设发展总公司'},
            {'contact': '--', 'expertName': '白树海', 'expertType': '社会评委（正常）', 'id': 9, 'idNo': '--', 'major': '--', 'workUnit': '天津市市政景观设计有限公司'},
        ],
        'title': '字段详情 - 评标专家信息'
    },
    'supervisor_info': {
        'columns': [
            {'editable': True, 'key': 'name', 'title': '监督人姓名'},
            {'editable': True, 'key': 'department', 'title': '所属单位'},
            {'editable': True, 'key': 'duty', 'title': '职务'},
            {'editable': True, 'key': 'contact', 'title': '联系方式'},
        ],
        'records': [
            {'contact': '13800001234', 'department': '天津市城市管理委员会', 'duty': '项目监督员', 'id': 1, 'name': '张晨'},
            {'contact': '13900004321', 'department': '天津市城市管理委员会', 'duty': '质量监督员', 'id': 2, 'name': '刘洋'},
        ],
        'title': '字段详情 - 监督人信息'
    },
    'bid_ranking': {
        'columns': [
            {'editable': True, 'key': 'companyName', 'title': '投标单位'},
            {'editable': True, 'key': 'bidAmount', 'title': '投标报价'},
            {'editable': True, 'key': 'rank', 'title': '排名'},
            {'editable': True, 'key': 'remark', 'title': '备注'},
        ],
        'records': [
            {'bidAmount': '38296827', 'companyName': '天津路桥建设工程有限公司', 'id': 1, 'rank': '1', 'remark': '中标候选人'},
            {'bidAmount': '38900220', 'companyName': '天津市政工程集团', 'id': 2, 'rank': '2', 'remark': '--'},
            {'bidAmount': '39218800', 'companyName': '天津城建路桥有限公司', 'id': 3, 'rank': '3', 'remark': '--'},
        ],
        'title': '字段详情 - 投标报价排名'
    },
    'winning_candidate_notice': {
        'columns': [
            {'editable': True, 'key': 'candidateName', 'title': '候选人名称'},
            {'editable': True, 'key': 'publicDate', 'title': '公示日期'},
            {'editable': True, 'key': 'score', 'title': '综合得分'},
            {'editable': True, 'key': 'remark', 'title': '备注'},
        ],
        'records': [
            {'candidateName': '天津路桥建设工程有限公司', 'id': 1, 'publicDate': '2026-02-05', 'remark': '第一中标候选人', 'score': '97.50'},
            {'candidateName': '天津市政工程集团', 'id': 2, 'publicDate': '2026-02-05', 'remark': '第二中标候选人', 'score': '95.20'},
            {'candidateName': '天津城建路桥有限公司', 'id': 3, 'publicDate': '2026-02-05', 'remark': '第三中标候选人', 'score': '93.40'},
        ],
        'title': '字段详情 - 中标候选人公示'
    }
}


def _get_data_schedule_industry_table_config(industry):
    normalized = _normalize_data_schedule_detail_industry(industry)
    if normalized == DATA_SCHEDULE_DETAIL_SW:
        return {
            'bidSubmissionFieldName': '中标候选人公示',
            'bidSubmissionFieldKey': 'winning_candidate_notice',
            'bidSubmissionTable': DATA_SCHEDULE_SW_BID_SUBMISSION_TABLE,
            'detailProjectTable': DATA_SCHEDULE_SW_PROJECT_INFO_TABLE,
            'detailSectionTable': DATA_SCHEDULE_SW_SECTION_TABLE,
            'expertFieldName': '评标专家信息',
            'expertFieldKey': 'expert_info',
            'expertTable': DATA_SCHEDULE_SW_EXPERT_TABLE,
            'openingAttendeeFieldName': '',
            'openingAttendeeFieldKey': '',
            'openingAttendeeTable': '',
            'tenderAgentTable': DATA_SCHEDULE_SW_TENDER_AGENT_TABLE,
        }
    if normalized == DATA_SCHEDULE_DETAIL_JS:
        return {
            'bidSubmissionFieldName': '',
            'bidSubmissionFieldKey': '',
            'bidSubmissionTable': '',
            'detailProjectTable': DATA_SCHEDULE_JS_PROJECT_INFO_TABLE,
            'detailSectionTable': DATA_SCHEDULE_JS_SECTION_TABLE,
            'expertFieldName': '评标专家信息',
            'expertFieldKey': 'expert_info',
            'expertTable': DATA_SCHEDULE_JS_EXPERT_TABLE,
            'openingAttendeeFieldName': '开标人员信息',
            'openingAttendeeFieldKey': 'opening_attendee_info',
            'openingAttendeeTable': DATA_SCHEDULE_JS_OPENING_ATTENDEE_TABLE,
            'tenderAgentTable': DATA_SCHEDULE_JS_BID_SUBMISSION_TABLE,
        }
    return None


def _build_data_schedule_detail_context(task_id):
    task_row = _find_data_schedule_task_row(task_id)
    if not task_row:
        raise ValueError('任务不存在')

    file_prepare_project_id = _to_str(task_row.get('projectId')) or _resolve_data_schedule_project_id_from_task_id(task_id)
    project_row = _fetch_data_schedule_project_base_row(file_prepare_project_id)
    if not project_row:
        raise ValueError('项目不存在或无有效行业信息')

    industry = _normalize_data_schedule_detail_industry(project_row.get('industry'))
    if not industry:
        raise RuntimeError('项目行业为空，无法加载详情')

    kb_project_id = _resolve_data_schedule_kb_project_id_from_file_prepare_project_id(file_prepare_project_id)
    if not kb_project_id:
        raise ValueError('项目映射不存在')

    config = _get_data_schedule_industry_table_config(industry)
    if config is None:
        raise RuntimeError(f'暂不支持行业 {industry} 的详情查询')

    return {
        'agentIds': set(),
        'config': config,
        'filePrepareProjectId': file_prepare_project_id,
        'industry': industry,
        'projectId': kb_project_id,
        'projectName': _to_str(project_row.get('projectName')),
        'sectionIds': set(),
        'taskId': _to_str(task_row.get('taskId')) or _to_str(task_id),
        'taskRow': task_row,
    }


def _build_data_schedule_where_with_values(column_name, values):
    normalized_values = [_to_str(item) for item in values if _to_str(item)]
    if not normalized_values:
        return None, []

    placeholders = ', '.join(['%s'] * len(normalized_values))
    return f"`{column_name}` IN ({placeholders})", normalized_values


def _build_data_schedule_detail_table_rows(table_name, project_id):
    if not _table_exists(table_name):
        return []

    where_parts = []
    params = []
    if _table_has_column(table_name, 'project_id'):
        where_parts.append('`project_id` = %s')
        params.append(project_id)
    return _fetch_data_schedule_table_records(table_name, where_parts, params)


def _dedupe_data_schedule_field_rows(rows):
    deduped = []
    seen = set()
    for row in rows:
        if row.get('canDrilldown'):
            signature = ('drilldown', _to_str(row.get('fieldKey')))
        else:
            signature = (
                'field',
                _to_str(row.get('fieldName')),
                _to_str(row.get('fieldValueDisplay')),
            )
        if signature in seen:
            continue
        seen.add(signature)
        deduped.append(row)
    return deduped


def _append_data_schedule_drilldown_field(rows, field_key, field_name, status_value=DATA_SCHEDULE_SCOPE_COMPLETE):
    if not field_key or not field_name:
        return
    rows.append({
        'canDrilldown': True,
        'canEdit': False,
        'fieldKey': field_key,
        'fieldName': field_name,
        'fieldValueDisplay': '查看',
        'status': status_value,
    })


def _build_data_schedule_field_rows(context):
    project_id = _to_str((context or {}).get('projectId'))
    config = (context or {}).get('config') or {}
    if not project_id or not config:
        return []

    field_rows = []
    agent_ids = set()
    section_ids = set()
    for table_name in [
        _to_str(config.get('detailProjectTable')),
        _to_str(config.get('detailSectionTable')),
    ]:
        if not table_name:
            continue
        table_rows = _build_data_schedule_detail_table_rows(table_name, project_id)
        normalized_rows, extracted_agent_ids, extracted_section_ids = _normalize_data_schedule_detail_rows(table_name, table_rows)
        field_rows.extend(normalized_rows)
        agent_ids.update(extracted_agent_ids)
        section_ids.update(extracted_section_ids)

    context['agentIds'] = agent_ids
    context['sectionIds'] = section_ids

    tender_agent_status = DATA_SCHEDULE_SCOPE_COMPLETE if agent_ids else DATA_SCHEDULE_SCOPE_MISSING
    _append_data_schedule_drilldown_field(field_rows, 'tender_agent_info', '招标代理机构', tender_agent_status)
    _append_data_schedule_drilldown_field(
        field_rows,
        _to_str(config.get('expertFieldKey')),
        _to_str(config.get('expertFieldName')),
    )
    _append_data_schedule_drilldown_field(
        field_rows,
        _to_str(config.get('bidSubmissionFieldKey')),
        _to_str(config.get('bidSubmissionFieldName')),
    )
    _append_data_schedule_drilldown_field(
        field_rows,
        _to_str(config.get('openingAttendeeFieldKey')),
        _to_str(config.get('openingAttendeeFieldName')),
    )

    return _dedupe_data_schedule_field_rows(field_rows)


def _extract_data_schedule_bid_no(fields):
    for field in fields:
        field_name = _to_str(field.get('fieldName'))
        if '招标编号' not in field_name:
            continue
        field_value = _to_str(field.get('fieldValueDisplay'))
        if not _is_blank_text_value(field_value):
            return field_value
    return ''


def _build_data_schedule_detail_bundle(task_id):
    context = _build_data_schedule_detail_context(task_id)
    fields = _build_data_schedule_field_rows(context)
    counts = _build_data_schedule_counts(fields)
    task_row = context.get('taskRow') or {}

    summary = {
        'bidNo': _extract_data_schedule_bid_no(fields),
        'counts': counts,
        'createTime': _to_str(task_row.get('createTime')),
        'creator': _to_str(task_row.get('creator')),
        'progress': _to_int(task_row.get('progress')),
        'projectName': _to_str(task_row.get('projectName')) or _to_str(context.get('projectName')),
        'status': _to_str(task_row.get('taskStatus')) or DATA_SCHEDULE_TASK_STATUS_PENDING,
        'taskId': _to_str(task_row.get('taskId')) or _to_str(task_id),
    }

    detail = {
        **summary,
        'industry': _to_str(context.get('industry')),
        'projectId': _to_str(context.get('filePrepareProjectId')) or _to_str(context.get('projectId')),
    }

    return {
        'context': context,
        'detail': detail,
        'fields': fields,
        'summary': summary,
    }


def _build_data_schedule_drilldown_from_table(table_name, title, context, prefer_agent=False):
    if not _table_exists(table_name):
        return {
            'actions': {'canCreate': False, 'canDelete': False, 'canEdit': False},
            'columns': [],
            'defaultInsertValues': {},
            'editableColumns': [],
            'records': [],
            'sourceTable': table_name,
            'title': title,
        }

    where_parts = []
    params = []
    project_id = _to_str((context or {}).get('projectId'))
    section_ids = sorted([(item) for item in ((context or {}).get('sectionIds') or set()) if _to_str(item)])
    agent_ids = sorted([(item) for item in ((context or {}).get('agentIds') or set()) if _to_str(item)])

    if prefer_agent and _table_has_column(table_name, 'agent_id'):
        in_clause, in_params = _build_data_schedule_where_with_values('agent_id', agent_ids)
        if in_clause is None:
            return {
                'actions': {'canCreate': False, 'canDelete': False, 'canEdit': False},
                'columns': [],
                'defaultInsertValues': {},
                'editableColumns': [],
                'records': [],
                'sourceTable': table_name,
                'title': title,
            }
        where_parts.append(in_clause)
        params.extend(in_params)

    if _table_has_column(table_name, 'project_id') and project_id:
        where_parts.append('`project_id` = %s')
        params.append(project_id)

    if _table_has_column(table_name, 'prj_section_id') and section_ids:
        section_clause, section_params = _build_data_schedule_where_with_values('prj_section_id', section_ids)
        if section_clause:
            where_parts.append(section_clause)
            params.extend(section_params)

    rows = _fetch_data_schedule_table_records(table_name, where_parts, params)
    column_meta = [
        item
        for item in _get_table_column_meta(table_name)
        if item.get('normalizedName') not in DATA_SCHEDULE_DETAIL_EXCLUDED_COLUMN_SET
    ]
    columns = []
    normalized_to_column_name = {}
    editable_columns = []
    for item in column_meta:
        column_name = _to_str(item.get('name'))
        if not column_name:
            continue

        normalized_name = _to_str(item.get('normalizedName')).lower()
        if normalized_name and normalized_name not in normalized_to_column_name:
            normalized_to_column_name[normalized_name] = column_name

        editable = _is_data_schedule_drilldown_column_editable(normalized_name)
        columns.append({
            'editable': editable,
            'key': column_name,
            'title': _build_data_schedule_field_name(item),
        })
        if editable:
            editable_columns.append(column_name)

    records = []
    writable_rows = 0
    for index, row in enumerate(rows, start=1):
        normalized = {
            '__sourceTable': table_name,
            'id': index,
        }
        for column in columns:
            column_key = _to_str(column.get('key'))
            normalized[column_key] = _format_data_schedule_field_value((row or {}).get(column_key))
        source_locator = _build_data_schedule_row_locator(normalized_to_column_name, row)
        normalized['__sourceLocator'] = source_locator
        if source_locator:
            writable_rows += 1
        records.append(normalized)

    default_insert_values = {}
    can_create = bool(editable_columns)
    if _table_has_column(table_name, 'project_id'):
        if project_id:
            default_insert_values['project_id'] = project_id
        else:
            can_create = False
    if _table_has_column(table_name, 'prj_section_id'):
        if len(section_ids) == 1:
            default_insert_values['prj_section_id'] = section_ids[0]
        elif section_ids:
            can_create = False
    if prefer_agent and _table_has_column(table_name, 'agent_id'):
        if len(agent_ids) == 1:
            default_insert_values['agent_id'] = agent_ids[0]
        elif agent_ids:
            can_create = False

    return {
        'actions': {
            'canCreate': can_create,
            'canDelete': writable_rows > 0,
            'canEdit': bool(editable_columns) and writable_rows > 0
        },
        'columns': columns,
        'defaultInsertValues': default_insert_values,
        'editableColumns': editable_columns,
        'records': records,
        'sourceTable': table_name,
        'title': title,
    }


def _get_data_schedule_drilldown_payload(task_id, field_key):
    bundle = _build_data_schedule_detail_bundle(task_id)
    field_item = _get_data_schedule_field_item(task_id, field_key, bundle=bundle)
    if not field_item:
        return None, None, '字段不存在', status.HTTP_404_NOT_FOUND, ERROR_CODE_NOT_FOUND
    if not field_item.get('canDrilldown'):
        return None, None, '当前字段不支持下钻', status.HTTP_400_BAD_REQUEST, ERROR_CODE_INVALID_PARAMS

    context = bundle.get('context') or {}
    config = context.get('config') or {}
    normalized_field_key = _to_str(field_item.get('fieldKey'))

    table_name = ''
    title = f"字段详情 - {field_item.get('fieldName') or normalized_field_key}"
    prefer_agent = False
    if normalized_field_key == 'tender_agent_info':
        table_name = _to_str(config.get('tenderAgentTable'))
        prefer_agent = True
    elif normalized_field_key == _to_str(config.get('expertFieldKey')):
        table_name = _to_str(config.get('expertTable'))
    elif normalized_field_key == _to_str(config.get('bidSubmissionFieldKey')):
        table_name = _to_str(config.get('bidSubmissionTable'))
    elif normalized_field_key == _to_str(config.get('openingAttendeeFieldKey')):
        table_name = _to_str(config.get('openingAttendeeTable'))

    if not table_name:
        return None, None, '下钻数据不存在', status.HTTP_404_NOT_FOUND, ERROR_CODE_NOT_FOUND

    payload = _build_data_schedule_drilldown_from_table(table_name, title, context, prefer_agent=prefer_agent)
    return field_item, payload, '', status.HTTP_200_OK, SUCCESS_CODE


def _get_data_schedule_summary(task_id):
    return _build_data_schedule_detail_bundle(task_id).get('summary') or {}


def _get_data_schedule_detail(task_id):
    return _build_data_schedule_detail_bundle(task_id).get('detail') or {}


def _get_data_schedule_fields(task_id):
    return _build_data_schedule_detail_bundle(task_id).get('fields') or []


def _normalize_data_schedule_source_locator(source_table, source_locator):
    normalized_locator = []
    for locator_item in source_locator:
        if not isinstance(locator_item, dict):
            continue
        locator_column = _to_str(locator_item.get('column'))
        locator_value = locator_item.get('value')
        if not locator_column or _is_blank_text_value(locator_value):
            continue
        if not _is_safe_sql_identifier(locator_column):
            raise ValueError('字段数据源不合法')
        if not _table_has_column(source_table, locator_column):
            raise ValueError('字段数据源不存在')
        normalized_locator.append((locator_column, locator_value))
    if not normalized_locator:
        raise ValueError('字段缺少可更新的数据源信息')
    return normalized_locator


def _ensure_single_data_schedule_locator_target(source_table, where_sql, locator_params):
    with connection.cursor() as cursor:
        cursor.execute(
            f"SELECT COUNT(1) FROM `{source_table}` WHERE {where_sql}",
            locator_params,
        )
        count_row = cursor.fetchone() or [0]

    match_count = _to_int(count_row[0], default=0)
    if match_count <= 0:
        raise ValueError('更新目标不存在')
    if match_count > 1:
        raise ValueError('更新目标不唯一，暂不支持编辑')


def _create_data_schedule_drilldown_row_to_db(payload, row_data):
    source_table = _to_str((payload or {}).get('sourceTable'))
    editable_columns = [item for item in ((payload or {}).get('editableColumns') or []) if _to_str(item)]
    default_insert_values = (payload or {}).get('defaultInsertValues') or {}

    if not source_table or not _is_safe_sql_identifier(source_table):
        raise ValueError('字段数据源不合法')
    if not _table_exists(source_table):
        raise ValueError('字段数据源不存在')

    insert_values = {}
    for key, value in (default_insert_values or {}).items():
        column_name = _to_str(key)
        if not column_name:
            continue
        if not _is_safe_sql_identifier(column_name):
            raise ValueError('字段数据源不合法')
        if not _table_has_column(source_table, column_name):
            raise ValueError('字段数据源不存在')
        insert_values[column_name] = value

    for column_name in editable_columns:
        normalized_column = _to_str(column_name)
        if not normalized_column:
            continue
        if not _is_safe_sql_identifier(normalized_column):
            raise ValueError('字段数据源不合法')
        if not _table_has_column(source_table, normalized_column):
            raise ValueError('字段数据源不存在')
        if normalized_column in (row_data or {}):
            insert_values[normalized_column] = (row_data or {}).get(normalized_column)

    if not insert_values:
        raise ValueError('没有可新增的列')

    insert_columns = list(insert_values.keys())
    insert_params = [insert_values[item] for item in insert_columns]
    column_sql = ', '.join([f"`{item}`" for item in insert_columns])
    placeholder_sql = ', '.join(['%s'] * len(insert_columns))
    sql = f"INSERT INTO `{source_table}` ({column_sql}) VALUES ({placeholder_sql})"

    with connection.cursor() as cursor:
        cursor.execute(sql, insert_params)


def _build_data_schedule_db_write_error_message(exc, action_name='操作'):
    normalized_action_name = _to_str(action_name) or '操作'
    error_message = _to_str(exc).lower()

    if isinstance(exc, IntegrityError):
        if any(token in error_message for token in ['duplicate', 'unique']):
            return f'{normalized_action_name}失败：存在重复数据（唯一约束冲突）'
        if any(token in error_message for token in ['cannot be null', "doesn't have a default value", 'null value']):
            missing_column = ''
            missing_column_match = re.search(r"column '([^']+)'", _to_str(exc), flags=re.IGNORECASE)
            if missing_column_match:
                missing_column = _to_str(missing_column_match.group(1))
            if missing_column:
                return f'{normalized_action_name}失败：必填字段缺失（{missing_column}）'
            return f'{normalized_action_name}失败：必填字段缺失'
        return f'{normalized_action_name}失败：字段约束校验未通过'

    if isinstance(exc, DataError):
        return f'{normalized_action_name}失败：字段值不合法'

    if isinstance(exc, OperationalError):
        if 'access denied' in error_message:
            return f'{normalized_action_name}失败：数据库账号认证失败'
        if 'command denied' in error_message:
            return f'{normalized_action_name}失败：数据库账号缺少写权限'
        return f'{normalized_action_name}失败：数据库连接不可用'

    return f'{normalized_action_name}失败，请检查数据库连接'


def _update_data_schedule_drilldown_row_to_db(payload, row, row_data):
    source_table = _to_str((payload or {}).get('sourceTable'))
    editable_columns = [item for item in ((payload or {}).get('editableColumns') or []) if _to_str(item)]
    source_locator = (row or {}).get('__sourceLocator') or []

    if not source_table or not _is_safe_sql_identifier(source_table):
        raise ValueError('字段数据源不合法')
    if not _table_exists(source_table):
        raise ValueError('字段数据源不存在')
    if not editable_columns:
        raise ValueError('当前字段不支持编辑')

    normalized_locator = _normalize_data_schedule_source_locator(source_table, source_locator)
    where_sql = ' AND '.join([f"`{column_name}` = %s" for column_name, _ in normalized_locator])
    locator_params = [locator_value for _, locator_value in normalized_locator]
    _ensure_single_data_schedule_locator_target(source_table, where_sql, locator_params)

    set_parts = []
    set_params = []
    for item in editable_columns:
        column_name = _to_str(item)
        if not column_name:
            continue
        if not _is_safe_sql_identifier(column_name):
            raise ValueError('字段数据源不合法')
        if not _table_has_column(source_table, column_name):
            raise ValueError('字段数据源不存在')
        if column_name not in (row_data or {}):
            continue

        set_parts.append(f"`{column_name}` = %s")
        set_params.append((row_data or {}).get(column_name))
    if not set_parts:
        raise ValueError('没有可更新的列')

    sql = (
        f"UPDATE `{source_table}` "
        f"SET {', '.join(set_parts)} "
        f"WHERE {where_sql}"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, [*set_params, *locator_params])


def _delete_data_schedule_drilldown_row_from_db(payload, row):
    source_table = _to_str((payload or {}).get('sourceTable'))
    source_locator = (row or {}).get('__sourceLocator') or []

    if not source_table or not _is_safe_sql_identifier(source_table):
        raise ValueError('字段数据源不合法')
    if not _table_exists(source_table):
        raise ValueError('字段数据源不存在')

    normalized_locator = _normalize_data_schedule_source_locator(source_table, source_locator)
    where_sql = ' AND '.join([f"`{column_name}` = %s" for column_name, _ in normalized_locator])
    locator_params = [locator_value for _, locator_value in normalized_locator]
    _ensure_single_data_schedule_locator_target(source_table, where_sql, locator_params)

    sql = f"DELETE FROM `{source_table}` WHERE {where_sql}"
    with connection.cursor() as cursor:
        cursor.execute(sql, locator_params)


def _update_data_schedule_field_value_to_db(field_item, field_value):
    source_table = _to_str((field_item or {}).get('sourceTable'))
    source_column = _to_str((field_item or {}).get('sourceColumn'))
    source_locator = (field_item or {}).get('sourceLocator') or []

    if not source_table or not source_column or not isinstance(source_locator, list):
        raise ValueError('字段缺少可更新的数据源信息')
    if not _is_data_schedule_detail_column_editable(source_column.lower()):
        raise ValueError('当前字段不支持编辑')
    if not (_is_safe_sql_identifier(source_table) and _is_safe_sql_identifier(source_column)):
        raise ValueError('字段数据源不合法')
    if not _table_exists(source_table):
        raise ValueError('字段数据源不存在')
    if not _table_has_column(source_table, source_column):
        raise ValueError('字段数据源不存在')

    normalized_locator = _normalize_data_schedule_source_locator(source_table, source_locator)

    where_sql = ' AND '.join([f"`{column_name}` = %s" for column_name, _ in normalized_locator])
    locator_params = [locator_value for _, locator_value in normalized_locator]
    _ensure_single_data_schedule_locator_target(source_table, where_sql, locator_params)

    sql = (
        f"UPDATE `{source_table}` "
        f"SET `{source_column}` = %s "
        f"WHERE {where_sql}"
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, [field_value, *locator_params])


def _get_data_schedule_field_item(task_id, field_key, bundle=None):
    normalized_field_key = str(field_key or '').strip()
    if not normalized_field_key:
        return None

    if bundle and isinstance(bundle, dict):
        all_fields = bundle.get('fields') or []
    else:
        all_fields = _get_data_schedule_fields(task_id)
    field_map = {item.get('fieldKey'): item for item in all_fields}
    return field_map.get(normalized_field_key)


def _normalize_drilldown_row_data(row_data, columns):
    if row_data is None:
        row_data = {}
    if not isinstance(row_data, dict):
        return None, 'rowData 格式错误'

    normalized = {}
    for column in columns:
        column_key = str(column.get('key') or '').strip()
        if not column_key:
            continue

        raw_value = row_data.get(column_key, '--')
        if raw_value is None:
            value = '--'
        else:
            value = str(raw_value)

        if len(value) > DATA_SCHEDULE_DRILLDOWN_ROW_VALUE_MAX_LENGTH:
            return None, f'{column_key} 长度不能超过{DATA_SCHEDULE_DRILLDOWN_ROW_VALUE_MAX_LENGTH}个字符'

        normalized[column_key] = value
    return normalized, ''


def _get_next_drilldown_row_id(records):
    max_id = 0
    for item in records:
        try:
            max_id = max(max_id, int(item.get('id')))
        except (TypeError, ValueError):
            continue
    return max_id + 1


def _find_drilldown_row(records, row_id):
    for item in records:
        try:
            if int(item.get('id')) == int(row_id):
                return item
        except (TypeError, ValueError):
            continue
    return None


def _filter_data_schedule_fields_by_scope(fields, scope):
    if scope == DATA_SCHEDULE_SCOPE_COMPLETE:
        return [item for item in fields if item.get('status') == DATA_SCHEDULE_SCOPE_COMPLETE]
    if scope == DATA_SCHEDULE_SCOPE_MISSING:
        return [item for item in fields if item.get('status') == DATA_SCHEDULE_SCOPE_MISSING]
    return fields


def _build_data_schedule_counts(fields):
    complete = len([item for item in fields if item.get('status') == DATA_SCHEDULE_SCOPE_COMPLETE])
    missing = len([item for item in fields if item.get('status') == DATA_SCHEDULE_SCOPE_MISSING])
    return {'all': len(fields), 'complete': complete, 'missing': missing}


def _resolve_data_schedule_log_project_id(task_id, project_id=''):
    normalized_project_id = _to_str(project_id)
    if normalized_project_id:
        return normalized_project_id
    return _resolve_data_schedule_project_id_from_task_id(task_id)


def _fetch_data_schedule_project_queue_rows(project_id):
    normalized_project_id = _to_str(project_id)
    if not normalized_project_id:
        return []
    if not _table_exists('c_r_cm_task_item_queue'):
        return []

    select_parts = [
        "resource_id AS resourceId",
        "project_id AS projectId",
        "file_id AS fileId",
        "status AS status",
        "step_no AS stepNo" if _table_has_column('c_r_cm_task_item_queue', 'step_no') else "NULL AS stepNo",
        "last_error AS lastError" if _table_has_column('c_r_cm_task_item_queue', 'last_error') else "NULL AS lastError",
        "updated_at AS updatedAt" if _table_has_column('c_r_cm_task_item_queue', 'updated_at') else "NULL AS updatedAt",
        "locked_by AS lockedBy" if _table_has_column('c_r_cm_task_item_queue', 'locked_by') else "NULL AS lockedBy",
        "locked_at AS lockedAt" if _table_has_column('c_r_cm_task_item_queue', 'locked_at') else "NULL AS lockedAt",
        "created_at AS createdAt" if _table_has_column('c_r_cm_task_item_queue', 'created_at') else "NULL AS createdAt",
        "create_time AS createTime" if _table_has_column('c_r_cm_task_item_queue', 'create_time') else "NULL AS createTime",
    ]

    order_time_columns = []
    if _table_has_column('c_r_cm_task_item_queue', 'locked_at'):
        order_time_columns.append('locked_at')
    if _table_has_column('c_r_cm_task_item_queue', 'updated_at'):
        order_time_columns.append('updated_at')
    if _table_has_column('c_r_cm_task_item_queue', 'created_at'):
        order_time_columns.append('created_at')
    if _table_has_column('c_r_cm_task_item_queue', 'create_time'):
        order_time_columns.append('create_time')

    if order_time_columns:
        order_sql = f" ORDER BY COALESCE({', '.join(order_time_columns)}) DESC, resource_id DESC"
    else:
        order_sql = " ORDER BY resource_id DESC"

    sql = (
        "SELECT "
        + ", ".join(select_parts)
        + " FROM c_r_cm_task_item_queue "
        "WHERE project_id = %s AND file_id LIKE 'PROJECT:%%'"
        + order_sql
    )
    with connection.cursor() as cursor:
        cursor.execute(sql, [normalized_project_id])
        return _dictfetchall(cursor)


def _infer_data_schedule_log_level(status_value, last_error):
    normalized_status = _to_str(status_value).upper()
    normalized_last_error = _to_str(last_error)
    if normalized_last_error or normalized_status in {'FAILED', 'PARTIAL_FAILED', 'ERROR'}:
        return DATA_SCHEDULE_LOG_LEVEL_ERROR
    if normalized_status in {'PAUSED', 'STOPPED', 'CANCELLED', 'CANCEL'}:
        return DATA_SCHEDULE_LOG_LEVEL_WARN
    return DATA_SCHEDULE_LOG_LEVEL_INFO


def _build_data_schedule_queue_log_timestamp(row):
    for key in ('updatedAt', 'lockedAt', 'createdAt', 'createTime'):
        parsed = _parse_datetime_value((row or {}).get(key))
        if parsed:
            return parsed
    return None


def _build_data_schedule_queue_log_message(row, project_id):
    message_parts = [f'project_id={_to_str(project_id)}']
    task_id = _to_str((row or {}).get('resourceId'))
    if task_id:
        message_parts.append(f'task_id={task_id}')
    file_id = _to_str((row or {}).get('fileId'))
    if file_id:
        message_parts.append(f'file_id={file_id}')
    step_no = (row or {}).get('stepNo')
    if step_no not in (None, ''):
        message_parts.append(f'step_no={step_no}')
    status_value = _to_str((row or {}).get('status'))
    if status_value:
        message_parts.append(f'status={status_value}')
    locked_by = _to_str((row or {}).get('lockedBy'))
    if locked_by:
        message_parts.append(f'node={locked_by}')
    last_error = _to_str((row or {}).get('lastError'))
    if last_error:
        message_parts.append(f'error={last_error}')
    return '; '.join(message_parts)


def _map_data_schedule_queue_rows_to_logs(rows, project_id):
    records = []
    for index, row in enumerate((rows or []), start=1):
        timestamp = _build_data_schedule_queue_log_timestamp(row)
        time_text = timestamp.strftime(DATA_SCHEDULE_LOG_TIME_FORMAT) if timestamp else '--'
        records.append({
            '_timestamp': timestamp,
            'component': _to_str((row or {}).get('resourceId')) or 'task-item-queue',
            'id': index,
            'level': _infer_data_schedule_log_level((row or {}).get('status'), (row or {}).get('lastError')),
            'message': _build_data_schedule_queue_log_message(row, project_id),
            'time': time_text
        })
    return records


def _get_data_schedule_log_meta(task_id, project_id=''):
    normalized_project_id = _resolve_data_schedule_log_project_id(task_id, project_id)
    project_row = _fetch_data_schedule_project_base_row(normalized_project_id) if normalized_project_id else None
    queue_rows = _fetch_data_schedule_project_queue_rows(normalized_project_id) if normalized_project_id else []
    latest_row = queue_rows[0] if queue_rows else {}

    normalized_task_id = _to_str(task_id) or _to_str(latest_row.get('resourceId'))
    if not normalized_task_id and normalized_project_id:
        normalized_task_id = f'PROJECT:{normalized_project_id}'

    return {
        'nodeName': _to_str(latest_row.get('lockedBy')) or '--',
        'projectName': _to_str((project_row or {}).get('projectName')) or normalized_project_id or '--',
        'serviceName': 'c_r_cm_task_item_queue',
        'taskId': normalized_task_id,
    }


def _build_data_schedule_log_records(task_id, project_id=''):
    normalized_project_id = _resolve_data_schedule_log_project_id(task_id, project_id)
    if not normalized_project_id:
        return []
    queue_rows = _fetch_data_schedule_project_queue_rows(normalized_project_id)
    return _map_data_schedule_queue_rows_to_logs(queue_rows, normalized_project_id)


def _parse_data_schedule_log_datetime(raw_value, *, end_of_day=False):
    value = str(raw_value or '').strip()
    if not value:
        return None, ''

    formats = [
        '%Y-%m-%d %H:%M:%S',
        '%Y-%m-%d %H:%M',
        '%Y-%m-%d',
    ]
    for fmt in formats:
        try:
            parsed = datetime.strptime(value, fmt)
        except ValueError:
            continue

        if fmt == '%Y-%m-%d' and end_of_day:
            parsed = parsed.replace(hour=23, minute=59, second=59)

        return parsed, ''

    return None, '时间格式错误，请使用 YYYY-MM-DD 或 YYYY-MM-DD HH:mm:ss'


def _parse_data_schedule_log_filters(request):
    level = str(request.query_params.get('level') or '').strip().lower()
    if level and level not in DATA_SCHEDULE_LOG_LEVEL_SET:
        return None, 'level 参数无效，仅支持 info/warn/error'

    order_by = str(request.query_params.get('orderBy') or DATA_SCHEDULE_LOG_ORDER_BY_TIME).strip().lower()
    if not order_by:
        order_by = DATA_SCHEDULE_LOG_ORDER_BY_TIME
    if order_by != DATA_SCHEDULE_LOG_ORDER_BY_TIME:
        return None, 'orderBy 参数无效，仅支持 time'

    order_direction = str(request.query_params.get('orderDirection') or DATA_SCHEDULE_LOG_ORDER_DIRECTION_DESC).strip().lower()
    if not order_direction:
        order_direction = DATA_SCHEDULE_LOG_ORDER_DIRECTION_DESC
    if order_direction not in DATA_SCHEDULE_LOG_ORDER_DIRECTION_SET:
        return None, 'orderDirection 参数无效，仅支持 asc/desc'

    count = _parse_positive_int(
        request.query_params.get('count'),
        default=DATA_SCHEDULE_LOG_DEFAULT_COUNT,
        max_value=DATA_SCHEDULE_LOG_MAX_COUNT
    )
    keyword = str(request.query_params.get('keyword') or '').strip()

    start_time, start_error = _parse_data_schedule_log_datetime(request.query_params.get('startTime'))
    if start_error:
        return None, start_error
    end_time, end_error = _parse_data_schedule_log_datetime(request.query_params.get('endTime'), end_of_day=True)
    if end_error:
        return None, end_error

    if start_time and end_time and start_time > end_time:
        return None, '开始时间不能晚于结束时间'

    return {
        'count': count,
        'endTime': end_time,
        'keyword': keyword,
        'level': level,
        'orderBy': order_by,
        'orderDirection': order_direction,
        'startTime': start_time
    }, ''


def _filter_data_schedule_logs(records, filters):
    keyword = (filters.get('keyword') or '').lower()
    level = filters.get('level') or ''
    start_time = filters.get('startTime')
    end_time = filters.get('endTime')

    filtered = []
    for record in records:
        if level and record.get('level') != level:
            continue

        if keyword:
            message = str(record.get('message') or '').lower()
            component = str(record.get('component') or '').lower()
            if keyword not in message and keyword not in component:
                continue

        record_time = record.get('_timestamp')
        if start_time and record_time and record_time < start_time:
            continue
        if end_time and record_time and record_time > end_time:
            continue

        filtered.append(record)

    return filtered


def _sort_data_schedule_logs(records, order_direction):
    reverse = order_direction != DATA_SCHEDULE_LOG_ORDER_DIRECTION_ASC
    return sorted(records, key=lambda item: item.get('_timestamp') or datetime.min, reverse=reverse)


def _serialize_data_schedule_logs(records):
    return [
        {
            'component': record.get('component') or '',
            'id': record.get('id'),
            'level': record.get('level') or DATA_SCHEDULE_LOG_LEVEL_INFO,
            'message': record.get('message') or '',
            'time': record.get('time') or ''
        }
        for record in records
    ]


def _build_data_schedule_log_export_binary(task_meta, filters, records):
    try:
        from openpyxl import Workbook
    except Exception as exc:
        logger.exception('openpyxl import failed when export log: %s', exc)
        raise

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = '数据提取日志'

    worksheet.append(['字段', '内容'])
    worksheet.append(['任务标识', task_meta.get('taskId') or ''])
    worksheet.append(['项目名称', task_meta.get('projectName') or ''])
    worksheet.append(['服务名称', task_meta.get('serviceName') or ''])
    worksheet.append(['执行节点', task_meta.get('nodeName') or ''])
    worksheet.append(['日志级别', filters.get('level') or '全部'])
    worksheet.append(['排序方向', filters.get('orderDirection') or DATA_SCHEDULE_LOG_ORDER_DIRECTION_DESC])
    worksheet.append(['查询条数', filters.get('count') or DATA_SCHEDULE_LOG_DEFAULT_COUNT])
    worksheet.append(['开始时间', (filters.get('startTime') or '').strftime(DATA_SCHEDULE_LOG_TIME_FORMAT) if filters.get('startTime') else '--'])
    worksheet.append(['结束时间', (filters.get('endTime') or '').strftime(DATA_SCHEDULE_LOG_TIME_FORMAT) if filters.get('endTime') else '--'])
    worksheet.append(['关键字', filters.get('keyword') or '--'])
    worksheet.append([])
    worksheet.append(['时间', '组件', '等级', '日志'])

    for record in records:
        worksheet.append([
            _to_excel_cell_value(record.get('time') or ''),
            _to_excel_cell_value(record.get('component') or ''),
            _to_excel_cell_value(record.get('level') or DATA_SCHEDULE_LOG_LEVEL_INFO),
            _to_excel_cell_value(record.get('message') or ''),
        ])

    _style_data_schedule_log_sheet(worksheet, data_start_row=13)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output.getvalue()


def _safe_excel_sheet_name(raw_name, used_names):
    invalid_chars = ['\\', '/', '*', '?', ':', '[', ']']
    safe = raw_name
    for char in invalid_chars:
        safe = safe.replace(char, '_')
    safe = (safe or 'sheet').strip()
    safe = safe[:31]
    if not safe:
        safe = 'sheet'

    candidate = safe
    suffix = 1
    while candidate in used_names:
        suffix_text = f'_{suffix}'
        candidate = f'{safe[:31 - len(suffix_text)]}{suffix_text}'
        suffix += 1

    used_names.add(candidate)
    return candidate


def _build_excel_internal_link(sheet_name, target_cell='A1'):
    normalized_sheet_name = str(sheet_name or '').replace("'", "''")
    normalized_cell = str(target_cell or 'A1').strip() or 'A1'
    return f"#'{normalized_sheet_name}'!{normalized_cell}"


def _style_data_schedule_summary_sheet(sheet):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    sheet.column_dimensions['A'].width = 18
    sheet.column_dimensions['B'].width = 56

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    label_fill = PatternFill(fill_type='solid', fgColor='EEF3FA')
    header_font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
    body_font = Font(name='Microsoft YaHei', color='1F2937')
    label_font = Font(name='Microsoft YaHei', bold=True, color='1F4E78')
    align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    for column in range(1, 3):
        cell = sheet.cell(row=1, column=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align
        cell.border = thin_border

    for row in range(2, sheet.max_row + 1):
        for column in range(1, 3):
            cell = sheet.cell(row=row, column=column)
            cell.font = body_font
            cell.alignment = align
            cell.border = thin_border
            if column == 1:
                cell.fill = label_fill
                cell.font = label_font

    sheet.sheet_view.showGridLines = False


def _style_data_schedule_field_sheet(sheet):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    sheet.column_dimensions['A'].width = 26
    sheet.column_dimensions['B'].width = 56
    sheet.column_dimensions['C'].width = 14

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    header_font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
    body_font = Font(name='Microsoft YaHei', color='111827')
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    for column in range(1, 4):
        cell = sheet.cell(row=1, column=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row in range(2, sheet.max_row + 1):
        for column in range(1, 4):
            cell = sheet.cell(row=row, column=column)
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = left_align if column in {1, 2} else center_align

        status_cell = sheet.cell(row=row, column=3)
        status_value = str(status_cell.value or '')
        if status_value == '提取完整':
            status_cell.fill = PatternFill(fill_type='solid', fgColor='E8F5E9')
            status_cell.font = Font(name='Microsoft YaHei', bold=True, color='1B5E20')
        elif status_value == '提取缺失':
            status_cell.fill = PatternFill(fill_type='solid', fgColor='FDECEA')
            status_cell.font = Font(name='Microsoft YaHei', bold=True, color='B42318')

        value_cell = sheet.cell(row=row, column=2)
        if str(value_cell.value or '') == '查看':
            value_cell.font = Font(name='Microsoft YaHei', color='0563C1', underline='single')
            value_cell.fill = PatternFill(fill_type='solid', fgColor='EEF6FF')

    sheet.auto_filter.ref = f'A1:C{max(1, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False


def _style_data_schedule_drilldown_sheet(sheet):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    header_fill = PatternFill(fill_type='solid', fgColor='2F5597')
    header_font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
    body_font = Font(name='Microsoft YaHei', color='111827')
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    has_action_column = str(sheet.cell(row=1, column=sheet.max_column).value or '') == '操作'

    for column in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=1, column=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row in range(2, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row, column=column)
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = center_align if has_action_column and column == sheet.max_column else left_align

    for column in range(1, sheet.max_column + 1):
        max_length = 0
        for row in range(1, min(sheet.max_row, 200) + 1):
            value = sheet.cell(row=row, column=column).value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(max_length + 4, 12), 48)

    sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{max(1, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False


def _style_data_schedule_log_sheet(sheet, data_start_row=13):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    sheet.column_dimensions['A'].width = 22
    sheet.column_dimensions['B'].width = 18
    sheet.column_dimensions['C'].width = 12
    sheet.column_dimensions['D'].width = 82

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    label_fill = PatternFill(fill_type='solid', fgColor='EEF3FA')
    header_font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
    body_font = Font(name='Microsoft YaHei', color='111827')
    label_font = Font(name='Microsoft YaHei', bold=True, color='1F4E78')
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    top_left_align = Alignment(horizontal='left', vertical='top', wrap_text=True)

    for column in range(1, 3):
        cell = sheet.cell(row=1, column=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row in range(2, 12):
        label_cell = sheet.cell(row=row, column=1)
        value_cell = sheet.cell(row=row, column=2)
        label_cell.fill = label_fill
        label_cell.font = label_font
        label_cell.alignment = left_align
        label_cell.border = thin_border
        value_cell.font = body_font
        value_cell.alignment = left_align
        value_cell.border = thin_border

    for column in range(1, 5):
        cell = sheet.cell(row=data_start_row, column=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row in range(data_start_row + 1, sheet.max_row + 1):
        for column in range(1, 5):
            cell = sheet.cell(row=row, column=column)
            cell.font = body_font
            cell.border = thin_border
            if column in {1, 2, 3}:
                cell.alignment = center_align if column != 2 else left_align
            else:
                cell.alignment = top_left_align

        level_cell = sheet.cell(row=row, column=3)
        level_value = str(level_cell.value or '').lower()
        if level_value == 'info':
            level_cell.fill = PatternFill(fill_type='solid', fgColor='E8F5E9')
            level_cell.font = Font(name='Microsoft YaHei', bold=True, color='1B5E20')
        elif level_value == 'warn':
            level_cell.fill = PatternFill(fill_type='solid', fgColor='FFF7E6')
            level_cell.font = Font(name='Microsoft YaHei', bold=True, color='9A6700')
        elif level_value == 'error':
            level_cell.fill = PatternFill(fill_type='solid', fgColor='FDECEA')
            level_cell.font = Font(name='Microsoft YaHei', bold=True, color='B42318')

    sheet.freeze_panes = f'A{data_start_row + 1}'
    sheet.auto_filter.ref = f'A{data_start_row}:D{max(data_start_row, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False


def _style_data_query_pivot_export_sheet(sheet):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    thin_border = Border(
        left=Side(style='thin', color='E5E7EB'),
        right=Side(style='thin', color='E5E7EB'),
        top=Side(style='thin', color='E5E7EB'),
        bottom=Side(style='thin', color='E5E7EB')
    )
    header_fill = PatternFill(fill_type='solid', fgColor='1F4E78')
    header_font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF')
    body_font = Font(name='Microsoft YaHei', color='111827')
    link_font = Font(name='Microsoft YaHei', color='0563C1', underline='single')
    center_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    for column in range(1, sheet.max_column + 1):
        cell = sheet.cell(row=1, column=column)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    for row in range(2, sheet.max_row + 1):
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row, column=column)
            cell.font = body_font
            cell.border = thin_border
            cell.alignment = left_align if column == 1 else center_align

            if str(cell.value or '') == '查看':
                cell.font = link_font
                cell.fill = PatternFill(fill_type='solid', fgColor='EEF6FF')

    for column in range(1, sheet.max_column + 1):
        max_length = 0
        for row in range(1, min(sheet.max_row, 200) + 1):
            value = sheet.cell(row=row, column=column).value
            if value is None:
                continue
            max_length = max(max_length, len(str(value)))
        sheet.column_dimensions[get_column_letter(column)].width = min(max(max_length + 4, 14), 48)

    sheet.auto_filter.ref = f'A1:{get_column_letter(sheet.max_column)}{max(1, sheet.max_row)}'
    sheet.sheet_view.showGridLines = False


def _normalize_data_schedule_scope(raw_scope):
    scope = str(raw_scope or DATA_SCHEDULE_SCOPE_ALL).strip().lower()
    if not scope:
        scope = DATA_SCHEDULE_SCOPE_ALL
    if scope not in DATA_SCHEDULE_SCOPE_SET:
        return None
    return scope


def _to_data_schedule_field_record(field):
    return {
        'canDrilldown': bool(field.get('canDrilldown')),
        'canEdit': bool(field.get('canEdit')),
        'drilldownLabel': '查看' if field.get('canDrilldown') else '',
        'fieldKey': field.get('fieldKey') or '',
        'fieldName': field.get('fieldName') or '',
        'fieldValueDisplay': field.get('fieldValueDisplay') or '--',
        'status': field.get('status') or DATA_SCHEDULE_SCOPE_MISSING
    }


def _paginate_data_schedule_records(records, current, size):
    total = len(records)
    start = (current - 1) * size
    end = start + size
    return {
        'current': current,
        'records': records[start:end],
        'size': size,
        'total': total
    }


def _build_data_schedule_export_binary(task_id, scope):
    try:
        from openpyxl import Workbook
        from openpyxl.worksheet.hyperlink import Hyperlink
    except Exception as exc:
        logger.exception('openpyxl import failed: %s', exc)
        raise

    all_fields = _get_data_schedule_fields(task_id)
    filtered_fields = _filter_data_schedule_fields_by_scope(all_fields, scope)
    summary = _get_data_schedule_summary(task_id)
    counts = _build_data_schedule_counts(all_fields)
    summary['counts'] = counts

    workbook = Workbook()
    used_sheet_names = set()
    summary_sheet = workbook.active
    summary_sheet.title = _safe_excel_sheet_name('任务概览', used_sheet_names)
    summary_sheet.append(['字段', '内容'])
    summary_rows = [
        ('任务标识', summary.get('taskId') or ''),
        ('项目名称', summary.get('projectName') or ''),
        ('招标编号', summary.get('bidNo') or ''),
        ('创建人', summary.get('creator') or ''),
        ('创建时间', summary.get('createTime') or ''),
        ('执行状态', summary.get('status') or ''),
        ('执行进度', f"{summary.get('progress') or 0}%"),
        ('全部字段', counts.get('all') or 0),
        ('提取完整', counts.get('complete') or 0),
        ('提取缺失', counts.get('missing') or 0),
    ]
    for row in summary_rows:
        summary_sheet.append(list(row))

    field_sheet = workbook.create_sheet(_safe_excel_sheet_name(f'字段详情_{scope}', used_sheet_names))
    field_sheet.append(['字段名称', '字段内容', '状态'])

    drilldown_sheet_map = {}
    for field in filtered_fields:
        if not field.get('canDrilldown'):
            continue
        field_key = field.get('fieldKey') or ''
        _, drilldown_payload, err_msg, _, _ = _get_data_schedule_drilldown_payload(task_id, field_key)
        if err_msg or not drilldown_payload:
            continue
        sheet_name = _safe_excel_sheet_name(f"下钻_{field.get('fieldName') or field_key}", used_sheet_names)
        drilldown_sheet_map[field_key] = sheet_name
        drilldown_sheet = workbook.create_sheet(sheet_name)
        columns = drilldown_payload.get('columns') or []
        drilldown_sheet.append([col.get('title') or col.get('key') or '' for col in columns])
        for row in drilldown_payload.get('records') or []:
            row_values = [_to_excel_cell_value(row.get(col.get('key') or '', '--')) for col in columns]
            drilldown_sheet.append(row_values)
        drilldown_sheet.freeze_panes = 'A2'
        _style_data_schedule_drilldown_sheet(drilldown_sheet)

    for index, field in enumerate(filtered_fields, start=2):
        status_value = field.get('status')
        if status_value == DATA_SCHEDULE_SCOPE_COMPLETE:
            status_text = '提取完整'
        elif status_value == DATA_SCHEDULE_SCOPE_MISSING:
            status_text = '提取缺失'
        else:
            status_text = status_value or ''

        field_sheet.append([
            _to_excel_cell_value(field.get('fieldName') or ''),
            _to_excel_cell_value(field.get('fieldValueDisplay') or '--'),
            _to_excel_cell_value(status_text)
        ])

        field_key = field.get('fieldKey') or ''
        if field_key in drilldown_sheet_map:
            value_cell = field_sheet.cell(row=index, column=2)
            value_cell.value = '查看'
            value_cell.hyperlink = Hyperlink(
                ref=value_cell.coordinate,
                location=_build_excel_internal_link(drilldown_sheet_map[field_key], 'A1').lstrip('#'),
                display='查看'
            )

    field_sheet.freeze_panes = 'A2'
    _style_data_schedule_summary_sheet(summary_sheet)
    _style_data_schedule_field_sheet(field_sheet)

    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output.getvalue()


def success_response(data=None, msg='ok', http_status=status.HTTP_200_OK):
    return Response({'code': SUCCESS_CODE, 'data': data, 'msg': msg}, status=http_status)


def error_response(msg, code=ERROR_CODE_INVALID_PARAMS, http_status=status.HTTP_400_BAD_REQUEST):
    return Response({'code': code, 'data': None, 'msg': msg}, status=http_status)

_WATERMARK_HIDDEN_MODE_COLUMN_EXISTS = None


def _supports_watermark_hidden_mode():
    global _WATERMARK_HIDDEN_MODE_COLUMN_EXISTS

    if _WATERMARK_HIDDEN_MODE_COLUMN_EXISTS is not None:
        return _WATERMARK_HIDDEN_MODE_COLUMN_EXISTS

    try:
        with connection.cursor() as cursor:
            columns = connection.introspection.get_table_description(cursor, WebsiteSetting._meta.db_table)
        _WATERMARK_HIDDEN_MODE_COLUMN_EXISTS = any(
            getattr(column, 'name', '') == 'watermark_hidden_mode'
            for column in columns
        )
    except Exception as exc:
        logger.warning('detect watermark_hidden_mode column failed, fallback to disabled: %s', exc)
        _WATERMARK_HIDDEN_MODE_COLUMN_EXISTS = False

    return _WATERMARK_HIDDEN_MODE_COLUMN_EXISTS


def _clean_username(raw_username):
    return (raw_username or '').strip()


def _normalize_system_manage_user_status(raw_status):
    if isinstance(raw_status, bool):
        return raw_status

    normalized = _to_str(raw_status).lower()
    if not normalized:
        return None

    if normalized in {'1', 'true', 'enabled', 'enable', 'active'}:
        return True
    if normalized in {'2', 'false', 'disabled', 'disable', 'inactive'}:
        return False

    return None


def _resolve_system_manage_user_update_time(user):
    for field_name in ('update_time', 'updated_at', 'modified_at', 'last_login', 'date_joined'):
        value = getattr(user, field_name, None)
        if value:
            return value
    return None


def _serialize_system_manage_user_row(user):
    create_time = _format_datetime_text(getattr(user, 'date_joined', None)) or ''
    update_time = _format_datetime_text(_resolve_system_manage_user_update_time(user)) or create_time

    return {
        'id': int(getattr(user, 'id', 0) or 0),
        'userName': _to_str(getattr(user, 'username', '')),
        'status': SYSTEM_MANAGE_USER_STATUS_ENABLED if bool(getattr(user, 'is_active', False)) else SYSTEM_MANAGE_USER_STATUS_DISABLED,
        'userEmail': _to_str(getattr(user, 'email', '')),
        'createTime': create_time,
        'updateTime': update_time,
    }


def _apply_system_manage_user_password(user, password):
    plain_text = str(password or '')
    if plain_text:
        user.set_password(plain_text)


def _validate_system_manage_user_payload(raw_payload, is_create):
    payload = raw_payload if isinstance(raw_payload, dict) else {}

    username = _clean_username(payload.get('userName') or payload.get('username'))
    if not username:
        return None, '用户名不能为空'
    if len(username) > 150:
        return None, '用户名长度不能超过150个字符'

    email = _to_str(payload.get('email') or payload.get('userEmail'))
    if len(email) > 254:
        return None, '邮箱长度不能超过254个字符'

    raw_status = payload.get('status')
    if raw_status is None:
        raw_status = payload.get('userStatus')
    status_value = _normalize_system_manage_user_status(raw_status)
    if status_value is None:
        return None, '状态参数无效，仅支持 1(可用)/2(禁用)'

    password = str(payload.get('password') or '')
    confirm_password = str(payload.get('confirmPassword') or payload.get('confirm_password') or '')
    if is_create and not password:
        return None, '密码不能为空'
    if password or confirm_password:
        if password != confirm_password:
            return None, '两次输入的密码不一致'
        try:
            validate_password(password)
        except ValidationError as exc:
            return None, exc.messages[0]

    return {
        'email': email,
        'is_active': status_value,
        'password': password,
        'username': username,
    }, ''


def _create_refresh_token(user):
    payload = {
        'passwordFingerprint': hashlib.sha256(user.password.encode('utf-8')).hexdigest(),
        'uid': user.id,
        'username': user.username
    }
    return signing.dumps(payload, salt=REFRESH_TOKEN_SALT)


def _parse_refresh_token(refresh_token):
    return signing.loads(refresh_token, max_age=REFRESH_TOKEN_MAX_AGE, salt=REFRESH_TOKEN_SALT)


def _issue_auth_tokens(user):
    # Rotate the access token on each issue to avoid long-lived leaked tokens.
    Token.objects.filter(user=user).delete()
    token = Token.objects.create(user=user)

    return {'refreshToken': _create_refresh_token(user), 'token': token.key}


def _generate_captcha_code():
    return ''.join(secrets.choice(string.digits) for _ in range(4))


def _create_captcha_id(code):
    nonce_chars = string.ascii_letters + string.digits
    nonce = ''.join(secrets.choice(nonce_chars) for _ in range(16))
    digest = hashlib.sha256(f'{code}:{nonce}:{settings.SECRET_KEY}'.encode('utf-8')).hexdigest()

    captcha_id = signing.dumps({'nonce': nonce, 'digest': digest}, salt=CAPTCHA_TOKEN_SALT)
    if CAPTCHA_SINGLE_USE:
        cache.set(_captcha_cache_key(nonce), 1, timeout=CAPTCHA_MAX_AGE)

    return captcha_id


def _captcha_cache_key(nonce):
    nonce_hash = hashlib.sha256(str(nonce).encode('utf-8')).hexdigest()
    return f'api:captcha:nonce:{nonce_hash}'


def _verify_captcha(captcha_id, captcha_code):
    if not captcha_id or not captcha_code:
        return False, '验证码不能为空'

    if len(captcha_code) != 4 or not captcha_code.isdigit():
        return False, '验证码格式错误'

    try:
        payload = signing.loads(captcha_id, max_age=CAPTCHA_MAX_AGE, salt=CAPTCHA_TOKEN_SALT)
    except signing.BadSignature:
        return False, '验证码无效或已过期'

    nonce = payload.get('nonce') or ''
    digest = payload.get('digest') or ''

    if CAPTCHA_SINGLE_USE:
        captcha_key = _captcha_cache_key(nonce)
        if not cache.get(captcha_key):
            return False, '验证码无效或已过期'
        # Consume captcha before verification to prevent brute-force and replay attacks.
        cache.delete(captcha_key)

    expected_digest = hashlib.sha256(f'{captcha_code}:{nonce}:{settings.SECRET_KEY}'.encode('utf-8')).hexdigest()
    if not hmac.compare_digest(str(digest), expected_digest):
        return False, '验证码错误'

    return True, ''


def _build_captcha_svg(code):
    width = 120
    height = 40

    lines = []
    for _ in range(5):
        x1, y1 = secrets.randbelow(width + 1), secrets.randbelow(height + 1)
        x2, y2 = secrets.randbelow(width + 1), secrets.randbelow(height + 1)
        lines.append(
            f'<line x1="{x1}" y1="{y1}" x2="{x2}" y2="{y2}" stroke="#c7d2fe" stroke-width="1" />'
        )

    texts = []
    for i, char in enumerate(code):
        x = 18 + i * 24 + (secrets.randbelow(5) - 2)
        y = 28 + (secrets.randbelow(5) - 2)
        rotate = secrets.randbelow(31) - 15
        colors = ['#1e3a8a', '#1d4ed8', '#334155', '#0f172a']
        color = colors[secrets.randbelow(len(colors))]
        texts.append(
            f'<text x="{x}" y="{y}" fill="{color}" font-size="22" '
            f'font-family="Arial, sans-serif" transform="rotate({rotate} {x} {y})">{char}</text>'
        )

    svg = (
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">'
        '<rect width="100%" height="100%" rx="6" ry="6" fill="#f8fafc" stroke="#cbd5e1" />'
        f'{"".join(lines)}{"".join(texts)}'
        '</svg>'
    )
    encoded = base64.b64encode(svg.encode('utf-8')).decode('utf-8')

    return f'data:image/svg+xml;base64,{encoded}'


def _build_user_info(user):
    if user.is_superuser:
        role = 'R_SUPER'
    elif user.is_staff:
        role = 'R_ADMIN'
    else:
        role = 'R_USER'

    return {
        'buttons': ROLE_BUTTONS_MAP.get(role, []),
        'roles': [role],
        'userId': str(user.id),
        'userName': user.username
    }


def _extract_auth_token(request):
    auth_header = (request.headers.get('Authorization') or '').strip()
    if not auth_header:
        return ''

    parts = auth_header.split()
    if len(parts) == 2 and parts[0].lower() in {'bearer', 'token'}:
        return parts[1]

    return ''


def _extract_client_ip(request):
    forwarded_for = (request.META.get('HTTP_X_FORWARDED_FOR') or '').split(',')
    if forwarded_for and forwarded_for[0].strip():
        return forwarded_for[0].strip()
    return (request.META.get('REMOTE_ADDR') or '').strip() or 'unknown'


def _build_login_subject(username, client_ip):
    raw = f'{(username or "").lower()}|{client_ip}'
    return hashlib.sha256(raw.encode('utf-8')).hexdigest()


def _login_fail_key(subject):
    return f'api:login:fail:{subject}'


def _login_lock_key(subject):
    return f'api:login:lock:{subject}'


def _get_login_lock_remaining(subject):
    lock_until = cache.get(_login_lock_key(subject))
    if not lock_until:
        return 0
    try:
        remaining = int(lock_until) - int(time.time())
    except (TypeError, ValueError):
        cache.delete(_login_lock_key(subject))
        return 0

    if remaining <= 0:
        cache.delete(_login_lock_key(subject))
        return 0

    return remaining


def _register_login_failure(subject):
    fail_key = _login_fail_key(subject)
    failures = int(cache.get(fail_key) or 0) + 1
    cache.set(fail_key, failures, timeout=LOGIN_FAIL_WINDOW_SECONDS)

    if failures >= LOGIN_MAX_FAIL_COUNT:
        lock_until = int(time.time()) + int(LOGIN_LOCK_SECONDS)
        cache.set(_login_lock_key(subject), lock_until, timeout=LOGIN_LOCK_SECONDS)
        cache.delete(fail_key)
        return LOGIN_LOCK_SECONDS

    return 0


def _clear_login_failure_state(subject):
    cache.delete(_login_fail_key(subject))
    cache.delete(_login_lock_key(subject))


def _validate_login_payload(username, password, captcha_id):
    if not username or not password:
        return '用户名和密码不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST

    if len(username) > LOGIN_MAX_USERNAME_LENGTH:
        return '账号长度超出限制', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST

    if len(password) > LOGIN_MAX_PASSWORD_LENGTH:
        return '密码长度超出限制', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST

    if len(captcha_id) > LOGIN_MAX_CAPTCHA_ID_LENGTH:
        return '验证码无效或已过期', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST

    return None


def _is_minio_enabled():
    return bool(getattr(settings, 'MINIO_ENABLED', False))


def _get_minio_client():
    return Minio(
        getattr(settings, 'MINIO_ENDPOINT', ''),
        access_key=getattr(settings, 'MINIO_ACCESS_KEY', ''),
        secret_key=getattr(settings, 'MINIO_SECRET_KEY', ''),
        secure=bool(getattr(settings, 'MINIO_SECURE', False))
    )


def _ensure_minio_bucket(client):
    bucket = getattr(settings, 'MINIO_BUCKET', '')
    if not bucket:
        raise ValueError('MINIO_BUCKET is empty')

    if cache.get(_MINIO_BUCKET_READY_CACHE_KEY):
        return

    if not client.bucket_exists(bucket):
        client.make_bucket(bucket)

    bucket_policy = {
        'Version': '2012-10-17',
        'Statement': [{
            'Effect': 'Allow',
            'Principal': {'AWS': ['*']},
            'Action': ['s3:GetObject'],
            'Resource': [f'arn:aws:s3:::{bucket}/*']
        }]
    }
    client.set_bucket_policy(bucket, json.dumps(bucket_policy))
    cache.set(_MINIO_BUCKET_READY_CACHE_KEY, 1, timeout=24 * 60 * 60)


def _minio_object_url(object_name):
    bucket = getattr(settings, 'MINIO_BUCKET', '')
    public_base_url = (getattr(settings, 'MINIO_PUBLIC_BASE_URL', '') or '').rstrip('/')
    if public_base_url:
        return f'{public_base_url}/{bucket}/{object_name}'

    scheme = 'https' if getattr(settings, 'MINIO_SECURE', False) else 'http'
    endpoint = (getattr(settings, 'MINIO_ENDPOINT', '') or '').rstrip('/')
    return f'{scheme}://{endpoint}/{bucket}/{object_name}'


def _guess_image_extension(content_type, filename=''):
    normalized_type = (content_type or '').lower()
    extension = WEBSITE_IMAGE_EXTENSION_MAP.get(normalized_type, '')
    if extension:
        return extension

    suffix = Path(filename or '').suffix.lower().lstrip('.')
    if suffix == 'jfif':
        return 'jpg'
    if suffix == 'jpeg':
        return 'jpg'
    if suffix in {'avif', 'bmp', 'gif', 'heic', 'heif', 'ico', 'jpg', 'png', 'svg', 'tif', 'tiff', 'webp'}:
        return suffix

    return ''


def _calculate_rgb_distance(color1, color2):
    return (
        abs(int(color1[0]) - int(color2[0]))
        + abs(int(color1[1]) - int(color2[1]))
        + abs(int(color1[2]) - int(color2[2]))
    )


def _estimate_background_rgb(image):
    width, height = image.size
    if width <= 0 or height <= 0:
        return 255, 255, 255

    step = max(min(width, height) // 120, 1)
    border_samples = []
    for x in range(0, width, step):
        border_samples.append(image.getpixel((x, 0)))
        border_samples.append(image.getpixel((x, height - 1)))
    for y in range(0, height, step):
        border_samples.append(image.getpixel((0, y)))
        border_samples.append(image.getpixel((width - 1, y)))

    if not border_samples:
        border_samples = [
            image.getpixel((0, 0)),
            image.getpixel((max(width - 1, 0), 0)),
            image.getpixel((0, max(height - 1, 0))),
            image.getpixel((max(width - 1, 0), max(height - 1, 0)))
        ]

    return (
        sum(int(pixel[0]) for pixel in border_samples) // len(border_samples),
        sum(int(pixel[1]) for pixel in border_samples) // len(border_samples),
        sum(int(pixel[2]) for pixel in border_samples) // len(border_samples)
    )


def _is_background_dominant_on_border(image, background_rgb, tolerance):
    width, height = image.size
    if width <= 0 or height <= 0:
        return False

    step = max(min(width, height) // 120, 1)
    border_points = []
    for x in range(0, width, step):
        border_points.append((x, 0))
        border_points.append((x, height - 1))
    for y in range(0, height, step):
        border_points.append((0, y))
        border_points.append((width - 1, y))

    if not border_points:
        return False

    matched = 0
    for x, y in border_points:
        pixel = image.getpixel((x, y))
        if _calculate_rgb_distance(pixel, background_rgb) <= tolerance:
            matched += 1

    return (matched / len(border_points)) >= 0.55


def _build_edge_seed_points(width, height, step):
    seeds = []
    for x in range(0, width, step):
        seeds.append((x, 0))
        seeds.append((x, height - 1))
    for y in range(0, height, step):
        seeds.append((0, y))
        seeds.append((width - 1, y))
    return seeds


def _transparentize_connected_background(image, background_rgb, tolerance):
    width, height = image.size
    if width <= 0 or height <= 0:
        return 0

    step = max(min(width, height) // 120, 1)
    seeds = _build_edge_seed_points(width, height, step)
    visited = set()
    queue = deque()

    for x, y in seeds:
        if (x, y) in visited:
            continue
        visited.add((x, y))
        red, green, blue, alpha = image.getpixel((x, y))
        if alpha >= 240 and _calculate_rgb_distance((red, green, blue), background_rgb) <= tolerance:
            queue.append((x, y))

    changed_pixels = 0
    while queue:
        x, y = queue.popleft()
        red, green, blue, alpha = image.getpixel((x, y))
        if alpha != 0:
            image.putpixel((x, y), (red, green, blue, 0))
            changed_pixels += 1

        for nx, ny in ((x - 1, y), (x + 1, y), (x, y - 1), (x, y + 1)):
            if nx < 0 or ny < 0 or nx >= width or ny >= height or (nx, ny) in visited:
                continue
            visited.add((nx, ny))
            nr, ng, nb, na = image.getpixel((nx, ny))
            if na >= 240 and _calculate_rgb_distance((nr, ng, nb), background_rgb) <= tolerance:
                queue.append((nx, ny))

    return changed_pixels


def _try_make_logo_background_transparent(binary, normalized_type, image_kind):
    if image_kind != 'logo' or not WEBSITE_LOGO_AUTO_TRANSPARENT:
        return binary, normalized_type, ''

    if normalized_type not in {
        'image/avif',
        'image/bmp',
        'image/gif',
        'image/heic',
        'image/heif',
        'image/jpeg',
        'image/png',
        'image/tiff',
        'image/webp'
    }:
        return binary, normalized_type, ''

    try:
        from PIL import Image
    except Exception as exc:
        logger.warning('pillow is not available, skip logo transparent processing: %s', exc)
        return binary, normalized_type, ''

    try:
        with Image.open(io.BytesIO(binary)) as raw_image:
            image = raw_image.convert('RGBA')
    except Exception as exc:
        logger.warning('logo transparent processing skipped, invalid image: %s', exc)
        return binary, normalized_type, ''

    background_rgb = _estimate_background_rgb(image)
    tolerance = max(min(WEBSITE_LOGO_TRANSPARENT_TOLERANCE, 120), 0)

    pixels = list(image.getdata())
    if not pixels:
        return binary, normalized_type, ''

    candidate_backgrounds = [(255, 255, 255), (250, 250, 250)]
    if _is_background_dominant_on_border(image, background_rgb, tolerance * 3):
        candidate_backgrounds.insert(0, background_rgb)

    unique_candidates = []
    seen_candidates = set()
    for color in candidate_backgrounds:
        if color in seen_candidates:
            continue
        seen_candidates.add(color)
        unique_candidates.append(color)

    best_image = image
    best_changed_pixels = 0
    best_background = None
    for bg_color in unique_candidates:
        candidate_image = image.copy()
        changed_pixels = _transparentize_connected_background(candidate_image, bg_color, tolerance * 3)
        if changed_pixels > best_changed_pixels:
            best_changed_pixels = changed_pixels
            best_image = candidate_image
            best_background = bg_color

    changed_ratio = best_changed_pixels / len(pixels)
    if changed_ratio < WEBSITE_LOGO_TRANSPARENT_MIN_RATIO:
        return binary, normalized_type, ''

    image = best_image
    output = io.BytesIO()
    image.save(output, format='PNG', optimize=True)
    processed_binary = output.getvalue()
    if not processed_binary:
        return binary, normalized_type, ''

    if len(processed_binary) > WEBSITE_IMAGE_MAX_BYTES:
        logger.warning('transparent logo exceeds max size after processing, keep original')
        return binary, normalized_type, ''

    logger.info(
        'logo transparent processing applied, bg=%s, ratio=%.4f, bytes_before=%s, bytes_after=%s',
        best_background,
        changed_ratio,
        len(binary),
        len(processed_binary)
    )
    return processed_binary, 'image/png', 'png'


def _guess_image_content_type(content_type, filename=''):
    normalized_type = (content_type or '').lower()
    if normalized_type in WEBSITE_IMAGE_ALLOWED_TYPES:
        return normalized_type

    suffix = Path(filename or '').suffix.lower().lstrip('.')
    if suffix == 'jpeg':
        suffix = 'jpg'

    return WEBSITE_IMAGE_MIME_MAP.get(suffix, '')


def _upload_binary_to_minio(binary, content_type, image_kind, filename=''):
    if image_kind not in WEBSITE_IMAGE_KIND_SET:
        return '', '图片类型不支持'

    if not _is_minio_enabled():
        return '', '对象存储未启用，无法上传网站图片'

    normalized_type = _guess_image_content_type(content_type, filename)
    if normalized_type not in WEBSITE_IMAGE_ALLOWED_TYPES:
        return '', '网站图片格式不正确，仅支持常见图片格式'

    if len(binary) > WEBSITE_IMAGE_MAX_BYTES:
        max_mb = WEBSITE_IMAGE_MAX_BYTES // (1024 * 1024)
        return '', f'网站图片不能超过{max_mb}MB'

    processed_binary, processed_type, processed_extension = _try_make_logo_background_transparent(
        binary,
        normalized_type,
        image_kind
    )
    binary = processed_binary
    normalized_type = processed_type

    extension = processed_extension or _guess_image_extension(normalized_type, filename)
    if not extension:
        return '', '网站图片格式不正确，仅支持常见图片格式'

    try:
        client = _get_minio_client()
        _ensure_minio_bucket(client)

        object_name = f'{WEBSITE_ASSET_OBJECT_PREFIX}/{image_kind}/{uuid.uuid4().hex}.{extension}'
        client.put_object(
            bucket_name=getattr(settings, 'MINIO_BUCKET', ''),
            object_name=object_name,
            data=io.BytesIO(binary),
            length=len(binary),
            content_type=normalized_type
        )
        return _minio_object_url(object_name), ''
    except (S3Error, ValueError, Exception) as exc:
        logger.exception('website image upload to minio failed, kind=%s, error=%s', image_kind, exc)
        return '', '网站图片上传失败，请稍后重试'


def _is_admin_user(user):
    return bool(user and (user.is_staff or user.is_superuser))


def _website_settings_to_dict(settings_obj):
    return {
        'favicon': settings_obj.favicon,
        'logo': settings_obj.logo,
        'watermark': _watermark_settings_to_dict(settings_obj),
        'websiteName': settings_obj.website_name
    }


def _website_brand_settings_to_dict(settings_obj):
    return {
        'favicon': settings_obj.favicon,
        'logo': settings_obj.logo,
        'websiteName': settings_obj.website_name
    }


def _watermark_settings_to_dict(settings_obj):
    hidden_mode = bool(getattr(settings_obj, 'watermark_hidden_mode', False))
    if not _supports_watermark_hidden_mode():
        hidden_mode = False

    return {
        'enabled': settings_obj.watermark_enabled,
        'fontSize': settings_obj.watermark_font_size,
        'hiddenMode': hidden_mode,
        'opacity': settings_obj.watermark_opacity,
        'rotate': settings_obj.watermark_rotate,
        'text': settings_obj.watermark_text
    }


def _get_website_settings():
    defaults = {
        'favicon': '',
        'logo': '',
        'watermark_enabled': False,
        'watermark_font_size': 16,
        'watermark_opacity': 0.15,
        'watermark_rotate': -15,
        'watermark_text': 'zznode',
        'website_name': '智能审计系统'
    }
    if _supports_watermark_hidden_mode():
        defaults['watermark_hidden_mode'] = False

    settings_obj, _ = WebsiteSetting.objects.get_or_create(id=1, defaults=defaults)
    return settings_obj


def _get_website_settings_for_update():
    settings_obj = WebsiteSetting.objects.select_for_update().filter(id=1).first()
    if settings_obj:
        return settings_obj

    settings_obj = _get_website_settings()
    return WebsiteSetting.objects.select_for_update().get(id=settings_obj.id)


def _extract_minio_object_name(url):
    if not url:
        return ''

    parsed = urlparse(url)
    if parsed.scheme not in {'http', 'https'}:
        return ''

    allowed_netlocs = set()
    endpoint = (getattr(settings, 'MINIO_ENDPOINT', '') or '').strip()
    if endpoint:
        allowed_netlocs.add(endpoint)

    public_base_url = (getattr(settings, 'MINIO_PUBLIC_BASE_URL', '') or '').strip()
    if public_base_url:
        public_netloc = urlparse(public_base_url).netloc
        if public_netloc:
            allowed_netlocs.add(public_netloc)

    if allowed_netlocs and parsed.netloc not in allowed_netlocs:
        return ''

    bucket = (getattr(settings, 'MINIO_BUCKET', '') or '').strip()
    if not bucket:
        return ''

    expected_prefix = f'{bucket}/'
    object_path = (parsed.path or '').lstrip('/')
    if not object_path.startswith(expected_prefix):
        return ''

    object_name = object_path[len(expected_prefix):].strip()
    if not object_name.startswith(f'{WEBSITE_ASSET_OBJECT_PREFIX}/'):
        return ''

    return object_name


def _delete_minio_object(object_name):
    if not object_name or not _is_minio_enabled():
        return

    try:
        client = _get_minio_client()
        client.remove_object(getattr(settings, 'MINIO_BUCKET', ''), object_name)
    except (S3Error, ValueError, Exception) as exc:
        logger.warning('minio object cleanup failed, object=%s, error=%s', object_name, exc)


def _cleanup_stale_website_assets(previous_logo, previous_favicon, next_logo, next_favicon):
    stale_urls = []
    next_refs = {next_logo, next_favicon}

    if previous_logo and previous_logo not in next_refs:
        stale_urls.append(previous_logo)
    if previous_favicon and previous_favicon not in next_refs:
        stale_urls.append(previous_favicon)

    for stale_url in stale_urls:
        object_name = _extract_minio_object_name(stale_url)
        if object_name:
            _delete_minio_object(object_name)


def _is_valid_image_reference(value):
    if not value:
        return True

    if len(value) > WEBSITE_IMAGE_MAX_LENGTH:
        return False

    parsed = urlparse(value)
    return parsed.scheme in {'http', 'https'} and bool(parsed.netloc)


def _validate_website_brand_payload(data):
    website_name = (data.get('websiteName') or '').strip()
    logo = (data.get('logo') or '').strip()
    favicon = (data.get('favicon') or '').strip()

    if not website_name:
        return None, '网站名称不能为空'

    if len(website_name) > WEBSITE_NAME_MAX_LENGTH:
        return None, f'网站名称长度不能超过{WEBSITE_NAME_MAX_LENGTH}个字符'

    if any(ord(char) < 32 for char in website_name):
        return None, '网站名称包含非法字符'

    if not _is_valid_image_reference(logo):
        return None, '网站 Logo 格式不正确'

    if not _is_valid_image_reference(favicon):
        return None, '网站图标格式不正确'

    return {
        'favicon': favicon,
        'logo': logo,
        'website_name': website_name
    }, ''


def _validate_watermark_payload(watermark):
    if watermark is None:
        watermark = {}
    if not isinstance(watermark, dict):
        return None, '水印配置格式不正确'

    watermark_enabled = watermark.get('enabled', False)
    watermark_hidden_mode = watermark.get('hiddenMode', False)
    watermark_text = (watermark.get('text') or '').strip()
    watermark_font_size = watermark.get('fontSize', 16)
    watermark_opacity = watermark.get('opacity', 0.15)
    watermark_rotate = watermark.get('rotate', -15)

    if not isinstance(watermark_enabled, bool):
        return None, '水印开关格式不正确'

    if not isinstance(watermark_hidden_mode, bool):
        return None, '隐藏水印模式格式不正确'

    if not watermark_text:
        return None, '水印文字不能为空'
    if len(watermark_text) > WATERMARK_TEXT_MAX_LENGTH:
        return None, f'水印文字长度不能超过{WATERMARK_TEXT_MAX_LENGTH}个字符'
    if any(ord(char) < 32 for char in watermark_text):
        return None, '水印文字包含非法字符'

    try:
        watermark_font_size = int(watermark_font_size)
    except (TypeError, ValueError):
        return None, '水印字体大小格式不正确'
    if watermark_font_size < WATERMARK_FONT_SIZE_MIN or watermark_font_size > WATERMARK_FONT_SIZE_MAX:
        return None, f'水印字体大小必须在{WATERMARK_FONT_SIZE_MIN}-{WATERMARK_FONT_SIZE_MAX}之间'

    try:
        watermark_opacity = float(watermark_opacity)
    except (TypeError, ValueError):
        return None, '水印透明度格式不正确'
    if watermark_opacity < WATERMARK_OPACITY_MIN or watermark_opacity > WATERMARK_OPACITY_MAX:
        return None, f'水印透明度必须在{WATERMARK_OPACITY_MIN}-{WATERMARK_OPACITY_MAX}之间'

    try:
        watermark_rotate = int(watermark_rotate)
    except (TypeError, ValueError):
        return None, '水印旋转角度格式不正确'
    if watermark_rotate < WATERMARK_ROTATE_MIN or watermark_rotate > WATERMARK_ROTATE_MAX:
        return None, f'水印旋转角度必须在{WATERMARK_ROTATE_MIN}-{WATERMARK_ROTATE_MAX}之间'

    return {
        'watermark_enabled': watermark_enabled,
        'watermark_font_size': watermark_font_size,
        'watermark_hidden_mode': watermark_hidden_mode if _supports_watermark_hidden_mode() else False,
        'watermark_opacity': watermark_opacity,
        'watermark_rotate': watermark_rotate,
        'watermark_text': watermark_text
    }, ''


def _validate_website_settings_payload(data):
    brand_payload, brand_err_msg = _validate_website_brand_payload(data)
    if not brand_payload:
        return None, brand_err_msg

    watermark_payload, watermark_err_msg = _validate_watermark_payload(data.get('watermark'))
    if not watermark_payload:
        return None, watermark_err_msg

    payload = {}
    payload.update(brand_payload)
    payload.update(watermark_payload)
    return payload, ''


@swagger_auto_schema(
    method='post',
    operation_description='上传网站图片（Logo/Favicon）',
    manual_parameters=[
        openapi.Parameter('kind', openapi.IN_FORM, description='图片类型: logo 或 favicon', type=openapi.TYPE_STRING, required=True),
        openapi.Parameter('file', openapi.IN_FORM, description='图片文件', type=openapi.TYPE_FILE, required=True),
    ],
    responses={200: '上传成功', 400: '参数错误', 403: '权限不足'}
)
@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser])
def upload_website_image(request):
    if not _is_admin_user(request.user):
        return error_response('无权限上传网站图片', ERROR_CODE_FORBIDDEN, status.HTTP_403_FORBIDDEN)

    image_kind = (request.data.get('kind') or '').strip().lower()
    if image_kind not in WEBSITE_IMAGE_KIND_SET:
        return error_response('图片类型不支持', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    upload_file = request.FILES.get('file')
    if not upload_file:
        return error_response('上传文件不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    binary = upload_file.read()
    if not binary:
        return error_response('上传文件不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    image_url, err_msg = _upload_binary_to_minio(binary, upload_file.content_type, image_kind, upload_file.name)
    if err_msg:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    return success_response({'url': image_url}, '上传成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取图形验证码',
    responses={
        200: openapi.Response(
            '获取成功',
            openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'code': openapi.Schema(type=openapi.TYPE_STRING),
                    'data': openapi.Schema(
                        type=openapi.TYPE_OBJECT,
                        properties={
                            'captchaId': openapi.Schema(type=openapi.TYPE_STRING),
                            'captchaImage': openapi.Schema(type=openapi.TYPE_STRING)
                        }
                    ),
                    'msg': openapi.Schema(type=openapi.TYPE_STRING)
                }
            )
        )
    }
)
@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def get_captcha(_request):
    code = _generate_captcha_code()
    captcha_id = _create_captcha_id(code)

    return success_response(
        {
            'captchaId': captcha_id,
            'captchaImage': _build_captcha_svg(code)
        },
        '获取验证码成功'
    )


@swagger_auto_schema(
    method='get',
    operation_description='健康检查接口',
    responses={
        200: openapi.Response(
            '成功',
            openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'code': openapi.Schema(type=openapi.TYPE_STRING),
                    'data': openapi.Schema(
                        type=openapi.TYPE_OBJECT,
                        properties={'status': openapi.Schema(type=openapi.TYPE_STRING)}
                    ),
                    'msg': openapi.Schema(type=openapi.TYPE_STRING)
                }
            )
        )
    }
)
@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def health_check(_request):
    return success_response({'status': 'ok'})


@swagger_auto_schema(
    method='get',
    operation_description='获取项目状态汇总（参考 ai-audit-agent 口径）',
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_project_status_summary(_request):
    try:
        rows = _build_project_status_rows()
        return success_response(_build_project_summary_payload(rows), '获取成功')
    except Exception as exc:
        logger.exception('get project status summary failed: %s', exc)
        return error_response('项目状态汇总查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='get',
    operation_description='获取单项目状态明细（参考 ai-audit-agent 口径）',
    responses={200: '获取成功', 404: '项目不存在', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_project_status_detail(_request, project_id):
    normalized_project_id = _to_str(project_id)
    if not normalized_project_id:
        return error_response('projectId 不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        rows = _build_project_status_rows()
    except Exception as exc:
        logger.exception('get project status detail failed: %s', exc)
        return error_response('项目状态查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    matched = next((item for item in rows if item.get('projectId') == normalized_project_id), None)
    if matched is None:
        return error_response('项目不存在', ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)

    return success_response(matched, '获取成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取项目状态排名（用于数据概览 TOP20）',
    manual_parameters=[
        openapi.Parameter(
            'category',
            openapi.IN_QUERY,
            description='分类：completed/running/remaining/abnormal',
            type=openapi.TYPE_STRING
        ),
        openapi.Parameter('startTime', openapi.IN_QUERY, description='开始时间，支持 YYYY-MM-DD 或 YYYY-MM-DD HH:mm:ss', type=openapi.TYPE_STRING),
        openapi.Parameter('endTime', openapi.IN_QUERY, description='结束时间，支持 YYYY-MM-DD 或 YYYY-MM-DD HH:mm:ss', type=openapi.TYPE_STRING),
        openapi.Parameter('limit', openapi.IN_QUERY, description='返回条数，默认20，最大100', type=openapi.TYPE_INTEGER),
    ],
    responses={200: '获取成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_project_status_ranking(request):
    category = _normalize_ranking_category(request.query_params.get('category'))
    if category is None:
        return error_response(
            'category 参数无效，仅支持 completed/running/remaining/abnormal',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_400_BAD_REQUEST
        )

    raw_start_time = request.query_params.get('startTime')
    raw_end_time = request.query_params.get('endTime')
    start_time = _parse_datetime_value(raw_start_time)
    end_time = _parse_datetime_value(raw_end_time)
    if raw_start_time and start_time is None:
        return error_response('startTime 格式无效', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    if raw_end_time and end_time is None:
        return error_response('endTime 格式无效', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    # Date-only endTime should cover the full day.
    if raw_end_time and end_time and len(str(raw_end_time).strip()) == 10:
        end_time = end_time + timedelta(days=1) - timedelta(seconds=1)
    if start_time and end_time and start_time > end_time:
        return error_response('startTime 不能晚于 endTime', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    limit = _parse_positive_int(
        request.query_params.get('limit'),
        default=PROJECT_STATUS_RANKING_DEFAULT_LIMIT,
        max_value=PROJECT_STATUS_RANKING_MAX_LIMIT
    )

    try:
        rows = _build_project_status_rows()
    except Exception as exc:
        logger.exception('get project status ranking failed: %s', exc)
        return error_response('项目状态排名查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    category_rows = [row for row in rows if _status_match_ranking_category(category, _to_str(row.get('status')).upper())]
    category_rows = _filter_project_status_rows_by_time(category_rows, start_time, end_time)

    if category == PROJECT_STATUS_CATEGORY_REMAINING:
        category_rows.sort(key=lambda item: (_to_str(item.get('projectName')) or _to_str(item.get('projectId'))))
    else:
        category_rows.sort(
            key=lambda item: (
                _parse_datetime_value(item.get('latestTaskUpdatedAt')) or datetime.min,
                _to_str(item.get('projectId'))
            ),
            reverse=True
        )

    records = []
    for idx, item in enumerate(category_rows[:limit], start=1):
        records.append({
            'rank': idx,
            'projectId': item.get('projectId'),
            'name': _to_str(item.get('projectName')) or _to_str(item.get('projectId')),
            'completeTime': item.get('latestTaskUpdatedAt'),
            'status': item.get('status'),
        })

    return success_response(
        {
            'category': category,
            'total': len(category_rows),
            'records': records,
        },
        '获取成功'
    )


@swagger_auto_schema(
    method='get',
    operation_description='获取数据查询行业选项',
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_query_industries(_request):
    try:
        options = _build_data_query_industry_options()
        return success_response(options, '获取成功')
    except Exception as exc:
        logger.exception('get data query industries failed: %s', exc)
        return error_response(
            '项目行业查询失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@swagger_auto_schema(
    method='get',
    operation_description='获取行业维度项目字段透视列表',
    manual_parameters=[
        openapi.Parameter('industry', openapi.IN_QUERY, description='项目行业（必填，支持 SW/JS）', type=openapi.TYPE_STRING),
        openapi.Parameter('projectName', openapi.IN_QUERY, description='项目名称关键字（可选）', type=openapi.TYPE_STRING),
        openapi.Parameter('current', openapi.IN_QUERY, description='页码，默认1', type=openapi.TYPE_INTEGER),
        openapi.Parameter('size', openapi.IN_QUERY, description='每页条数，默认10，最大50', type=openapi.TYPE_INTEGER),
    ],
    responses={200: '获取成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_query_industry_pivot(request):
    filters, err_msg = _parse_data_query_industry_pivot_filters(request)
    if err_msg:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        payload = _build_data_query_industry_pivot_payload(filters)
        return success_response(payload, '获取成功')
    except Exception as exc:
        logger.exception('get data query industry pivot failed: %s', exc)
        return error_response(
            '数据查询列表加载失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@swagger_auto_schema(
    method='get',
    operation_description='导出行业维度项目字段透视列表（xlsx）',
    manual_parameters=[
        openapi.Parameter('industry', openapi.IN_QUERY, description='项目行业（必填，支持 SW/JS）', type=openapi.TYPE_STRING),
        openapi.Parameter('projectName', openapi.IN_QUERY, description='项目名称关键字（可选）', type=openapi.TYPE_STRING),
    ],
    responses={200: '导出成功', 400: '参数错误', 500: '导出失败', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_data_query_industry_pivot(request):
    filters, err_msg = _parse_data_query_industry_pivot_filters(request)
    if err_msg:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        binary = _build_data_query_export_binary(filters)
    except Exception as exc:
        logger.exception('build data query export failed: %s', exc)
        if isinstance(exc, (ImportError, ModuleNotFoundError)):
            return error_response(
                '导出失败：Excel 导出依赖缺失，请检查 openpyxl 安装',
                ERROR_CODE_INVALID_PARAMS,
                status.HTTP_500_INTERNAL_SERVER_ERROR
            )
        return error_response('导出失败，请稍后重试', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    industry = _to_str(filters.get('industry')) or 'unknown'
    project_name = _to_str(filters.get('projectName'))
    safe_project_name = re.sub(r'[\\/:*?"<>|]+', '_', project_name).strip('_')
    file_suffix = datetime.now().strftime('%Y%m%d%H%M%S')
    if safe_project_name:
        filename = quote(f'data-query-{industry}-{safe_project_name}-{file_suffix}.xlsx')
    else:
        filename = quote(f'data-query-{industry}-{file_suffix}.xlsx')
    response = HttpResponse(
        binary,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
    return response


@swagger_auto_schema(
    method='get',
    operation_description='获取数据查询字段下钻明细（只读）',
    manual_parameters=[
        openapi.Parameter('projectId', openapi.IN_QUERY, description='项目ID（必填）', type=openapi.TYPE_STRING),
        openapi.Parameter('industry', openapi.IN_QUERY, description='项目行业（必填，支持 SW/JS）', type=openapi.TYPE_STRING),
        openapi.Parameter(
            'fieldKey',
            openapi.IN_QUERY,
            description='下钻字段 key（必填，支持 tender_agent_info/expert_info/opening_attendee_info）',
            type=openapi.TYPE_STRING
        ),
        openapi.Parameter('current', openapi.IN_QUERY, description='页码，默认1', type=openapi.TYPE_INTEGER),
        openapi.Parameter('size', openapi.IN_QUERY, description='每页条数，默认10，最大200', type=openapi.TYPE_INTEGER),
    ],
    responses={200: '获取成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_query_drilldown(request):
    filters, err_msg = _parse_data_query_drilldown_filters(request)
    if err_msg:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        payload = _build_data_query_drilldown_payload(filters)
        return success_response(payload, '获取成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('get data query drilldown failed: %s', exc)
        return error_response(
            '数据查询下钻加载失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度创建人下拉选项（来源 auth_user）',
    manual_parameters=[
        openapi.Parameter('keyword', openapi.IN_QUERY, description='用户名关键字（可选）', type=openapi.TYPE_STRING),
    ],
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_creators(request):
    keyword = _to_str(request.query_params.get('keyword'))
    options = _build_data_schedule_creator_options(keyword)
    return success_response(options, '获取成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度选择数据弹窗树结构',
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_select_data(_request):
    try:
        payload = _build_data_schedule_select_data_payload()
        return success_response(payload, '获取成功')
    except Exception as exc:
        logger.exception('get data schedule select data failed: %s', exc)
        return error_response(
            '选择数据目录查询失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@swagger_auto_schema(
    method='post',
    operation_description='执行数据调度任务动作（重新执行/继续执行/暂停/停止/删除/刷新）',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['action'],
        properties={
            'action': openapi.Schema(
                type=openapi.TYPE_STRING,
                description='动作标识：reExecute/continue/pause/stop/delete/refresh'
            ),
            'taskIds': openapi.Schema(
                type=openapi.TYPE_ARRAY,
                items=openapi.Items(type=openapi.TYPE_STRING),
                description='任务ID列表（action=refresh 时可为空）'
            ),
        }
    ),
    responses={200: '执行成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def operate_data_schedule_tasks(request):
    action = _normalize_data_schedule_task_action(request.data.get('action'))
    if action not in DATA_SCHEDULE_TASK_ACTION_SET:
        return error_response(
            'action 参数无效，仅支持 reExecute/continue/pause/stop/delete/refresh',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_400_BAD_REQUEST
        )

    task_ids = _normalize_data_schedule_task_ids(request.data.get('taskIds'))
    if action != DATA_SCHEDULE_TASK_ACTION_REFRESH and not task_ids:
        return error_response('taskIds 不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    if len(task_ids) > DATA_SCHEDULE_TASK_ACTION_MAX_SIZE:
        return error_response(
            f'taskIds 数量不能超过 {DATA_SCHEDULE_TASK_ACTION_MAX_SIZE}',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_400_BAD_REQUEST
        )

    try:
        task_rows = _fetch_data_schedule_action_task_rows(task_ids)
        valid_task_id_set = {_to_str(item.get('taskId')) for item in task_rows if _to_str(item.get('taskId'))}
        invalid_task_ids = [task_id for task_id in task_ids if task_id not in valid_task_id_set]

        operate_result = _operate_data_schedule_tasks(action, task_rows)

        payload = {
            'action': action,
            'deleted': _to_int(operate_result.get('deleted')),
            'fileReset': _to_int(operate_result.get('fileReset')),
            'invalidTaskIds': invalid_task_ids,
            'requested': len(task_ids),
            'updated': _to_int(operate_result.get('updated')),
        }
        if action == DATA_SCHEDULE_TASK_ACTION_REFRESH:
            payload.update(_build_data_schedule_task_refresh_payload())

        return success_response(payload, '执行成功')
    except Exception as exc:
        logger.exception('operate data schedule tasks failed: %s', exc)
        return error_response(
            '数据调度任务动作执行失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度任务列表',
    manual_parameters=[
        openapi.Parameter('current', openapi.IN_QUERY, description='页码，默认1', type=openapi.TYPE_INTEGER),
        openapi.Parameter('size', openapi.IN_QUERY, description=f'每页条数，默认{DATA_SCHEDULE_TASK_DEFAULT_SIZE}，最大{DATA_SCHEDULE_TASK_MAX_SIZE}', type=openapi.TYPE_INTEGER),
        openapi.Parameter('projectName', openapi.IN_QUERY, description='项目名称（模糊匹配）', type=openapi.TYPE_STRING),
        openapi.Parameter('taskStatus', openapi.IN_QUERY, description='任务状态（可重复传参）', type=openapi.TYPE_STRING),
        openapi.Parameter('creator', openapi.IN_QUERY, description='创建人（可重复传参）', type=openapi.TYPE_STRING),
        openapi.Parameter('matchRangeMin', openapi.IN_QUERY, description='匹配度下限（0-100）', type=openapi.TYPE_INTEGER),
        openapi.Parameter('matchRangeMax', openapi.IN_QUERY, description='匹配度上限（0-100）', type=openapi.TYPE_INTEGER),
        openapi.Parameter('createStartTime', openapi.IN_QUERY, description='创建时间起', type=openapi.TYPE_STRING),
        openapi.Parameter('createEndTime', openapi.IN_QUERY, description='创建时间止', type=openapi.TYPE_STRING),
    ],
    responses={200: '获取成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_tasks(request):
    filters, err_msg = _parse_data_schedule_task_filters(request)
    if err_msg:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        rows = _build_data_schedule_task_rows()
    except Exception as exc:
        logger.exception('get data schedule tasks failed: %s', exc)
        return error_response(
            '数据调度任务列表查询失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )

    filtered_rows = _filter_data_schedule_task_rows(rows, filters)
    current = filters.get('current') or 1
    size = filters.get('size') or DATA_SCHEDULE_TASK_DEFAULT_SIZE
    total = len(filtered_rows)
    start_idx = (current - 1) * size
    end_idx = start_idx + size
    paged_records = filtered_rows[start_idx:end_idx]

    return success_response(
        {
            'current': current,
            'size': size,
            'total': total,
            'records': paged_records,
        },
        '获取成功'
    )


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度任务详情',
    responses={200: '获取成功', 400: '参数错误', 404: '任务不存在', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_task_detail(_request, task_id):
    try:
        detail = _get_data_schedule_detail(task_id)
        return success_response(detail, '获取成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('get data schedule task detail failed: %s', exc)
        return error_response('任务详情查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度任务提取结果概览',
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_extract_summary(_request, task_id):
    try:
        bundle = _build_data_schedule_detail_bundle(task_id)
        all_fields = bundle.get('fields') or []
        summary = bundle.get('summary') or {}
        summary['counts'] = _build_data_schedule_counts(all_fields)
        return success_response(summary, '获取成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('get data schedule extract summary failed: %s', exc)
        return error_response('任务概览查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度任务提取字段结果列表',
    manual_parameters=[
        openapi.Parameter('scope', openapi.IN_QUERY, description='筛选范围: all/complete/missing', type=openapi.TYPE_STRING),
        openapi.Parameter('current', openapi.IN_QUERY, description='页码，默认1', type=openapi.TYPE_INTEGER),
        openapi.Parameter('size', openapi.IN_QUERY, description='每页条数，默认20，最大200', type=openapi.TYPE_INTEGER),
    ],
    responses={200: '获取成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_extract_fields(request, task_id):
    scope = _normalize_data_schedule_scope(request.query_params.get('scope'))
    if scope is None:
        return error_response('scope 参数无效，仅支持 all/complete/missing', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    current = _parse_positive_int(request.query_params.get('current'), default=1, max_value=100000)
    size = _parse_positive_int(request.query_params.get('size'), default=20, max_value=200)

    try:
        all_fields = (_build_data_schedule_detail_bundle(task_id).get('fields') or [])
        filtered_fields = _filter_data_schedule_fields_by_scope(all_fields, scope)
        field_records = [_to_data_schedule_field_record(item) for item in filtered_fields]
        paginated_data = _paginate_data_schedule_records(field_records, current, size)
        paginated_data['scope'] = scope
        paginated_data['counts'] = _build_data_schedule_counts(all_fields)
        return success_response(paginated_data, '获取成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('get data schedule extract fields failed: %s', exc)
        return error_response('字段列表查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='patch',
    operation_description='更新数据调度任务字段内容',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['fieldValue'],
        properties={
            'fieldValue': openapi.Schema(type=openapi.TYPE_STRING, description='字段内容'),
        }
    ),
    responses={200: '更新成功', 400: '参数错误', 404: '字段不存在', 401: '未认证'}
)
@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def update_data_schedule_extract_field(request, task_id, field_key):
    normalized_field_key = str(field_key or '').strip()
    if not normalized_field_key:
        return error_response('fieldKey 不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    all_fields = _get_data_schedule_fields(task_id)
    field_map = {item.get('fieldKey'): item for item in all_fields}
    field_item = field_map.get(normalized_field_key)
    if not field_item:
        return error_response('字段不存在', ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    if not field_item.get('canEdit'):
        return error_response('当前字段不支持编辑', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    raw_field_value = request.data.get('fieldValue')
    if raw_field_value is None:
        return error_response('fieldValue 不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    field_value = str(raw_field_value)
    if len(field_value) > DATA_SCHEDULE_FIELD_VALUE_MAX_LENGTH:
        return error_response(
            f'fieldValue 长度不能超过{DATA_SCHEDULE_FIELD_VALUE_MAX_LENGTH}个字符',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_400_BAD_REQUEST
        )

    try:
        _update_data_schedule_field_value_to_db(field_item, field_value)
    except ValueError as exc:
        error_msg = str(exc) or '更新失败'
        if error_msg == '更新目标不存在':
            return error_response(error_msg, ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
        return error_response(error_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('update data schedule extract field failed: %s', exc)
        return error_response('字段更新失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    field_item['fieldValueDisplay'] = field_value
    return success_response(_to_data_schedule_field_record(field_item), '更新成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取字段下钻详情数据',
    manual_parameters=[
        openapi.Parameter('fieldKey', openapi.IN_QUERY, description='字段标识', type=openapi.TYPE_STRING, required=True),
        openapi.Parameter('scope', openapi.IN_QUERY, description='筛选范围: all/complete/missing', type=openapi.TYPE_STRING),
        openapi.Parameter('current', openapi.IN_QUERY, description='页码，默认1', type=openapi.TYPE_INTEGER),
        openapi.Parameter('size', openapi.IN_QUERY, description='每页条数，默认20，最大200', type=openapi.TYPE_INTEGER),
    ],
    responses={200: '获取成功', 400: '参数错误', 404: '字段不存在', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_extract_drilldown(request, task_id):
    field_key = str(request.query_params.get('fieldKey') or '').strip()
    if not field_key:
        return error_response('fieldKey 不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    scope = _normalize_data_schedule_scope(request.query_params.get('scope'))
    if scope is None:
        return error_response('scope 参数无效，仅支持 all/complete/missing', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    current = _parse_positive_int(request.query_params.get('current'), default=1, max_value=100000)
    size = _parse_positive_int(request.query_params.get('size'), default=20, max_value=200)

    try:
        field_item, payload, err_msg, http_status, err_code = _get_data_schedule_drilldown_payload(task_id, field_key)
        if err_msg:
            return error_response(err_msg, err_code, http_status)

        records = payload.get('records') or []
        if scope != DATA_SCHEDULE_SCOPE_ALL and field_item.get('status') != scope:
            records = []

        paginated_data = _paginate_data_schedule_records(records, current, size)
        paginated_data.update(
            {
                'actions': {
                    'canCreate': bool((payload.get('actions') or {}).get('canCreate')),
                    'canDelete': bool((payload.get('actions') or {}).get('canDelete')),
                    'canEdit': bool((payload.get('actions') or {}).get('canEdit'))
                },
                'columns': payload.get('columns') or [],
                'fieldKey': field_key,
                'scope': scope,
                'title': payload.get('title') or f"字段详情 - {field_item.get('fieldName') or field_key}",
            }
        )

        return success_response(paginated_data, '获取成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('get data schedule extract drilldown failed: %s', exc)
        return error_response('字段下钻查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='post',
    operation_description='新增下钻行',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['fieldKey'],
        properties={
            'fieldKey': openapi.Schema(type=openapi.TYPE_STRING, description='字段标识'),
            'rowData': openapi.Schema(type=openapi.TYPE_OBJECT, description='行数据，key 为下钻列 key'),
        }
    ),
    responses={200: '新增成功', 400: '参数错误', 404: '字段不存在', 401: '未认证'}
)
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def create_data_schedule_extract_drilldown_row(request, task_id):
    try:
        field_key = str(request.data.get('fieldKey') or '').strip()
        field_item, payload, err_msg, http_status, err_code = _get_data_schedule_drilldown_payload(task_id, field_key)
        if err_msg:
            return error_response(err_msg, err_code, http_status)
        if not (payload.get('actions') or {}).get('canCreate'):
            return error_response('当前字段不支持新增', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        columns = [item for item in (payload.get('columns') or []) if item.get('editable')]
        normalized_row, normalize_err = _normalize_drilldown_row_data(request.data.get('rowData'), columns)
        if normalize_err:
            return error_response(normalize_err, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        _create_data_schedule_drilldown_row_to_db(payload, normalized_row)

        records = payload.get('records') or []
        row_id = _get_next_drilldown_row_id(records)
        created_row = {'id': row_id, **normalized_row}

        response_data = {
            'fieldKey': field_item.get('fieldKey') or field_key,
            'record': created_row
        }
        return success_response(response_data, '新增成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except (IntegrityError, DataError) as exc:
        logger.exception('create data schedule drilldown row failed: %s', exc)
        return error_response(
            _build_data_schedule_db_write_error_message(exc, action_name='新增'),
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_400_BAD_REQUEST
        )
    except OperationalError as exc:
        logger.exception('create data schedule drilldown row failed: %s', exc)
        return error_response(
            _build_data_schedule_db_write_error_message(exc, action_name='新增'),
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )
    except Exception as exc:
        logger.exception('create data schedule drilldown row failed: %s', exc)
        return error_response('新增失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='patch',
    operation_description='编辑下钻行',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['fieldKey'],
        properties={
            'fieldKey': openapi.Schema(type=openapi.TYPE_STRING, description='字段标识'),
            'rowData': openapi.Schema(type=openapi.TYPE_OBJECT, description='行数据，key 为下钻列 key'),
        }
    ),
    responses={200: '更新成功', 400: '参数错误', 404: '字段或行不存在', 401: '未认证'}
)
@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def update_data_schedule_extract_drilldown_row(request, task_id, row_id):
    try:
        field_key = str(request.data.get('fieldKey') or '').strip()
        field_item, payload, err_msg, http_status, err_code = _get_data_schedule_drilldown_payload(task_id, field_key)
        if err_msg:
            return error_response(err_msg, err_code, http_status)
        if not (payload.get('actions') or {}).get('canEdit'):
            return error_response('当前字段不支持编辑', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        records = payload.get('records') or []
        target_row = _find_drilldown_row(records, row_id)
        if not target_row:
            return error_response('下钻行不存在', ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)

        columns = [item for item in (payload.get('columns') or []) if item.get('editable')]
        normalized_row, normalize_err = _normalize_drilldown_row_data(request.data.get('rowData'), columns)
        if normalize_err:
            return error_response(normalize_err, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        _update_data_schedule_drilldown_row_to_db(payload, target_row, normalized_row)
        target_row.update(normalized_row)
        response_data = {
            'fieldKey': field_item.get('fieldKey') or field_key,
            'record': target_row
        }
        return success_response(response_data, '更新成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('update data schedule drilldown row failed: %s', exc)
        return error_response('更新失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='delete',
    operation_description='删除下钻行',
    manual_parameters=[
        openapi.Parameter('fieldKey', openapi.IN_QUERY, description='字段标识', type=openapi.TYPE_STRING, required=True),
    ],
    responses={200: '删除成功', 400: '参数错误', 404: '字段或行不存在', 401: '未认证'}
)
@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def delete_data_schedule_extract_drilldown_row(request, task_id, row_id):
    try:
        field_key = str(request.query_params.get('fieldKey') or '').strip()
        field_item, payload, err_msg, http_status, err_code = _get_data_schedule_drilldown_payload(task_id, field_key)
        if err_msg:
            return error_response(err_msg, err_code, http_status)
        if not (payload.get('actions') or {}).get('canDelete'):
            return error_response('当前字段不支持删除', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        records = payload.get('records') or []
        target_row = _find_drilldown_row(records, row_id)
        if not target_row:
            return error_response('下钻行不存在', ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)

        _delete_data_schedule_drilldown_row_from_db(payload, target_row)
        response_data = {
            'fieldKey': field_item.get('fieldKey') or field_key,
            'rowId': row_id
        }
        return success_response(response_data, '删除成功')
    except ValueError as exc:
        return error_response(str(exc), ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
    except RuntimeError as exc:
        return error_response(str(exc), ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    except Exception as exc:
        logger.exception('delete data schedule drilldown row failed: %s', exc)
        return error_response('删除失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)


@swagger_auto_schema(
    method='get',
    operation_description='导出数据调度提取结果（xlsx，支持下钻 Sheet 跳转）',
    manual_parameters=[
        openapi.Parameter('scope', openapi.IN_QUERY, description='筛选范围: all/complete/missing', type=openapi.TYPE_STRING),
    ],
    responses={200: '导出成功', 400: '参数错误', 500: '导出失败', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_data_schedule_extract_result(request, task_id):
    scope = _normalize_data_schedule_scope(request.query_params.get('scope'))
    if scope is None:
        return error_response('scope 参数无效，仅支持 all/complete/missing', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    query_snapshot = _get_request_query_snapshot(request)
    _append_data_schedule_export_log(
        'info',
        'extract_export_started',
        task_id=task_id,
        scope=scope,
        request_params=query_snapshot
    )

    try:
        binary = _build_data_schedule_export_binary(task_id, scope)
    except Exception as exc:
        logger.exception('build data schedule export failed, task_id=%s, scope=%s', task_id, scope)
        _append_data_schedule_export_log(
            'error',
            'extract_export_failed',
            task_id=task_id,
            scope=scope,
            request_params=query_snapshot,
            error=exc
        )
        if isinstance(exc, (ImportError, ModuleNotFoundError)):
            return error_response(
                '导出失败：Excel 导出依赖缺失，请检查 openpyxl 安装',
                ERROR_CODE_INVALID_PARAMS,
                status.HTTP_500_INTERNAL_SERVER_ERROR
            )
        return error_response('导出失败，请稍后重试', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    _append_data_schedule_export_log(
        'info',
        'extract_export_succeeded',
        task_id=task_id,
        scope=scope,
        request_params=query_snapshot,
        extra={'bytes': len(binary)}
    )

    filename = quote(f'data-schedule-{task_id}-{scope}.xlsx')
    response = HttpResponse(
        binary,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
    return response


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度日志元信息（支持 projectId 优先查询）',
    manual_parameters=[
        openapi.Parameter('projectId', openapi.IN_QUERY, description='项目ID（可选，优先使用）', type=openapi.TYPE_STRING),
    ],
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_logs_meta(request, task_id):
    project_id = _to_str(request.query_params.get('projectId'))
    try:
        return success_response(_get_data_schedule_log_meta(task_id, project_id), '获取成功')
    except Exception as exc:
        logger.exception('get data schedule logs meta failed: %s', exc)
        return error_response(
            '数据提取日志元信息加载失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度日志列表（按条数查询）',
    manual_parameters=[
        openapi.Parameter('projectId', openapi.IN_QUERY, description='项目ID（可选，优先使用）', type=openapi.TYPE_STRING),
        openapi.Parameter('level', openapi.IN_QUERY, description='日志级别: info/warn/error', type=openapi.TYPE_STRING),
        openapi.Parameter('keyword', openapi.IN_QUERY, description='日志关键字', type=openapi.TYPE_STRING),
        openapi.Parameter('startTime', openapi.IN_QUERY, description='开始时间', type=openapi.TYPE_STRING),
        openapi.Parameter('endTime', openapi.IN_QUERY, description='结束时间', type=openapi.TYPE_STRING),
        openapi.Parameter('orderBy', openapi.IN_QUERY, description='排序字段，固定 time', type=openapi.TYPE_STRING),
        openapi.Parameter('orderDirection', openapi.IN_QUERY, description='排序方向: asc/desc', type=openapi.TYPE_STRING),
        openapi.Parameter(
            'count',
            openapi.IN_QUERY,
            description=f'查询条数，默认{DATA_SCHEDULE_LOG_DEFAULT_COUNT}，最大{DATA_SCHEDULE_LOG_MAX_COUNT}',
            type=openapi.TYPE_INTEGER
        ),
    ],
    responses={200: '获取成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_logs(request, task_id):
    filters, err_msg = _parse_data_schedule_log_filters(request)
    if err_msg:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    project_id = _to_str(request.query_params.get('projectId'))
    try:
        records = _build_data_schedule_log_records(task_id, project_id)
        filtered_records = _filter_data_schedule_logs(records, filters)
        sorted_records = _sort_data_schedule_logs(filtered_records, filters.get('orderDirection'))
        count = filters.get('count') or DATA_SCHEDULE_LOG_DEFAULT_COUNT
        sliced_records = sorted_records[:count]

        data = {
            'count': count,
            'records': _serialize_data_schedule_logs(sliced_records),
            'total': len(sorted_records),
        }
        return success_response(data, '获取成功')
    except Exception as exc:
        logger.exception('get data schedule logs failed: %s', exc)
        return error_response(
            '数据提取日志加载失败，请检查数据库连接',
            ERROR_CODE_INVALID_PARAMS,
            status.HTTP_500_INTERNAL_SERVER_ERROR
        )


@swagger_auto_schema(
    method='get',
    operation_description='导出数据调度日志',
    manual_parameters=[
        openapi.Parameter('projectId', openapi.IN_QUERY, description='项目ID（可选，优先使用）', type=openapi.TYPE_STRING),
        openapi.Parameter('level', openapi.IN_QUERY, description='日志级别: info/warn/error', type=openapi.TYPE_STRING),
        openapi.Parameter('keyword', openapi.IN_QUERY, description='日志关键字', type=openapi.TYPE_STRING),
        openapi.Parameter('startTime', openapi.IN_QUERY, description='开始时间', type=openapi.TYPE_STRING),
        openapi.Parameter('endTime', openapi.IN_QUERY, description='结束时间', type=openapi.TYPE_STRING),
        openapi.Parameter('orderBy', openapi.IN_QUERY, description='排序字段，固定 time', type=openapi.TYPE_STRING),
        openapi.Parameter('orderDirection', openapi.IN_QUERY, description='排序方向: asc/desc', type=openapi.TYPE_STRING),
        openapi.Parameter(
            'count',
            openapi.IN_QUERY,
            description=f'导出条数，默认{DATA_SCHEDULE_LOG_DEFAULT_COUNT}，最大{DATA_SCHEDULE_LOG_MAX_COUNT}',
            type=openapi.TYPE_INTEGER
        ),
    ],
    responses={200: '导出成功', 400: '参数错误', 500: '导出失败', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def export_data_schedule_logs(request, task_id):
    filters, err_msg = _parse_data_schedule_log_filters(request)
    if err_msg:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)
    project_id = _to_str(request.query_params.get('projectId'))

    query_snapshot = _get_request_query_snapshot(request)
    _append_data_schedule_export_log(
        'info',
        'task_logs_export_started',
        task_id=task_id,
        scope='logs',
        request_params=query_snapshot
    )

    try:
        task_meta = _get_data_schedule_log_meta(task_id, project_id)
        records = _build_data_schedule_log_records(task_id, project_id)
        filtered_records = _filter_data_schedule_logs(records, filters)
        sorted_records = _sort_data_schedule_logs(filtered_records, filters.get('orderDirection'))
        count = filters.get('count') or DATA_SCHEDULE_LOG_DEFAULT_COUNT
        sliced_records = sorted_records[:count]
        binary = _build_data_schedule_log_export_binary(task_meta, filters, sliced_records)
    except Exception as exc:
        logger.exception('build data schedule logs export failed, task_id=%s, project_id=%s', task_id, project_id)
        _append_data_schedule_export_log(
            'error',
            'task_logs_export_failed',
            task_id=task_id,
            scope='logs',
            request_params=query_snapshot,
            error=exc
        )
        if isinstance(exc, (ImportError, ModuleNotFoundError)):
            return error_response(
                '导出失败：Excel 导出依赖缺失，请检查 openpyxl 安装',
                ERROR_CODE_INVALID_PARAMS,
                status.HTTP_500_INTERNAL_SERVER_ERROR
            )
        return error_response('导出失败，请稍后重试', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    _append_data_schedule_export_log(
        'info',
        'task_logs_export_succeeded',
        task_id=task_id,
        scope='logs',
        request_params=query_snapshot,
        extra={'bytes': len(binary)}
    )

    filename = quote(f'data-schedule-log-{task_id}.xlsx')
    response = HttpResponse(
        binary,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f"attachment; filename*=UTF-8''{filename}"
    return response


@swagger_auto_schema(
    method='get',
    operation_description='获取数据调度导出调试日志（按行读取）',
    manual_parameters=[
        openapi.Parameter('lines', openapi.IN_QUERY, description='返回日志行数，默认200，最大2000', type=openapi.TYPE_INTEGER),
        openapi.Parameter('keyword', openapi.IN_QUERY, description='关键字过滤（可选）', type=openapi.TYPE_STRING),
    ],
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_data_schedule_export_debug_logs(request):
    lines = _parse_positive_int(request.query_params.get('lines'), default=200, max_value=2000)
    keyword = str(request.query_params.get('keyword') or '').strip()
    records = _read_data_schedule_export_log_records(lines, keyword)
    data = {
        'file': str(DATA_SCHEDULE_EXPORT_LOG_FILE),
        'records': records,
        'total': len(records)
    }
    return success_response(data, '获取成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取后端 API 失败日志（接口调用失败追踪）',
    manual_parameters=[
        openapi.Parameter('lines', openapi.IN_QUERY, description='返回日志行数，默认200，最大2000', type=openapi.TYPE_INTEGER),
        openapi.Parameter('keyword', openapi.IN_QUERY, description='关键字过滤（可选）', type=openapi.TYPE_STRING),
    ],
    responses={200: '获取成功', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_api_failure_logs(request):
    lines = _parse_positive_int(request.query_params.get('lines'), default=200, max_value=2000)
    keyword = str(request.query_params.get('keyword') or '').strip()
    records = _read_api_failure_log_records(lines, keyword)
    data = {
        'file': str(API_FAILURE_LOG_FILE),
        'records': records,
        'total': len(records)
    }
    return success_response(data, '获取成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取网站设置',
    responses={200: '获取成功'}
)
@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def get_website_settings(_request):
    settings_obj = _get_website_settings()
    return success_response(_website_settings_to_dict(settings_obj), '获取成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取网站基础设置（名称/Logo/Favicon）',
    responses={200: '获取成功'}
)
@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def get_website_brand_settings(_request):
    settings_obj = _get_website_settings()
    return success_response(_website_brand_settings_to_dict(settings_obj), '获取成功')


@swagger_auto_schema(
    method='post',
    operation_description='更新网站基础设置（名称/Logo/Favicon）',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['websiteName'],
        properties={
            'websiteName': openapi.Schema(type=openapi.TYPE_STRING, description='网站名称'),
            'logo': openapi.Schema(type=openapi.TYPE_STRING, description='网站 Logo URL'),
            'favicon': openapi.Schema(type=openapi.TYPE_STRING, description='网站图标 URL')
        }
    ),
    responses={200: '更新成功', 400: '参数错误', 403: '权限不足'}
)
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def update_website_brand_settings(request):
    if not _is_admin_user(request.user):
        return error_response('无权限修改网站设置', ERROR_CODE_FORBIDDEN, status.HTTP_403_FORBIDDEN)

    payload, err_msg = _validate_website_brand_payload(request.data or {})
    if not payload:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    with transaction.atomic():
        settings_obj = _get_website_settings_for_update()
        previous_logo = settings_obj.logo
        previous_favicon = settings_obj.favicon

        is_changed = (
            settings_obj.website_name != payload['website_name']
            or settings_obj.logo != payload['logo']
            or settings_obj.favicon != payload['favicon']
        )
        if not is_changed:
            return success_response(True, '设置未发生变化')

        settings_obj.website_name = payload['website_name']
        settings_obj.logo = payload['logo']
        settings_obj.favicon = payload['favicon']
        settings_obj.save(update_fields=['website_name', 'logo', 'favicon', 'updated_at'])

    _cleanup_stale_website_assets(previous_logo, previous_favicon, payload['logo'], payload['favicon'])

    return success_response(True, '更新成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取水印设置',
    responses={200: '获取成功'}
)
@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def get_watermark_settings(_request):
    settings_obj = _get_website_settings()
    return success_response(_watermark_settings_to_dict(settings_obj), '获取成功')


@swagger_auto_schema(
    method='post',
    operation_description='更新水印设置',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        properties={
            'watermark': openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'enabled': openapi.Schema(type=openapi.TYPE_BOOLEAN, description='是否启用水印'),
                    'hiddenMode': openapi.Schema(type=openapi.TYPE_BOOLEAN, description='是否启用隐藏水印模式'),
                    'text': openapi.Schema(type=openapi.TYPE_STRING, description='水印文字'),
                    'fontSize': openapi.Schema(type=openapi.TYPE_INTEGER, description='水印字号'),
                    'opacity': openapi.Schema(type=openapi.TYPE_NUMBER, description='水印透明度'),
                    'rotate': openapi.Schema(type=openapi.TYPE_INTEGER, description='水印旋转角度')
                }
            ),
            'enabled': openapi.Schema(type=openapi.TYPE_BOOLEAN, description='是否启用水印（兼容字段）'),
            'hiddenMode': openapi.Schema(type=openapi.TYPE_BOOLEAN, description='是否启用隐藏水印模式（兼容字段）'),
            'text': openapi.Schema(type=openapi.TYPE_STRING, description='水印文字（兼容字段）'),
            'fontSize': openapi.Schema(type=openapi.TYPE_INTEGER, description='水印字号（兼容字段）'),
            'opacity': openapi.Schema(type=openapi.TYPE_NUMBER, description='水印透明度（兼容字段）'),
            'rotate': openapi.Schema(type=openapi.TYPE_INTEGER, description='水印旋转角度（兼容字段）')
        }
    ),
    responses={200: '更新成功', 400: '参数错误', 403: '权限不足'}
)
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def update_watermark_settings(request):
    if not _is_admin_user(request.user):
        return error_response('无权限修改水印设置', ERROR_CODE_FORBIDDEN, status.HTTP_403_FORBIDDEN)

    request_data = request.data or {}
    watermark_data = request_data.get('watermark') if isinstance(request_data, dict) else None
    if watermark_data is None:
        watermark_data = request_data

    payload, err_msg = _validate_watermark_payload(watermark_data)
    if not payload:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    with transaction.atomic():
        settings_obj = _get_website_settings_for_update()
        supports_hidden_mode = _supports_watermark_hidden_mode()

        is_changed = (
            settings_obj.watermark_enabled != payload['watermark_enabled']
            or (supports_hidden_mode and settings_obj.watermark_hidden_mode != payload['watermark_hidden_mode'])
            or settings_obj.watermark_text != payload['watermark_text']
            or settings_obj.watermark_font_size != payload['watermark_font_size']
            or settings_obj.watermark_opacity != payload['watermark_opacity']
            or settings_obj.watermark_rotate != payload['watermark_rotate']
        )
        if not is_changed:
            return success_response(True, '设置未发生变化')

        settings_obj.watermark_enabled = payload['watermark_enabled']
        settings_obj.watermark_text = payload['watermark_text']
        settings_obj.watermark_font_size = payload['watermark_font_size']
        settings_obj.watermark_opacity = payload['watermark_opacity']
        settings_obj.watermark_rotate = payload['watermark_rotate']

        update_fields = [
            'watermark_enabled',
            'watermark_text',
            'watermark_font_size',
            'watermark_opacity',
            'watermark_rotate',
            'updated_at'
        ]
        if supports_hidden_mode:
            settings_obj.watermark_hidden_mode = payload['watermark_hidden_mode']
            update_fields.insert(1, 'watermark_hidden_mode')

        settings_obj.save(update_fields=update_fields)

    return success_response(True, '更新成功')


@swagger_auto_schema(
    method='post',
    operation_description='更新网站设置',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['websiteName'],
        properties={
            'websiteName': openapi.Schema(type=openapi.TYPE_STRING, description='网站名称'),
            'logo': openapi.Schema(type=openapi.TYPE_STRING, description='网站 Logo URL'),
            'favicon': openapi.Schema(type=openapi.TYPE_STRING, description='网站图标 URL'),
            'watermark': openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'enabled': openapi.Schema(type=openapi.TYPE_BOOLEAN, description='是否启用水印'),
                    'hiddenMode': openapi.Schema(type=openapi.TYPE_BOOLEAN, description='是否启用隐藏水印模式'),
                    'text': openapi.Schema(type=openapi.TYPE_STRING, description='水印文字'),
                    'fontSize': openapi.Schema(type=openapi.TYPE_INTEGER, description='水印字号'),
                    'opacity': openapi.Schema(type=openapi.TYPE_NUMBER, description='水印透明度'),
                    'rotate': openapi.Schema(type=openapi.TYPE_INTEGER, description='水印旋转角度')
                }
            )
        }
    ),
    responses={200: '更新成功', 400: '参数错误', 403: '权限不足'}
)
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def update_website_settings(request):
    if not _is_admin_user(request.user):
        return error_response('无权限修改网站设置', ERROR_CODE_FORBIDDEN, status.HTTP_403_FORBIDDEN)

    payload, err_msg = _validate_website_settings_payload(request.data or {})
    if not payload:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    with transaction.atomic():
        settings_obj = _get_website_settings_for_update()
        previous_logo = settings_obj.logo
        previous_favicon = settings_obj.favicon
        supports_hidden_mode = _supports_watermark_hidden_mode()

        is_changed = (
            settings_obj.website_name != payload['website_name']
            or settings_obj.logo != payload['logo']
            or settings_obj.favicon != payload['favicon']
            or settings_obj.watermark_enabled != payload['watermark_enabled']
            or (supports_hidden_mode and settings_obj.watermark_hidden_mode != payload['watermark_hidden_mode'])
            or settings_obj.watermark_text != payload['watermark_text']
            or settings_obj.watermark_font_size != payload['watermark_font_size']
            or settings_obj.watermark_opacity != payload['watermark_opacity']
            or settings_obj.watermark_rotate != payload['watermark_rotate']
        )
        if not is_changed:
            return success_response(True, '设置未发生变化')

        settings_obj.website_name = payload['website_name']
        settings_obj.logo = payload['logo']
        settings_obj.favicon = payload['favicon']
        settings_obj.watermark_enabled = payload['watermark_enabled']
        settings_obj.watermark_text = payload['watermark_text']
        settings_obj.watermark_font_size = payload['watermark_font_size']
        settings_obj.watermark_opacity = payload['watermark_opacity']
        settings_obj.watermark_rotate = payload['watermark_rotate']

        update_fields = [
            'website_name',
            'logo',
            'favicon',
            'watermark_enabled',
            'watermark_text',
            'watermark_font_size',
            'watermark_opacity',
            'watermark_rotate',
            'updated_at'
        ]
        if supports_hidden_mode:
            settings_obj.watermark_hidden_mode = payload['watermark_hidden_mode']
            update_fields.insert(4, 'watermark_hidden_mode')

        settings_obj.save(update_fields=update_fields)

    _cleanup_stale_website_assets(previous_logo, previous_favicon, payload['logo'], payload['favicon'])

    return success_response(True, '更新成功')


@swagger_auto_schema(
    method='get',
    operation_description='系统设置-用户管理列表（auth_user）',
    manual_parameters=[
        openapi.Parameter('current', openapi.IN_QUERY, description='页码，默认1', type=openapi.TYPE_INTEGER),
        openapi.Parameter(
            'size',
            openapi.IN_QUERY,
            description=f'每页条数，默认{SYSTEM_MANAGE_USER_DEFAULT_SIZE}，最大{SYSTEM_MANAGE_USER_MAX_SIZE}',
            type=openapi.TYPE_INTEGER
        ),
        openapi.Parameter('userName', openapi.IN_QUERY, description='用户名（模糊匹配）', type=openapi.TYPE_STRING),
        openapi.Parameter('userEmail', openapi.IN_QUERY, description='邮箱（模糊匹配）', type=openapi.TYPE_STRING),
        openapi.Parameter('status', openapi.IN_QUERY, description='状态：1=可用，2=禁用', type=openapi.TYPE_STRING),
    ],
    responses={200: '获取成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_system_manage_users(request):
    current = _parse_positive_int(request.query_params.get('current'), default=1, max_value=100000)
    size = _parse_positive_int(
        request.query_params.get('size'),
        default=SYSTEM_MANAGE_USER_DEFAULT_SIZE,
        max_value=SYSTEM_MANAGE_USER_MAX_SIZE
    )
    username_keyword = _clean_username(request.query_params.get('userName') or request.query_params.get('username'))
    email_keyword = _to_str(request.query_params.get('userEmail') or request.query_params.get('email'))
    raw_status = request.query_params.get('status')

    normalized_status = None
    if raw_status is not None and str(raw_status).strip() != '':
        normalized_status = _normalize_system_manage_user_status(raw_status)
        if normalized_status is None:
            return error_response('status 参数无效，仅支持 1(可用)/2(禁用)', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        queryset = User.objects.all().order_by('-id')
        if username_keyword:
            queryset = queryset.filter(username__icontains=username_keyword)
        if email_keyword:
            queryset = queryset.filter(email__icontains=email_keyword)
        if normalized_status is not None:
            queryset = queryset.filter(is_active=normalized_status)

        total = queryset.count()
        start_index = (current - 1) * size
        end_index = start_index + size
        users = list(queryset[start_index:end_index])
    except Exception as exc:
        logger.exception('get system manage users failed: %s', exc)
        return error_response('用户列表查询失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    records = [_serialize_system_manage_user_row(user) for user in users]
    return success_response(
        {
            'current': current,
            'size': size,
            'total': total,
            'records': records,
        },
        '获取成功'
    )


@swagger_auto_schema(
    method='post',
    operation_description='系统设置-新增用户（auth_user）',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['userName', 'password', 'confirmPassword', 'status'],
        properties={
            'userName': openapi.Schema(type=openapi.TYPE_STRING, description='用户名'),
            'password': openapi.Schema(type=openapi.TYPE_STRING, description='密码'),
            'confirmPassword': openapi.Schema(type=openapi.TYPE_STRING, description='确认密码'),
            'email': openapi.Schema(type=openapi.TYPE_STRING, description='邮箱'),
            'status': openapi.Schema(type=openapi.TYPE_STRING, description='状态：1=可用，2=禁用'),
        }
    ),
    responses={200: '新增成功', 400: '参数错误', 401: '未认证'}
)
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def create_system_manage_user(request):
    payload, err_msg = _validate_system_manage_user_payload(request.data or {}, is_create=True)
    if payload is None:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    username = payload.get('username')
    try:
        if User.objects.filter(username__iexact=username).exists():
            return error_response('用户名已存在', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        user = User(
            email=payload.get('email') or '',
            is_active=bool(payload.get('is_active')),
            username=username
        )
        _apply_system_manage_user_password(user, payload.get('password'))
        user.save()
    except Exception as exc:
        logger.exception('create system manage user failed: %s', exc)
        return error_response('新增用户失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    return success_response(_serialize_system_manage_user_row(user), '新增成功', status.HTTP_201_CREATED)


@swagger_auto_schema(
    method='put',
    operation_description='系统设置-编辑用户（auth_user）',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['userName', 'status'],
        properties={
            'userName': openapi.Schema(type=openapi.TYPE_STRING, description='用户名'),
            'password': openapi.Schema(type=openapi.TYPE_STRING, description='密码（留空则不修改）'),
            'confirmPassword': openapi.Schema(type=openapi.TYPE_STRING, description='确认密码（留空则不修改）'),
            'email': openapi.Schema(type=openapi.TYPE_STRING, description='邮箱'),
            'status': openapi.Schema(type=openapi.TYPE_STRING, description='状态：1=可用，2=禁用'),
        }
    ),
    responses={200: '更新成功', 400: '参数错误', 404: '用户不存在', 401: '未认证'}
)
@api_view(['PUT'])
@permission_classes([IsAuthenticated])
def update_system_manage_user(request, user_id):
    payload, err_msg = _validate_system_manage_user_payload(request.data or {}, is_create=False)
    if payload is None:
        return error_response(err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        user = User.objects.filter(id=user_id).first()
        if user is None:
            return error_response('用户不存在', ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)

        username = payload.get('username')
        if User.objects.filter(username__iexact=username).exclude(id=user.id).exists():
            return error_response('用户名已存在', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        user.username = username
        user.email = payload.get('email') or ''
        user.is_active = bool(payload.get('is_active'))
        _apply_system_manage_user_password(user, payload.get('password'))
        user.save()
    except Exception as exc:
        logger.exception('update system manage user failed: %s', exc)
        return error_response('更新用户失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    return success_response(_serialize_system_manage_user_row(user), '更新成功')


@swagger_auto_schema(
    method='delete',
    operation_description='系统设置-删除用户（auth_user）',
    responses={200: '删除成功', 404: '用户不存在', 401: '未认证'}
)
@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def delete_system_manage_user(request, user_id):
    try:
        user = User.objects.filter(id=user_id).first()
        if user is None:
            return error_response('用户不存在', ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)
        if request.user and request.user.id == user.id:
            return error_response('不允许删除当前登录用户', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

        user.delete()
    except Exception as exc:
        logger.exception('delete system manage user failed: %s', exc)
        return error_response('删除用户失败，请检查数据库连接', ERROR_CODE_INVALID_PARAMS, status.HTTP_500_INTERNAL_SERVER_ERROR)

    return success_response(True, '删除成功')


@swagger_auto_schema(
    method='post',
    operation_description='用户登录',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['userName', 'password', 'captchaId', 'captchaCode'],
        properties={
            'userName': openapi.Schema(type=openapi.TYPE_STRING, description='用户名'),
            'password': openapi.Schema(type=openapi.TYPE_STRING, description='密码'),
            'captchaId': openapi.Schema(type=openapi.TYPE_STRING, description='验证码标识'),
            'captchaCode': openapi.Schema(type=openapi.TYPE_STRING, description='验证码')
        }
    ),
    responses={
        200: openapi.Response(
            '登录结果（成功或业务错误）',
            openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'code': openapi.Schema(type=openapi.TYPE_STRING),
                    'data': openapi.Schema(
                        type=openapi.TYPE_OBJECT,
                        properties={
                            'token': openapi.Schema(type=openapi.TYPE_STRING),
                            'refreshToken': openapi.Schema(type=openapi.TYPE_STRING)
                        }
                    ),
                    'msg': openapi.Schema(type=openapi.TYPE_STRING)
                }
            )
        ),
        400: '参数错误',
        403: '用户被禁用'
    }
)
@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def user_login(request):
    username = _clean_username(request.data.get('userName'))
    password = request.data.get('password') or ''
    captcha_id = (request.data.get('captchaId') or '').strip()
    captcha_code = (request.data.get('captchaCode') or '').strip()
    client_ip = _extract_client_ip(request)
    login_subject = _build_login_subject(username, client_ip)

    validation_error = _validate_login_payload(username, password, captcha_id)
    if validation_error:
        msg, code, http_status = validation_error
        return error_response(msg, code, http_status)

    lock_remaining = _get_login_lock_remaining(login_subject)
    if lock_remaining:
        logger.warning('login blocked by lock, subject=%s, ip=%s, remaining=%s', login_subject, client_ip, lock_remaining)
        return error_response(f'登录失败次数过多，请{lock_remaining}秒后重试', ERROR_CODE_AUTH_FAILED, status.HTTP_200_OK)

    is_captcha_valid, captcha_err_msg = _verify_captcha(captcha_id, captcha_code)
    if not is_captcha_valid:
        logger.warning('login failed captcha, subject=%s, ip=%s, reason=%s', login_subject, client_ip, captcha_err_msg)
        return error_response(captcha_err_msg, ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    matched_user = User.objects.filter(username=username).first()
    if not matched_user:
        # Fall back to case-insensitive lookup for legacy users while preferring stable ordering.
        matched_user = User.objects.filter(username__iexact=username).order_by('id').first()
    auth_username = matched_user.username if matched_user else username
    user = authenticate(username=auth_username, password=password)
    if not user:
        lock_seconds = _register_login_failure(login_subject)
        logger.warning('login failed credential, subject=%s, ip=%s', login_subject, client_ip)
        if lock_seconds:
            return error_response(
                f'账号或密码错误，已临时锁定{lock_seconds}秒，请稍后重试',
                ERROR_CODE_AUTH_FAILED,
                status.HTTP_200_OK
            )
        # Keep login credential mismatch as business error under HTTP 200 for frontend contract compatibility.
        return error_response('账号或密码错误', ERROR_CODE_AUTH_FAILED, status.HTTP_200_OK)

    if not user.is_active:
        logger.warning('login rejected inactive user, uid=%s, ip=%s', user.id, client_ip)
        return error_response('用户已被禁用', ERROR_CODE_FORBIDDEN, status.HTTP_403_FORBIDDEN)

    _clear_login_failure_state(login_subject)
    logger.info('login success, uid=%s, subject=%s, ip=%s', user.id, login_subject, client_ip)

    return success_response(_issue_auth_tokens(user), '登录成功')


@swagger_auto_schema(
    method='post',
    operation_description='用户注册',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['username', 'password'],
        properties={
            'username': openapi.Schema(type=openapi.TYPE_STRING, description='用户名'),
            'userName': openapi.Schema(type=openapi.TYPE_STRING, description='用户名（兼容字段）'),
            'password': openapi.Schema(type=openapi.TYPE_STRING, description='密码'),
            'email': openapi.Schema(type=openapi.TYPE_STRING, description='邮箱（可选）')
        }
    ),
    responses={
        201: openapi.Response(
            '注册成功',
            openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'code': openapi.Schema(type=openapi.TYPE_STRING),
                    'data': openapi.Schema(
                        type=openapi.TYPE_OBJECT,
                        properties={
                            'token': openapi.Schema(type=openapi.TYPE_STRING),
                            'refreshToken': openapi.Schema(type=openapi.TYPE_STRING)
                        }
                    ),
                    'msg': openapi.Schema(type=openapi.TYPE_STRING)
                }
            )
        ),
        400: '参数错误或用户已存在'
    }
)
@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def user_register(request):
    username = _clean_username(request.data.get('username') or request.data.get('userName'))
    password = request.data.get('password') or ''
    email = (request.data.get('email') or '').strip()

    if not username or not password:
        return error_response('用户名和密码不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    if User.objects.filter(username__iexact=username).exists():
        return error_response('用户名已存在', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        validate_password(password)
    except ValidationError as exc:
        return error_response(exc.messages[0], ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    user = User.objects.create_user(username=username, password=password, email=email)

    return success_response(_issue_auth_tokens(user), '注册成功', status.HTTP_201_CREATED)


@swagger_auto_schema(
    method='post',
    operation_description='刷新登录令牌',
    request_body=openapi.Schema(
        type=openapi.TYPE_OBJECT,
        required=['refreshToken'],
        properties={
            'refreshToken': openapi.Schema(type=openapi.TYPE_STRING, description='刷新令牌')
        }
    ),
    responses={
        200: '刷新成功',
        401: '刷新令牌无效或过期'
    }
)
@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def refresh_token(request):
    refresh_token_value = request.data.get('refreshToken') or ''
    if not refresh_token_value:
        return error_response('refreshToken 不能为空', ERROR_CODE_INVALID_PARAMS, status.HTTP_400_BAD_REQUEST)

    try:
        payload = _parse_refresh_token(refresh_token_value)
    except signing.BadSignature:
        return error_response('refreshToken 无效或已过期', ERROR_CODE_AUTH_FAILED, status.HTTP_401_UNAUTHORIZED)

    user = User.objects.filter(id=payload.get('uid')).first()
    if not user:
        return error_response('用户不存在', ERROR_CODE_NOT_FOUND, status.HTTP_404_NOT_FOUND)

    if not user.is_active:
        return error_response('用户已被禁用', ERROR_CODE_FORBIDDEN, status.HTTP_403_FORBIDDEN)

    password_fingerprint = hashlib.sha256(user.password.encode('utf-8')).hexdigest()
    if payload.get('passwordFingerprint') != password_fingerprint:
        return error_response('refreshToken 已失效，请重新登录', ERROR_CODE_AUTH_FAILED, status.HTTP_401_UNAUTHORIZED)

    return success_response(_issue_auth_tokens(user), '刷新成功')


@swagger_auto_schema(
    method='get',
    operation_description='获取当前登录用户信息',
    responses={
        200: openapi.Response(
            '获取成功',
            openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'code': openapi.Schema(type=openapi.TYPE_STRING),
                    'data': openapi.Schema(
                        type=openapi.TYPE_OBJECT,
                        properties={
                            'buttons': openapi.Schema(type=openapi.TYPE_ARRAY, items=openapi.Items(type=openapi.TYPE_STRING)),
                            'roles': openapi.Schema(type=openapi.TYPE_ARRAY, items=openapi.Items(type=openapi.TYPE_STRING)),
                            'userId': openapi.Schema(type=openapi.TYPE_STRING),
                            'userName': openapi.Schema(type=openapi.TYPE_STRING)
                        }
                    ),
                    'msg': openapi.Schema(type=openapi.TYPE_STRING)
                }
            )
        ),
        401: '未认证'
    }
)
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_user_info(request):
    return success_response(_build_user_info(request.user), '获取用户信息成功')


@swagger_auto_schema(
    method='post',
    operation_description='用户登出',
    responses={
        200: openapi.Response(
            '登出成功',
            openapi.Schema(
                type=openapi.TYPE_OBJECT,
                properties={
                    'code': openapi.Schema(type=openapi.TYPE_STRING),
                    'data': openapi.Schema(type=openapi.TYPE_OBJECT),
                    'msg': openapi.Schema(type=openapi.TYPE_STRING)
                }
            )
        )
    }
)
@api_view(['POST'])
@authentication_classes([])
@permission_classes([AllowAny])
def user_logout(request):
    token = _extract_auth_token(request)
    if token:
        Token.objects.filter(key=token).delete()

    return success_response({}, '退出成功')


@swagger_auto_schema(
    method='get',
    operation_description='返回自定义后端错误（用于前端调试）',
    manual_parameters=[
        openapi.Parameter('code', openapi.IN_QUERY, description='错误码', type=openapi.TYPE_STRING, required=True),
        openapi.Parameter('msg', openapi.IN_QUERY, description='错误信息', type=openapi.TYPE_STRING, required=True)
    ],
    responses={400: '自定义错误'}
)
@api_view(['GET'])
@authentication_classes([])
@permission_classes([AllowAny])
def custom_error(request):
    code = str(request.query_params.get('code') or ERROR_CODE_INVALID_PARAMS)
    msg = (request.query_params.get('msg') or '自定义错误').strip()

    return error_response(msg, code=code, http_status=status.HTTP_400_BAD_REQUEST)
